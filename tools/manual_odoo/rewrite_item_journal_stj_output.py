#!/usr/bin/env python3
"""Guarded manual fix: rewrite Item Journal output columns from one picking."""

from __future__ import annotations

import argparse
import asyncio
import json
from pathlib import Path
import shutil
from typing import Any, Sequence

from openpyxl import load_workbook

from common import (
    build_move_to_journal_map,
    fetch_move_ids_for_picking,
    fetch_move_rows_with_product_keys,
    fetch_picking,
    m2o_id,
    normalize_qty,
    open_manual_connection,
    read_journal_names,
    write_manual_artifact,
    add_manual_common_args,
)
from smartscc_tools.services.odoo.gateway import build_company_context


DEFAULT_HEADERS = {
    "product_key": "Product key (kode/nama)",
    "qty": "Qty",
    "uom": "UoM",
    "out_picking": "Output No Internal Transfer",
    "out_stj": "Output No Item Journal",
    "out_error": "Output Error",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Rewrite Item Journal output cells from one successful stock picking.")
    parser.add_argument("--workbook", required=True, help="Workbook path.")
    parser.add_argument("--picking", required=True, help="Stock picking number.")
    parser.add_argument("--sheet", default="Item Journal", help="Worksheet name.")
    parser.add_argument("--backup", default="", help="Backup workbook path. Default: beside workbook with .before_<picking>.")
    add_manual_common_args(parser)
    return parser.parse_args()


def workbook_header_map(workbook_path: Path, sheet_name: str) -> dict[str, int]:
    wb = load_workbook(workbook_path, read_only=True, keep_vba=True, data_only=False)
    try:
        ws = wb[sheet_name]
        values = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        return {str(value or "").strip(): index for index, value in enumerate(values, start=1) if str(value or "").strip()}
    finally:
        wb.close()


def collect_error_rows(workbook_path: Path, sheet_name: str, picking: str, columns: dict[str, int]) -> list[dict[str, Any]]:
    wb = load_workbook(workbook_path, read_only=True, keep_vba=True, data_only=False)
    try:
        ws = wb[sheet_name]
        rows: list[dict[str, Any]] = []
        for row_idx, row_values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            out_picking = str(row_values[columns["out_picking"] - 1] or "").strip()
            out_stj = str(row_values[columns["out_stj"] - 1] or "").strip()
            out_error = str(row_values[columns["out_error"] - 1] or "").strip()
            if out_picking != "[Error]" or out_stj or picking not in out_error:
                continue
            rows.append(
                {
                    "row": row_idx,
                    "prod_key": str(row_values[columns["product_key"] - 1] or "").strip(),
                    "qty": normalize_qty(row_values[columns["qty"] - 1]),
                    "uom": str(row_values[columns["uom"] - 1] or "").strip(),
                    "existing_error": out_error,
                }
            )
        return rows
    finally:
        wb.close()


def compare_sequences(excel_rows: Sequence[dict[str, Any]], move_rows: Sequence[dict[str, Any]]) -> list[dict[str, Any]]:
    excel_keys = [(row["prod_key"], row["qty"], row["uom"]) for row in excel_rows]
    asc_keys = [(row["prod_key"], row["qty"], row["uom"]) for row in move_rows]
    desc_rows = list(reversed(move_rows))
    desc_keys = [(row["prod_key"], row["qty"], row["uom"]) for row in desc_rows]
    if excel_keys == asc_keys:
        return list(move_rows)
    if excel_keys == desc_keys:
        return desc_rows
    for index, (excel_key, move_key) in enumerate(zip(excel_keys, asc_keys), start=1):
        if excel_key != move_key:
            raise RuntimeError(f"Mismatch row ke-{index}: excel={excel_key} odoo={move_key}")
    raise RuntimeError("Urutan row Excel tidak cocok dengan stock.move Odoo.")


def backup_path_for(workbook_path: Path, picking: str, raw_backup: str) -> Path:
    if raw_backup:
        return Path(raw_backup).expanduser().resolve()
    suffix = workbook_path.suffix or ".xlsm"
    safe = picking.replace("/", "_").replace("\\", "_")
    return workbook_path.with_name(f"{workbook_path.stem}.before_{safe}{suffix}")


def write_via_excel(workbook_path: Path, sheet_name: str, columns: dict[str, int], mapping_rows: Sequence[dict[str, Any]]) -> None:
    try:
        import pythoncom
        import win32com.client
    except ImportError as exc:
        raise RuntimeError("pywin32 belum tersedia. Install pywin32 sebelum apply workbook rewrite.") from exc

    pythoncom.CoInitialize()
    excel = None
    workbook = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        workbook = excel.Workbooks.Open(str(workbook_path), UpdateLinks=0, ReadOnly=False)
        sheet = workbook.Worksheets(sheet_name)
        output_width = int(columns["out_error"]) - int(columns["out_picking"]) + 1
        if output_width < 3:
            raise RuntimeError(f"Layout output tidak dikenali. Lebar output={output_width}")
        for item in mapping_rows:
            row = int(item["row"])
            values = [["" for _ in range(output_width)]]
            values[0][0] = str(item["picking"])
            values[0][1] = str(item["journal_name"])
            values[0][-1] = ""
            target = sheet.Range(sheet.Cells(row, columns["out_picking"]), sheet.Cells(row, columns["out_error"]))
            target.Value = values
        workbook.Save()
    finally:
        if workbook is not None:
            workbook.Close(SaveChanges=False)
        if excel is not None:
            excel.Quit()
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass


async def amain() -> int:
    args = parse_args()
    workbook_path = Path(args.workbook).expanduser().resolve()
    if not workbook_path.exists():
        raise SystemExit(f"Workbook tidak ditemukan: {workbook_path}")
    conn = await open_manual_connection(args)
    try:
        header_map = workbook_header_map(workbook_path, args.sheet)
        missing = [label for label in DEFAULT_HEADERS.values() if label not in header_map]
        if missing:
            raise RuntimeError(f"Header workbook tidak ditemukan: {missing}")
        columns = {key: header_map[label] for key, label in DEFAULT_HEADERS.items()}
        excel_rows = collect_error_rows(workbook_path, args.sheet, args.picking, columns)

        picking = await fetch_picking(conn.client, args.picking)
        company_id = m2o_id(picking.get("company_id"))
        context = build_company_context(company_id)
        move_ids = await fetch_move_ids_for_picking(conn.client, int(picking["id"]), context)
        if len(excel_rows) != len(move_ids):
            raise RuntimeError(f"Jumlah row Excel ({len(excel_rows)}) tidak sama dengan stock.move Odoo ({len(move_ids)}).")
        move_rows = await fetch_move_rows_with_product_keys(conn.client, move_ids, context)
        ordered_move_rows = compare_sequences(excel_rows, move_rows)
        move_to_journal = await build_move_to_journal_map(conn.client, move_ids, context)
        journal_names = await read_journal_names(conn.client, sorted(set(move_to_journal.values())), context)

        mapping_rows: list[dict[str, Any]] = []
        for excel_row, move_row in zip(excel_rows, ordered_move_rows):
            journal_id = move_to_journal.get(int(move_row["move_id"]), 0)
            journal_name = journal_names.get(journal_id, "")
            if not journal_name:
                raise RuntimeError(f"Journal STJ tidak ditemukan untuk stock.move {move_row['move_id']}.")
            mapping_rows.append({**excel_row, **move_row, "picking": args.picking, "journal_id": journal_id, "journal_name": journal_name})

        plan = {
            "action": "rewrite_item_journal_stj_output",
            "mode": "apply" if args.apply else "dry_run",
            "database": conn.config.database,
            "workbook": str(workbook_path),
            "sheet": args.sheet,
            "picking": args.picking,
            "row_count": len(mapping_rows),
            "mapping_rows": mapping_rows,
        }
        if args.apply:
            backup_path = backup_path_for(workbook_path, args.picking, args.backup)
            backup_path.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(workbook_path, backup_path)
            write_via_excel(workbook_path, args.sheet, columns, mapping_rows)
            plan["backup"] = str(backup_path)

        print(json.dumps(plan, ensure_ascii=False, indent=2, default=str))
        path = write_manual_artifact(plan, args, f"rewrite_item_journal_{args.picking}")
        if path:
            print(f"Artefact: {path}")
        return 0
    finally:
        await conn.close()


if __name__ == "__main__":
    raise SystemExit(asyncio.run(amain()))

