#!/usr/bin/env python3
"""Build a reusable Purchase Cycle Balance report for a target bill list.

This is a read-only wrapper around the existing SVL Dashboard service.  It is
kept under tools/ so it can be called directly today and wrapped by external
AI adapters later without duplicating business logic.
"""

from __future__ import annotations

import argparse
import asyncio
import json
from collections import Counter, defaultdict
from dataclasses import asdict, is_dataclass
from datetime import datetime
from pathlib import Path
import sys
from typing import Any, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from smartscc_tools.core.global_config import GlobalPersistentState
from smartscc_tools.features.item_journal.utils import normalize_text
from smartscc_tools.features.svl_fix_je.config import PURCHASE_CYCLE_BALANCE_DATASET_MODE
from smartscc_tools.features.svl_fix_je.dashboard_service import SvlDashboardServiceAsync
from smartscc_tools.features.svl_fix_je.models import SvlDashboardRequest, SvlDashboardSnapshot
from smartscc_tools.services.odoo.profiles import FOLLOW_GLOBAL_PROFILE_ID
from tools.odoo_inspector.common import (
    add_common_connection_args,
    build_base_payload,
    open_tool_connection,
    write_json_artifact,
)


NAVY = "1F4E79"
BLUE = "5B9BD5"
GREEN = "70AD47"
PURPLE = "7030A0"
LIGHT_BLUE = "DDEBF7"
LIGHT_GREEN = "E2F0D9"
LIGHT_AMBER = "FFF2CC"
LIGHT_RED = "FCE4D6"
WHITE = "FFFFFF"
TEXT = "1F1F1F"

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _norm(value: Any) -> str:
    return normalize_text(value).strip()


def _text_values(values: Any) -> list[str]:
    if values is None:
        return []
    if isinstance(values, str):
        value = _norm(values)
        return [value] if value else []
    try:
        iterator = iter(values)
    except TypeError:
        value = _norm(values)
        return [value] if value else []
    result: list[str] = []
    for value in iterator:
        text = _norm(value)
        if text:
            result.append(text)
    return result


def _join(values: Any, sep: str = ", ") -> str:
    return sep.join(_text_values(values))


def _unique_texts(values: Any) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in _text_values(values):
        key = value.upper()
        if key in seen:
            continue
        seen.add(key)
        result.append(value)
    return result


def _money(value: Any) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _safe_stem(value: str) -> str:
    text = _norm(value)
    for char in '\\/:*?"<>|':
        text = text.replace(char, "_")
    return text.strip(" ._") or "purchase_cycle_balance"


def _json_ready(value: Any) -> Any:
    if is_dataclass(value):
        return _json_ready(asdict(value))
    if isinstance(value, dict):
        return {str(key): _json_ready(val) for key, val in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_ready(item) for item in value]
    if isinstance(value, (set, frozenset)):
        return [_json_ready(item) for item in sorted(value, key=lambda item: repr(item))]
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _parse_codes(value: str, fallback: str) -> frozenset[str]:
    source = _norm(value) or _norm(fallback)
    return frozenset(part.strip() for part in source.split(",") if part.strip())


def _parse_bill_values(raw_values: Sequence[str], bill_file: str) -> list[str]:
    values: list[str] = []
    if bill_file:
        path = Path(bill_file).expanduser()
        if not path.is_absolute():
            path = PROJECT_ROOT / path
        values.extend(path.read_text(encoding="utf-8").splitlines())
    values.extend(raw_values or [])
    seen: set[str] = set()
    clean_values: list[str] = []
    for value in values:
        clean = _norm(value).replace("\ufeff", "").upper()
        if not clean or clean in seen:
            continue
        seen.add(clean)
        clean_values.append(clean)
    if not clean_values:
        raise SystemExit("Isi --bill-file atau minimal satu --bill.")
    return clean_values


def _configured_inventory_coa_codes() -> list[str]:
    codes: list[str] = []
    seen: set[str] = set()
    global_settings = GlobalPersistentState().load()
    for entry in getattr(global_settings, "inventory_coa_entries", []) or []:
        code = _norm(getattr(entry, "coa_code", "")).upper()
        if not code or code in seen:
            continue
        seen.add(code)
        codes.append(code)
    return codes


def _target_bill_hits(cycle: dict[str, Any], targets: set[str]) -> set[str]:
    hits: set[str] = set()
    for value in cycle.get("bill_refs", []) or []:
        clean = _norm(value).upper()
        if clean in targets:
            hits.add(clean)
    for item in cycle.get("item_rows", []) or []:
        for value in item.get("bill_refs", []) or []:
            clean = _norm(value).upper()
            if clean in targets:
                hits.add(clean)
    for line in cycle.get("raw_lines", []) or []:
        for value in line.values():
            clean = _norm(value).upper()
            if clean in targets:
                hits.add(clean)
    return hits


def _filter_snapshot(snapshot: SvlDashboardSnapshot, target_bills: list[str]) -> dict[str, Any]:
    payload = _json_ready(snapshot)
    targets = {bill.upper() for bill in target_bills}
    matched_cycles: list[dict[str, Any]] = []
    found: set[str] = set()
    for cycle in payload.get("purchase_cycles", []) or []:
        hits = _target_bill_hits(cycle, targets)
        if not hits:
            continue
        cycle["_target_bill_hits"] = sorted(hits)
        matched_cycles.append(cycle)
        found.update(hits)
    payload["purchase_cycles"] = matched_cycles
    payload["target_bills"] = list(target_bills)
    payload["found_target_bills"] = sorted(found)
    payload["missing_target_bills"] = sorted(targets - found)
    return payload


async def _fetch_bill_headers(
    *,
    service: SvlDashboardServiceAsync,
    company_id: int,
    target_bills: list[str],
    context: dict[str, Any],
) -> list[dict[str, Any]]:
    account_move_fields = await service._fields_get_cached("account.move")  # noqa: SLF001 - tool-side schema compatibility
    fields = [
        field
        for field in [
            "id",
            "name",
            "date",
            "invoice_date",
            "invoice_origin",
            "move_type",
            "state",
            "partner_id",
            "amount_total",
            "amount_residual",
            "payment_state",
        ]
        if field == "id" or field in account_move_fields
    ]
    return await service.rpc.search_read(
        "account.move",
        [("company_id", "=", int(company_id)), ("name", "in", list(target_bills))],
        fields=fields,
        limit=max(1, len(target_bills) + 10),
        context=context,
        stage="PCB_TARGET_BILL_HEADERS",
    )


def _status_fill(status: str) -> str:
    clean = _norm(status).lower()
    if clean == "problem":
        return LIGHT_RED
    if clean == "partial":
        return LIGHT_AMBER
    if clean == "healthy":
        return LIGHT_GREEN
    return WHITE


def _account_map(rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    return {_norm(row.get("code")): row for row in rows or [] if _norm(row.get("code"))}


def _account_net(rows: list[dict[str, Any]], code: str) -> float:
    return _money(_account_map(rows).get(code, {}).get("net_balance"))


def _problem_account_summary(rows: list[dict[str, Any]]) -> str:
    parts: list[str] = []
    for row in rows or []:
        status = _norm(row.get("status")).lower()
        net = _money(row.get("net_balance"))
        if status == "problem" or abs(net) >= 0.005:
            code = _norm(row.get("code"))
            if code:
                parts.append(f"{code} {net:,.2f}")
    return "; ".join(parts)


def _set_title(ws, title: str, subtitle: str, last_col: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    cell = ws.cell(row=1, column=1, value=title)
    cell.font = Font(name="Calibri", bold=True, size=14, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    sub = ws.cell(row=2, column=1, value=subtitle)
    sub.font = Font(name="Calibri", italic=True, size=10, color=TEXT)
    sub.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    sub.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 22


def _write_table(ws, start_row: int, headers: list[str], rows: list[list[Any]], *, header_fill: str = BLUE) -> None:
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=header)
        cell.font = Font(name="Calibri", bold=True, size=10, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=header_fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    for row_offset, row in enumerate(rows, start=1):
        excel_row = start_row + row_offset
        for col, value in enumerate(row, start=1):
            cell = ws.cell(row=excel_row, column=col, value=value)
            cell.font = Font(name="Calibri", size=9, color=TEXT)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
    if rows:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(rows)}"


def _format_sheet(ws, widths: dict[int, float], *, freeze: str = "A4") -> None:
    ws.freeze_panes = freeze
    ws.sheet_view.showGridLines = False
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, (int, float)):
                header = _norm(ws.cell(row=3, column=cell.column).value).lower()
                if any(token in header for token in ("amount", "balance", "debit", "credit", "net", "saldo", "total", "value", "price", "nilai")):
                    cell.number_format = "#,##0.00"
                    cell.alignment = Alignment(horizontal="right", vertical="top")
                elif any(token in header for token in ("qty", "quantity")):
                    cell.number_format = "#,##0.####"
                    cell.alignment = Alignment(horizontal="right", vertical="top")


def _apply_status_rows(ws, status_col: int, first_data_row: int, last_data_row: int) -> None:
    for row in range(first_data_row, last_data_row + 1):
        fill = PatternFill("solid", fgColor=_status_fill(ws.cell(row=row, column=status_col).value))
        for col in range(1, ws.max_column + 1):
            ws.cell(row=row, column=col).fill = fill


def _bill_po_refs_from_headers(bill_headers: list[dict[str, Any]]) -> dict[str, list[str]]:
    refs_by_bill: dict[str, list[str]] = {}
    for row in bill_headers or []:
        bill = _norm(row.get("name")).upper()
        origin = _norm(row.get("invoice_origin"))
        if not bill or not origin:
            continue
        refs_by_bill.setdefault(bill, [])
        refs_by_bill[bill].extend(_unique_texts([origin]))
    return {bill: _unique_texts(refs) for bill, refs in refs_by_bill.items()}


def _bill_context_po_refs(
    *,
    line: dict[str, Any],
    cycle_target_bills: list[str],
    bill_po_refs_by_name: dict[str, list[str]],
) -> list[str]:
    refs: list[str] = []
    transaction = _norm(line.get("kode_transaksi")).upper()
    if transaction:
        refs.extend(bill_po_refs_by_name.get(transaction, []))
    for bill in cycle_target_bills or []:
        refs.extend(bill_po_refs_by_name.get(_norm(bill).upper(), []))
    return _unique_texts(refs)


def _po_source_label(*, raw_po: str, bill_po_refs: list[str], cycle_po_refs: list[str], po_enrich_company_id: int) -> str:
    if raw_po:
        return "row_item_po"
    if bill_po_refs:
        return f"bill_company_{po_enrich_company_id}_po" if po_enrich_company_id > 0 else "bill_po"
    if cycle_po_refs:
        return "cycle_bill_po"
    return ""


def _build_excel_report(
    *,
    payload: dict[str, Any],
    bill_headers: list[dict[str, Any]],
    po_enrich_bill_headers: list[dict[str, Any]],
    po_enrich_company_id: int,
    output_path: Path,
) -> None:
    cycles = payload.get("purchase_cycles", []) or []
    target_bills = payload.get("target_bills", []) or []
    bill_header_map = {_norm(row.get("name")).upper(): row for row in bill_headers}
    po_enrich_bill_header_map = {_norm(row.get("name")).upper(): row for row in po_enrich_bill_headers}
    po_enrich_refs_by_bill = _bill_po_refs_from_headers(po_enrich_bill_headers)

    cycles_by_bill: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for cycle in cycles:
        for bill in cycle.get("_target_bill_hits", []) or []:
            cycles_by_bill[_norm(bill).upper()].append(cycle)

    cycle_rows: list[list[Any]] = []
    item_rows: list[list[Any]] = []
    raw_rows: list[list[Any]] = []
    raw_po_rows: list[list[Any]] = []

    for cycle in cycles:
        account_rows = cycle.get("account_rows", []) or []
        target_hits = _join(cycle.get("_target_bill_hits", []))
        cycle_rows.append([
            target_hits,
            _norm(cycle.get("picking_name")),
            _join(cycle.get("picking_names")),
            _norm(cycle.get("gr_date"))[:10],
            _norm(cycle.get("partner_name")),
            _join(cycle.get("purchase_orders")),
            _join(cycle.get("bill_refs")),
            _norm(cycle.get("cycle_status")),
            _norm(cycle.get("document_classification_label")) or _norm(cycle.get("document_classification")),
            _norm(cycle.get("primary_case")) or "-",
            _join(cycle.get("issue_patterns"), "; "),
            len(cycle.get("item_rows", []) or []),
            _money(cycle.get("total_debit")),
            _money(cycle.get("total_credit")),
            int(cycle.get("problem_account_count") or 0),
            _account_net(account_rows, "2103006"),
            _account_net(account_rows, "1108099"),
            _account_net(account_rows, "11120003"),
            _account_net(account_rows, "1105003"),
            _account_net(account_rows, "2102002"),
            _problem_account_summary(account_rows),
        ])
        for item in cycle.get("item_rows", []) or []:
            i_accounts = item.get("account_rows", []) or []
            item_rows.append([
                target_hits,
                _norm(cycle.get("picking_name")),
                _norm(cycle.get("gr_date"))[:10],
                _norm(cycle.get("partner_name")),
                _norm(item.get("default_code")),
                _norm(item.get("product_name")),
                _money(item.get("gr_quantity")),
                _money(item.get("bill_quantity")),
                _money(item.get("standard_price")),
                _join(item.get("bill_refs")),
                _join(item.get("stj_refs")),
                _norm(item.get("primary_case")) or "-",
                _norm(item.get("bill_hit_role")) or "-",
                "Y" if item.get("has_item_bill") else "N",
                "Y" if item.get("has_item_stj") else "N",
                _account_net(i_accounts, "2103006"),
                _account_net(i_accounts, "1108099"),
                _account_net(i_accounts, "11120003"),
                _account_net(i_accounts, "1105003"),
                _account_net(i_accounts, "2102002"),
                _problem_account_summary(i_accounts),
            ])
        cycle_po_refs = _text_values(cycle.get("purchase_orders", []))
        cycle_po_label = _join(cycle_po_refs)
        cycle_bill_label = _join(cycle.get("bill_refs", []))
        cycle_target_bills = _text_values(cycle.get("_target_bill_hits", []))
        for line in cycle.get("raw_lines", []) or []:
            raw_po = _norm(line.get("no_po"))
            bill_po_refs = _bill_context_po_refs(
                line=line,
                cycle_target_bills=cycle_target_bills,
                bill_po_refs_by_name=po_enrich_refs_by_bill,
            )
            po_fallback_refs = bill_po_refs or cycle_po_refs
            po_value = raw_po or _join(po_fallback_refs)
            po_source = _po_source_label(
                raw_po=raw_po,
                bill_po_refs=bill_po_refs,
                cycle_po_refs=cycle_po_refs,
                po_enrich_company_id=po_enrich_company_id,
            )
            row = [
                target_hits,
                _norm(cycle.get("picking_name")),
                _norm(cycle.get("cycle_status")),
                _norm(cycle.get("primary_case")) or "-",
                _norm(line.get("kode_transaksi")),
                _norm(line.get("tanggal"))[:10],
                _norm(line.get("jenis")),
                _norm(line.get("akun_code")),
                _norm(line.get("akun_name")),
                _norm(line.get("kode_item")),
                _norm(line.get("nama_item")),
                _norm(line.get("uom")),
                _money(line.get("qty_item")),
                _norm(line.get("kategori_produk")),
                po_value,
                _norm(line.get("partner")),
                _money(line.get("debit")),
                _money(line.get("kredit")),
                _money(line.get("saldo")),
                _norm(line.get("komunikasi")),
                _norm(line.get("matching")),
            ]
            raw_rows.append(row)
            raw_po_rows.append([*row, raw_po, po_source, cycle_po_label, cycle_bill_label, _join(bill_po_refs)])

    bill_rows: list[list[Any]] = []
    for bill in target_bills:
        bill_key = _norm(bill).upper()
        matched = cycles_by_bill.get(bill_key, [])
        header = bill_header_map.get(bill_key, {})
        po_enrich_header = po_enrich_bill_header_map.get(bill_key, {})
        bill_rows.append([
            bill_key,
            "covered" if matched else "not in filtered cycles",
            len(matched),
            _join([cycle.get("picking_name") for cycle in matched]),
            _join(sorted({_norm(cycle.get("cycle_status")) for cycle in matched if _norm(cycle.get("cycle_status"))})),
            _join(sorted({_norm(cycle.get("primary_case")) or "-" for cycle in matched})),
            _join(sorted({po for cycle in matched for po in _text_values(cycle.get("purchase_orders", [])) if _norm(po)})),
            _join(sorted({_norm(cycle.get("partner_name")) for cycle in matched if _norm(cycle.get("partner_name"))})),
            _norm(header.get("date"))[:10],
            _norm(header.get("invoice_date"))[:10],
            _norm(header.get("invoice_origin")),
            _norm(po_enrich_header.get("invoice_origin")),
            po_enrich_company_id if po_enrich_header else "",
            _money(header.get("amount_total")),
            _money(header.get("amount_residual")),
            _norm(header.get("payment_state")),
            _norm(header.get("state")),
            _norm(header.get("move_type")),
        ])

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    _set_title(
        ws,
        "Purchase Cycle Balance - Target Bills",
        "Filtered to requested bills; headless source: SvlDashboardServiceAsync.",
        8,
    )
    status_counts = Counter(_norm(cycle.get("cycle_status")) for cycle in cycles)
    case_counts = Counter(_norm(cycle.get("primary_case")) or "-" for cycle in cycles)
    kpis = [
        ["Company", payload.get("company_name"), "Company ID", payload.get("company_id")],
        ["Database", payload.get("database"), "Period", payload.get("period")],
        ["Target Bills", len(target_bills), "Covered Bills", len(payload.get("found_target_bills", []))],
        ["Related Cycles", len(cycles), "Problem Cycles", status_counts.get("problem", 0)],
        ["Healthy Cycles", status_counts.get("healthy", 0), "Partial Cycles", status_counts.get("partial", 0)],
        ["Dominant Case", case_counts.most_common(1)[0][0] if case_counts else "-", "Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
        ["PO Enrich Company ID", po_enrich_company_id or "", "PO Enrich Bills Found", len(po_enrich_bill_header_map)],
    ]
    _write_table(ws, 4, ["Metric", "Value", "Metric", "Value"], kpis, header_fill=NAVY)

    def _write_small_table(start_col: int, headers: list[str], rows: list[list[Any]], fill: str) -> None:
        for col_offset, header in enumerate(headers):
            cell = ws.cell(row=13, column=start_col + col_offset, value=header)
            cell.font = Font(name="Calibri", bold=True, size=10, color=WHITE)
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER
        for row_offset, row_values in enumerate(rows, start=1):
            for col_offset, value in enumerate(row_values):
                cell = ws.cell(row=13 + row_offset, column=start_col + col_offset, value=value)
                cell.font = Font(name="Calibri", size=9, color=TEXT)
                cell.border = BORDER
                cell.alignment = Alignment(vertical="center")

    _write_small_table(1, ["Status", "Cycle Count"], [[key or "-", value] for key, value in status_counts.most_common()], GREEN)
    _write_small_table(4, ["Case", "Cycle Count"], [[key or "-", value] for key, value in case_counts.most_common()], PURPLE)
    for col, width in {"A": 18, "B": 24, "C": 18, "D": 18, "E": 14}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False

    ws_bill = wb.create_sheet("Bill Coverage")
    _set_title(ws_bill, "Bill Coverage", "One row per requested bill.", 18)
    _write_table(
        ws_bill,
        3,
        [
            "Bill", "Coverage", "Cycle Count", "Cycle(s)", "Statuses", "Cases", "PO(s)", "Partner(s)",
            "Move Date", "Invoice Date", "Invoice Origin", "PO Enrich Origin", "PO Enrich Company ID",
            "Amount Total", "Amount Residual", "Payment State", "State", "Move Type",
        ],
        bill_rows,
        header_fill=BLUE,
    )
    _format_sheet(ws_bill, {1: 20, 2: 18, 3: 12, 4: 36, 5: 14, 6: 12, 7: 28, 8: 34, 9: 12, 10: 12, 11: 24, 12: 26, 13: 18, 14: 16, 15: 16, 16: 16, 17: 12, 18: 12})

    ws_cycle = wb.create_sheet("Cycle Overview")
    _set_title(ws_cycle, "Cycle Overview", "Only purchase cycles related to the requested bills.", 21)
    _write_table(
        ws_cycle,
        3,
        [
            "Target Bill Hits", "Picking", "All Pickings", "GR Date", "Vendor", "PO(s)", "Bill Refs",
            "Cycle Status", "Document Class", "Primary Case", "Issue Patterns", "Item Count", "Total Debit",
            "Total Credit", "Problem Account Count", "2103006 Net", "1108099 Net", "11120003 Net",
            "1105003 Net", "2102002 Net", "Non-zero / Problem Accounts",
        ],
        cycle_rows,
        header_fill=NAVY,
    )
    _apply_status_rows(ws_cycle, 8, 4, 3 + len(cycle_rows))
    _format_sheet(ws_cycle, {1: 28, 2: 22, 3: 32, 4: 12, 5: 34, 6: 30, 7: 36, 8: 14, 9: 18, 10: 12, 11: 46, 12: 10, 13: 16, 14: 16, 15: 14, 16: 16, 17: 16, 18: 16, 19: 16, 20: 16, 21: 48})

    ws_item = wb.create_sheet("Item Detail")
    _set_title(ws_item, "Item Detail", "All item rows inside related cycles.", 21)
    _write_table(
        ws_item,
        3,
        [
            "Target Bill Hits", "Picking", "GR Date", "Vendor", "Item Code", "Product", "GR Qty", "Bill Qty",
            "Standard Price", "Bill Refs Item", "STJ Refs Item", "Item Case", "Bill Hit Role", "Has Bill",
            "Has STJ", "2103006 Net", "1108099 Net", "11120003 Net", "1105003 Net", "2102002 Net",
            "Non-zero / Problem Accounts",
        ],
        item_rows,
        header_fill=GREEN,
    )
    _format_sheet(ws_item, {1: 28, 2: 20, 3: 12, 4: 34, 5: 18, 6: 42, 7: 12, 8: 12, 9: 16, 10: 30, 11: 30, 12: 12, 13: 14, 14: 10, 15: 10, 16: 16, 17: 16, 18: 16, 19: 16, 20: 16, 21: 48})

    ws_raw = wb.create_sheet("Raw Ledger")
    _set_title(ws_raw, "Raw Ledger", "Raw ledger lines for the related purchase cycles.", 21)
    _write_table(
        ws_raw,
        3,
        [
            "Target Bill Hits", "Picking", "Cycle Status", "Case", "Transaction", "Date", "Source",
            "Account Code", "Account Name", "Item Code", "Item Name", "UOM", "Qty", "Product Category",
            "PO", "Partner", "Debit", "Credit", "Balance", "Communication", "Matching",
        ],
        raw_rows,
        header_fill=PURPLE,
    )
    _apply_status_rows(ws_raw, 3, 4, 3 + len(raw_rows))
    _format_sheet(ws_raw, {1: 28, 2: 20, 3: 14, 4: 10, 5: 24, 6: 12, 7: 12, 8: 14, 9: 36, 10: 16, 11: 38, 12: 16, 13: 12, 14: 20, 15: 24, 16: 30, 17: 16, 18: 16, 19: 16, 20: 42, 21: 20})

    ws_raw_po = wb.create_sheet("Raw Ledger PO Bill")
    _set_title(ws_raw_po, "Raw Ledger PO Bill", "Raw ledger with PO completed from bill PO enrichment or cycle context when row-level PO is blank.", 26)
    _write_table(
        ws_raw_po,
        3,
        [
            "Target Bill Hits", "Picking", "Cycle Status", "Case", "Transaction", "Date", "Source",
            "Account Code", "Account Name", "Item Code", "Item Name", "UOM", "Qty", "Product Category",
            "PO", "Partner", "Debit", "Credit", "Balance", "Communication", "Matching",
            "PO Raw", "PO Source", "Cycle PO(s)", "Cycle Bill(s)", "Bill PO Enrich(s)",
        ],
        raw_po_rows,
        header_fill=PURPLE,
    )
    _apply_status_rows(ws_raw_po, 3, 4, 3 + len(raw_po_rows))
    _format_sheet(ws_raw_po, {1: 28, 2: 20, 3: 14, 4: 10, 5: 24, 6: 12, 7: 12, 8: 14, 9: 36, 10: 16, 11: 38, 12: 16, 13: 12, 14: 20, 15: 24, 16: 30, 17: 16, 18: 16, 19: 16, 20: 42, 21: 20, 22: 24, 23: 22, 24: 30, 25: 30, 26: 32})

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    load_workbook(output_path, read_only=True, data_only=False).close()


def _default_output_path(payload: dict[str, Any], target_count: int) -> Path:
    company = _safe_stem(_norm(payload.get("company_name")) or f"company_{payload.get('company_id')}")
    period = _safe_stem(_norm(payload.get("period")) or datetime.now().strftime("%Y%m%d"))
    name = f"Purchase Cycle Balance - {company} - {target_count} Target Bills - {period}.xlsx"
    return PROJECT_ROOT / "output" / name


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Build a read-only Purchase Cycle Balance report for selected bills.")
    parser.add_argument("--company-id", type=int, required=True, help="Odoo res.company id.")
    parser.add_argument("--date-from", default="", help="GR seed date from, YYYY-MM-DD.")
    parser.add_argument("--date-to", default="", help="GR seed date to, YYYY-MM-DD.")
    parser.add_argument("--bill-file", default="", help="Text file containing one bill number per line.")
    parser.add_argument("--bill", action="append", default=[], help="Bill number. Repeat as needed.")
    parser.add_argument("--problem-codes", default="2103006,1108099", help="Comma-separated PCB problem account codes.")
    parser.add_argument("--info-codes", default="11120003", help="Comma-separated PCB info account codes.")
    parser.add_argument(
        "--po-enrich-company-id",
        type=int,
        default=0,
        help="Optional company ID used only to fetch target bill invoice_origin as PO enrichment, e.g. 1 for POOL.",
    )
    parser.add_argument("--output", default="", help="Excel output path. Default: output/Purchase Cycle Balance - ....xlsx")
    parser.add_argument("--json-output", default="", help="Optional filtered JSON output path.")
    parser.add_argument("--strict", action="store_true", help="Exit non-zero if any requested bill is not covered by a cycle.")
    add_common_connection_args(parser)
    parser.set_defaults(database_profile=FOLLOW_GLOBAL_PROFILE_ID)
    return parser


async def amain(argv: Sequence[str] | None = None) -> int:
    parser = _build_parser()
    args = parser.parse_args(list(argv) if argv is not None else None)
    target_bills = _parse_bill_values(args.bill, args.bill_file)
    conn = await open_tool_connection(
        database_profile_id=args.database_profile,
        max_concurrency=args.max_concurrency,
        verbose=args.verbose,
    )
    try:
        service = SvlDashboardServiceAsync(rpc=conn.client, logger=conn.logger)
        context = {"allowed_company_ids": [int(args.company_id)]}
        request = SvlDashboardRequest(
            database=conn.config.database,
            company_id=int(args.company_id),
            date_from=_norm(args.date_from),
            date_to=_norm(args.date_to),
            inventory_coa_codes=_configured_inventory_coa_codes(),
            dataset_mode=PURCHASE_CYCLE_BALANCE_DATASET_MODE,
            include_inventory_accounts=True,
            include_non_inventory_accounts=True,
            pcb_problem_codes=_parse_codes(args.problem_codes, "2103006,1108099"),
            pcb_info_codes=_parse_codes(args.info_codes, "11120003"),
        )
        po_enrich_company_id = int(args.po_enrich_company_id or 0)
        po_enrich_context = {"allowed_company_ids": [po_enrich_company_id]} if po_enrich_company_id > 0 else context
        snapshot, bill_headers, po_enrich_bill_headers = await asyncio.gather(
            service.analyze(request),
            _fetch_bill_headers(
                service=service,
                company_id=int(args.company_id),
                target_bills=target_bills,
                context=context,
            ),
            _fetch_bill_headers(
                service=service,
                company_id=po_enrich_company_id,
                target_bills=target_bills,
                context=po_enrich_context,
            )
            if po_enrich_company_id > 0
            else asyncio.sleep(0, result=[]),
        )
        payload = _filter_snapshot(snapshot, target_bills)
        payload.update(
            {
                "source": "tools/odoo_reports/purchase_cycle_bill_report.py",
                "database_profile_id": args.database_profile,
                "bill_header_count": len(bill_headers),
                "po_enrich_company_id": po_enrich_company_id,
                "po_enrich_bill_header_count": len(po_enrich_bill_headers),
            }
        )
        output_path = Path(args.output).expanduser() if _norm(args.output) else _default_output_path(payload, len(target_bills))
        if not output_path.is_absolute():
            output_path = PROJECT_ROOT / output_path
        _build_excel_report(
            payload=payload,
            bill_headers=bill_headers,
            po_enrich_bill_headers=po_enrich_bill_headers,
            po_enrich_company_id=po_enrich_company_id,
            output_path=output_path,
        )

        json_path = None
        if _norm(args.json_output):
            json_path = Path(args.json_output).expanduser()
            if not json_path.is_absolute():
                json_path = PROJECT_ROOT / json_path
            json_path.parent.mkdir(parents=True, exist_ok=True)
            json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        elif not args.no_artifact:
            artifact_payload = build_base_payload(conn, action="purchase_cycle_bill_report")
            artifact_payload.update(payload)
            json_path = write_json_artifact(
                artifact_payload,
                output_dir=args.output_dir,
                prefix="purchase_cycle_bill_report",
            )

        summary = {
            "database": payload.get("database"),
            "company_id": payload.get("company_id"),
            "company_name": payload.get("company_name"),
            "period": payload.get("period"),
            "target_bills": len(target_bills),
            "covered_bills": len(payload.get("found_target_bills", [])),
            "missing_bills": payload.get("missing_target_bills", []),
            "related_cycles": len(payload.get("purchase_cycles", []) or []),
            "po_enrich_company_id": po_enrich_company_id,
            "po_enrich_bill_headers": len(po_enrich_bill_headers),
            "excel_output": str(output_path),
            "json_output": str(json_path) if json_path else "",
        }
        print(json.dumps(summary, ensure_ascii=True, indent=2, default=str))
        if args.strict and summary["missing_bills"]:
            return 2
        return 0
    finally:
        await conn.close()


def main(argv: Sequence[str] | None = None) -> int:
    return asyncio.run(amain(argv))


if __name__ == "__main__":
    raise SystemExit(main())
