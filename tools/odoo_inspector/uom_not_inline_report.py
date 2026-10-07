#!/usr/bin/env python3
"""Export read-only UoM-not-inline purchase/SVL investigation to Excel."""

from __future__ import annotations

import argparse
import asyncio
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation
import json
import random
from pathlib import Path
import zipfile
from typing import Any, Iterable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from common import (
    DEFAULT_ARTIFACT_DIR,
    add_common_connection_args,
    build_base_payload,
    open_tool_connection,
    parse_csv_values,
    write_json_artifact,
)


MODEL_UOM = "uom.uom"
BUSINESS_EQUIVALENT_UOM_PAIRS = {frozenset({"units", "pcs (p)"})}
VENDOR_BILL_MOVE_TYPES = {"in_invoice", "in_refund"}

PO_LINE_FIELDS = [
    "id",
    "display_name",
    "order_id",
    "company_id",
    "product_id",
    "product_qty",
    "product_uom_id",
    "product_uom_qty",
    "qty_received",
    "qty_invoiced",
    "price_unit",
    "price_subtotal",
    "price_total",
    "invoice_lines",
    "move_ids",
    "name",
    "date_planned",
    "state",
]
PRODUCT_FIELDS = ["id", "display_name", "name", "default_code", "uom_id", "standard_price", "categ_id"]
UOM_FIELDS = [
    "id",
    "display_name",
    "name",
    "active",
    "relative_factor",
    "relative_uom_id",
    "factor",
    "rounding",
    "parent_path",
]
BILL_LINE_FIELDS = [
    "id",
    "display_name",
    "move_id",
    "move_name",
    "move_type",
    "ref",
    "name",
    "purchase_line_id",
    "purchase_order_id",
    "product_id",
    "product_uom_id",
    "quantity",
    "price_unit",
    "price_subtotal",
    "debit",
    "credit",
    "balance",
    "account_id",
    "date",
    "parent_state",
    "company_id",
]
STOCK_MOVE_FIELDS = [
    "id",
    "reference",
    "name",
    "product_id",
    "product_uom",
    "product_uom_qty",
    "product_qty",
    "quantity",
    "picking_id",
    "purchase_line_id",
    "state",
    "date",
    "company_id",
]
SVL_FIELDS = [
    "id",
    "reference",
    "description",
    "product_id",
    "quantity",
    "unit_cost",
    "value",
    "remaining_qty",
    "remaining_value",
    "uom_id",
    "stock_move_id",
    "account_move_id",
    "account_move_line_id",
    "company_id",
    "warehouse_id",
    "create_date",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Read-only export for purchase/bill UoM lines that are not inline with product/SVL UoM. "
            "The report starts from upstream purchase.order.line and account.move.line."
        )
    )
    parser.add_argument(
        "--company-keywords",
        default="dragon,PIK,Livehouse,Mimi",
        help="Comma-separated company name keywords. Default: dragon,PIK,Livehouse,Mimi.",
    )
    parser.add_argument("--company-limit", type=int, default=15, help="Number of matching companies to sample.")
    parser.add_argument("--random-seed", type=int, default=20260407, help="Deterministic random seed.")
    parser.add_argument("--page-size", type=int, default=500, help="Read page size.")
    parser.add_argument(
        "--max-purchase-lines",
        type=int,
        default=0,
        help="Safety limit for purchase.order.line scan. 0 means no explicit limit.",
    )
    parser.add_argument("--date-from", default="", help="Optional purchase.order.line date_planned lower bound.")
    parser.add_argument("--date-to", default="", help="Optional purchase.order.line date_planned upper bound.")
    parser.add_argument("--product-code", default="A-BELG-0024", help="Product default_code for focused SVL sheet.")
    parser.add_argument(
        "--bekasi-company-keyword",
        default="DRAGON BEKASI",
        help="Company keyword for focused product SVL sheet. Default: DRAGON BEKASI.",
    )
    parser.add_argument(
        "--output",
        default="",
        help="Excel output path. Default: logs/odoo_inspector/UoM Not Inline Transactions - YYYY-MM-DD.xlsx",
    )
    add_common_connection_args(parser)
    return parser.parse_args()


def _id(value: Any) -> int:
    if isinstance(value, (list, tuple)) and value:
        return _id(value[0])
    try:
        return int(value or 0)
    except (TypeError, ValueError):
        return 0


def _name(value: Any) -> str:
    if isinstance(value, (list, tuple)) and len(value) >= 2:
        return str(value[1] or "")
    if value is False or value is None:
        return ""
    return str(value)


def _float(value: Any) -> float:
    try:
        return float(value or 0.0)
    except (TypeError, ValueError):
        return 0.0


def _decimal(value: Any) -> Decimal:
    try:
        return Decimal(str(value or "0"))
    except (InvalidOperation, TypeError, ValueError):
        return Decimal("0")


def _format_decimal(value: Decimal) -> str:
    normalized = value.normalize()
    if normalized == normalized.to_integral():
        return str(normalized.quantize(Decimal(1)))
    return format(normalized, "f")


def _ids_from_x2many(value: Any) -> list[int]:
    if not isinstance(value, list):
        return []
    return [item_id for item_id in (_id(item) for item in value) if item_id]


def _unique_ints(values: Iterable[Any]) -> list[int]:
    return sorted({item_id for item_id in (_id(value) for value in values) if item_id > 0})


def _parent_path_ids(row: dict[str, Any] | None) -> list[int]:
    if not row:
        return []
    ids: list[int] = []
    for part in str(row.get("parent_path") or "").split("/"):
        if not part.strip():
            continue
        try:
            ids.append(int(part))
        except ValueError:
            continue
    return ids


def _uom_root_id(row: dict[str, Any] | None) -> int:
    ids = _parent_path_ids(row)
    if ids:
        return ids[0]
    return _id((row or {}).get("id"))


def _uom_name(row: dict[str, Any] | None) -> str:
    if not row:
        return ""
    return str(row.get("display_name") or row.get("name") or f"#{row.get('id')}")


def _business_uom_key(row: dict[str, Any] | None) -> str:
    return _uom_name(row).casefold().strip()


def _business_equivalent(source_uom: dict[str, Any] | None, target_uom: dict[str, Any] | None) -> bool:
    if not source_uom or not target_uom:
        return False
    return frozenset({_business_uom_key(source_uom), _business_uom_key(target_uom)}) in BUSINESS_EQUIVALENT_UOM_PAIRS


def _is_inline(source_uom: dict[str, Any] | None, target_uom: dict[str, Any] | None) -> bool:
    source_root = _uom_root_id(source_uom)
    target_root = _uom_root_id(target_uom)
    return bool((source_root and target_root and source_root == target_root) or _business_equivalent(source_uom, target_uom))


def _convert(source_uom: dict[str, Any] | None, target_uom: dict[str, Any] | None, qty: Any) -> str:
    if not _is_inline(source_uom, target_uom):
        return ""
    if _business_equivalent(source_uom, target_uom) and _uom_root_id(source_uom) != _uom_root_id(target_uom):
        return _format_decimal(_decimal(qty))
    target_factor = _decimal((target_uom or {}).get("factor"))
    if not target_factor:
        return ""
    result = _decimal(qty) * _decimal((source_uom or {}).get("factor")) / target_factor
    return _format_decimal(result)


def _date_to_datetime_start(value: str) -> str:
    clean = str(value or "").strip()
    return f"{clean} 00:00:00" if clean and len(clean) == 10 else clean


def _date_to_datetime_end(value: str) -> str:
    clean = str(value or "").strip()
    return f"{clean} 23:59:59" if clean and len(clean) == 10 else clean


async def _fields(conn: Any, model: str) -> set[str]:
    meta = await conn.client.fields_get(model, attributes=["type"], stage=f"UOM_NI_FIELDS_{model}")
    return set(meta.keys())


def _filter_fields(requested: Sequence[str], available: set[str]) -> list[str]:
    return [field for field in requested if field in available]


async def _search_count(conn: Any, model: str, domain: list[Any]) -> int:
    result = await conn.client.execute_kw(
        model=model,
        method="search_count",
        args=[domain],
        kwargs={},
        stage=f"UOM_NI_COUNT_{model}",
        mutating=False,
    )
    return int(result or 0)


async def _search_read_paged(
    conn: Any,
    model: str,
    domain: list[Any],
    fields: Sequence[str],
    *,
    page_size: int,
    order: str = "id",
    max_rows: int = 0,
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    offset = 0
    while True:
        remaining = max_rows - len(rows) if max_rows else page_size
        if max_rows and remaining <= 0:
            break
        limit = min(page_size, remaining) if max_rows else page_size
        result = await conn.client.execute_kw(
            model=model,
            method="search_read",
            args=[domain],
            kwargs={"fields": list(fields), "limit": int(limit), "offset": int(offset), "order": order},
            stage=f"UOM_NI_SEARCH_READ_{model}",
            mutating=False,
        )
        batch = result if isinstance(result, list) else []
        rows.extend(batch)
        if len(batch) < limit:
            break
        offset += limit
    return rows


async def _read_map(
    conn: Any,
    model: str,
    ids: Iterable[int],
    fields: Sequence[str],
    *,
    chunk_size: int = 500,
) -> dict[int, dict[str, Any]]:
    clean_ids = _unique_ints(ids)
    out: dict[int, dict[str, Any]] = {}
    for index in range(0, len(clean_ids), chunk_size):
        chunk = clean_ids[index : index + chunk_size]
        rows = await conn.client.read(model, chunk, fields=list(fields), stage=f"UOM_NI_READ_{model}")
        for row in rows:
            row_id = _id(row.get("id"))
            if row_id:
                out[row_id] = row
    return out


async def _fetch_uoms(conn: Any, ids: Iterable[int]) -> dict[int, dict[str, Any]]:
    rows = await _read_map(conn, MODEL_UOM, ids, UOM_FIELDS)
    root_ids = {_uom_root_id(row) for row in rows.values()}
    missing_roots = [root_id for root_id in root_ids if root_id and root_id not in rows]
    if missing_roots:
        rows.update(await _read_map(conn, MODEL_UOM, missing_roots, UOM_FIELDS))
    return rows


async def _select_companies(
    conn: Any,
    keywords: Sequence[str],
    *,
    company_limit: int,
    seed: int,
) -> list[dict[str, Any]]:
    clean_keywords = [keyword.strip() for keyword in keywords if keyword.strip()]
    if not clean_keywords:
        raise RuntimeError("At least one company keyword is required.")
    by_id: dict[int, dict[str, Any]] = {}
    for keyword in clean_keywords:
        rows = await conn.client.search_read(
            "res.company",
            [["name", "ilike", keyword]],
            fields=["id", "name"],
            limit=500,
            order="name",
            stage="UOM_NI_COMPANY_SEARCH",
        )
        for row in rows:
            row_id = _id(row.get("id"))
            if row_id:
                by_id[row_id] = row
    companies = list(by_id.values())
    rng = random.Random(seed)
    rng.shuffle(companies)
    return companies[: max(1, int(company_limit or 1))]


def _stock_move_names(rows: Sequence[dict[str, Any]]) -> str:
    names: list[str] = []
    seen: set[str] = set()
    for row in rows:
        name = str(row.get("reference") or _name(row.get("picking_id")) or row.get("name") or "")
        if name and name not in seen:
            seen.add(name)
            names.append(name)
    return ", ".join(names)


def _first_non_empty(values: Iterable[Any]) -> str:
    for value in values:
        text = str(value or "").strip()
        if text:
            return text
    return ""


def _join_unique(values: Iterable[Any]) -> str:
    out: list[str] = []
    seen: set[str] = set()
    for value in values:
        text = str(value or "").strip()
        if not text or text in seen:
            continue
        seen.add(text)
        out.append(text)
    return ", ".join(out)


def _signed_line_amount(row: dict[str, Any], field_name: str) -> float:
    sign = -1.0 if str(row.get("move_type") or "") == "in_refund" else 1.0
    return sign * _float(row.get(field_name))


def _is_vendor_bill_line(row: dict[str, Any]) -> bool:
    return str(row.get("move_type") or "") in VENDOR_BILL_MOVE_TYPES


def _summarize_bill_lines(
    line_bills: Sequence[dict[str, Any]],
    *,
    product_uom: dict[str, Any] | None,
    uoms: dict[int, dict[str, Any]],
) -> dict[str, Any]:
    vendor_lines = [row for row in line_bills if _is_vendor_bill_line(row)]
    other_lines = [row for row in line_bills if not _is_vendor_bill_line(row)]
    vendor_uoms = [uoms.get(_id(row.get("product_uom_id"))) for row in vendor_lines]
    vendor_roots = [uoms.get(_uom_root_id(uom)) for uom in vendor_uoms if uom]
    inline_checks = [_is_inline(uom, product_uom) for uom in vendor_uoms if uom]
    return {
        "Bill": _join_unique(_first_non_empty([row.get("move_name"), _name(row.get("move_id"))]) for row in vendor_lines),
        "Bill Line IDs": _join_unique(str(_id(row.get("id"))) for row in vendor_lines if _id(row.get("id"))),
        "Bill Date": _join_unique(row.get("date") for row in vendor_lines),
        "Bill Qty": sum(_signed_line_amount(row, "quantity") for row in vendor_lines) if vendor_lines else "",
        "Bill UoM": _join_unique(_uom_name(uom) for uom in vendor_uoms),
        "Bill UoM Root": _join_unique(_uom_name(root) for root in vendor_roots),
        "Bill Inline?": "Yes" if inline_checks and all(inline_checks) else "No" if inline_checks else "",
        "Bill Subtotal": sum(_signed_line_amount(row, "price_subtotal") for row in vendor_lines) if vendor_lines else "",
        "Other AML IDs": _join_unique(str(_id(row.get("id"))) for row in other_lines if _id(row.get("id"))),
        "Other Linked Moves": _join_unique(_first_non_empty([row.get("move_name"), _name(row.get("move_id"))]) for row in other_lines),
        "Other Linked Net": sum(_float(row.get("balance")) for row in other_lines) if other_lines else "",
    }


def _append_rows(sheet: Any, rows: Iterable[Sequence[Any]]) -> None:
    for row in rows:
        sheet.append(list(row))


def _coerce_numeric_cell_value(value: Any) -> int | float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return value
    text = str(value).strip()
    if not text:
        return None
    try:
        number = Decimal(text)
    except (InvalidOperation, ValueError):
        return None
    if number == number.to_integral():
        return int(number)
    return float(number)


def _apply_column_formats(
    sheet: Any,
    *,
    text_headers: set[str] | None = None,
    quantity_headers: set[str] | None = None,
    currency_headers: set[str] | None = None,
) -> None:
    text_headers = text_headers or set()
    quantity_headers = quantity_headers or set()
    currency_headers = currency_headers or set()
    headers = [str(cell.value or "") for cell in sheet[1]]
    for col_idx, header in enumerate(headers, start=1):
        if header in text_headers:
            for cell in sheet.iter_cols(min_col=col_idx, max_col=col_idx, min_row=2, max_row=sheet.max_row):
                for item in cell:
                    item.number_format = "@"
                    if item.value is not None:
                        item.value = str(item.value)
        elif header in quantity_headers:
            for cell in sheet.iter_cols(min_col=col_idx, max_col=col_idx, min_row=2, max_row=sheet.max_row):
                for item in cell:
                    numeric = _coerce_numeric_cell_value(item.value)
                    if numeric is not None:
                        item.value = numeric
                    item.number_format = '#,##0.########'
        elif header in currency_headers:
            for cell in sheet.iter_cols(min_col=col_idx, max_col=col_idx, min_row=2, max_row=sheet.max_row):
                for item in cell:
                    numeric = _coerce_numeric_cell_value(item.value)
                    if numeric is not None:
                        item.value = numeric
                    item.number_format = '#,##0.00'


def _style_sheet(
    sheet: Any,
    *,
    freeze: str = "A2",
    widths: dict[str, int] | None = None,
    text_headers: set[str] | None = None,
    quantity_headers: set[str] | None = None,
    currency_headers: set[str] | None = None,
    header_fills: dict[str, str] | None = None,
) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(bottom=thin)
    sheet.freeze_panes = freeze
    if sheet.max_row >= 1:
        for cell in sheet[1]:
            header_text = str(cell.value or "")
            cell.fill = PatternFill("solid", fgColor=(header_fills or {}).get(header_text, "1F4E78")) if header_fills else header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if sheet.max_row and sheet.max_column:
        sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    if sheet.max_row >= 2:
        _apply_column_formats(
            sheet,
            text_headers=text_headers,
            quantity_headers=quantity_headers,
            currency_headers=currency_headers,
        )
    for col_idx, column_cells in enumerate(sheet.columns, start=1):
        letter = get_column_letter(col_idx)
        if widths and letter in widths:
            sheet.column_dimensions[letter].width = widths[letter]
            continue
        max_len = 0
        for cell in column_cells:
            text = str(cell.value or "")
            max_len = max(max_len, min(len(text), 60))
        sheet.column_dimensions[letter].width = max(10, min(max_len + 2, 55))


def _write_excel(
    path: Path,
    *,
    summary: list[tuple[str, Any]],
    companies: list[dict[str, Any]],
    non_inline_rows: list[dict[str, Any]],
    a_belg_rows: list[dict[str, Any]],
    repair_guidance: list[tuple[str, str]],
) -> None:
    workbook = Workbook()
    try:
        summary_sheet = workbook.active
        summary_sheet.title = "Summary"
        _append_rows(summary_sheet, [("Metric", "Value"), *summary])
        _style_sheet(summary_sheet, freeze="A2", widths={"A": 36, "B": 90})

        company_sheet = workbook.create_sheet("Companies")
        _append_rows(company_sheet, [("Company ID", "Company Name")])
        _append_rows(company_sheet, [(row.get("id"), row.get("name")) for row in companies])
        _style_sheet(company_sheet, freeze="A2", text_headers={"Company ID"}, header_fills={"Company ID": "C65911"})

        detail_sheet = workbook.create_sheet("Non Inline Transactions")
        id_headers = {
            "PO Line ID",
            "Bill Line IDs",
            "Other AML IDs",
            "Stock Move IDs",
            "SVL IDs",
            "SVL JE IDs",
            "SVL ID",
            "SVL JE ID",
        }
        transaction_headers = {"PO", "Bill", "Other Linked Moves", "Stock Picking", "SVL JE", "SVL Ref", "Stock Move"}
        uom_headers = {
            "PO UoM",
            "PO UoM Root",
            "Product/SVL UoM",
            "Product/SVL UoM Root",
            "Bill UoM",
            "Bill UoM Root",
            "Move UoM",
        }
        quantity_headers = {"PO Qty", "PO Total Qty", "Correct Inline Qty", "Qty Gap", "Bill Qty", "SVL Qty", "Remaining Qty", "Move Qty"}
        currency_headers = {"Bill Subtotal", "Other Linked Net", "SVL Value", "SVL Unit Cost", "Remaining Value", "Unit Cost"}
        header_fills = {
            **{header: "C65911" for header in id_headers},
            **{header: "548235" for header in transaction_headers},
            **{header: "7F7F7F" for header in uom_headers},
            **{header: "8064A2" for header in currency_headers},
            **{header: "9E480E" for header in quantity_headers},
        }
        headers = [
            "Company",
            "PO Line ID",
            "Bill Line IDs",
            "Other AML IDs",
            "Stock Move IDs",
            "SVL IDs",
            "SVL JE IDs",
            "PO",
            "Bill",
            "Other Linked Moves",
            "Stock Picking",
            "SVL JE",
            "Product Code",
            "Product",
            "PO State",
            "Expected Arrival",
            "Bill Date",
            "PO UoM",
            "PO UoM Root",
            "Product/SVL UoM",
            "Product/SVL UoM Root",
            "Bill UoM",
            "Bill UoM Root",
            "Bill Inline?",
            "Bill Subtotal",
            "Other Linked Net",
            "SVL Value",
            "SVL Unit Cost",
            "PO Qty",
            "PO Total Qty",
            "Correct Inline Qty",
            "Qty Gap",
            "Bill Qty",
            "SVL Qty",
            "Risk",
            "Suggested Action",
        ]
        _append_rows(detail_sheet, [headers])
        for row in non_inline_rows:
            _append_rows(detail_sheet, [[row.get(header, "") for header in headers]])
        _style_sheet(
            detail_sheet,
            freeze="A2",
            text_headers={
                "PO Line ID",
                "Bill Line IDs",
                "Other AML IDs",
                "Stock Move IDs",
                "SVL IDs",
                "SVL JE IDs",
                "PO",
                "Bill",
                "Other Linked Moves",
                "Stock Picking",
                "SVL JE",
                "Product Code",
            },
            quantity_headers=quantity_headers,
            currency_headers=currency_headers,
            header_fills=header_fills,
        )

        product_sheet = workbook.create_sheet("A-BELG-0024 Bekasi SVL")
        product_headers = [
            "Company",
            "SVL ID",
            "PO Line ID",
            "Bill Line IDs",
            "Other AML IDs",
            "SVL JE ID",
            "SVL Ref",
            "SVL JE",
            "Stock Move",
            "PO",
            "Bill",
            "Other Linked Moves",
            "Product Code",
            "Product",
            "SVL Description",
            "SVL Create Date",
            "SVL/Product UoM",
            "Move UoM",
            "PO UoM",
            "PO UoM Root",
            "Inline PO->Product?",
            "Bill UoM",
            "Inline Bill->Product?",
            "SVL Value",
            "Remaining Value",
            "Other Linked Net",
            "Unit Cost",
            "SVL Qty",
            "Remaining Qty",
            "Move Qty",
            "PO Qty",
            "Bill Qty",
            "Observation",
        ]
        _append_rows(product_sheet, [product_headers])
        for row in a_belg_rows:
            _append_rows(product_sheet, [[row.get(header, "") for header in product_headers]])
        _style_sheet(
            product_sheet,
            freeze="A2",
            text_headers={
                "SVL ID",
                "PO Line ID",
                "Bill Line IDs",
                "Other AML IDs",
                "SVL JE ID",
                "SVL Ref",
                "SVL JE",
                "Stock Move",
                "PO",
                "Bill",
                "Other Linked Moves",
                "Product Code",
            },
            quantity_headers=quantity_headers,
            currency_headers=currency_headers,
            header_fills=header_fills,
        )

        guidance_sheet = workbook.create_sheet("Repair Guidance")
        _append_rows(guidance_sheet, [("Topic", "Guidance"), *repair_guidance])
        _style_sheet(guidance_sheet, freeze="A2", widths={"A": 28, "B": 110})

        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
    finally:
        workbook.close()


def _validate_xlsx(path: Path) -> None:
    with zipfile.ZipFile(path, "r") as zf:
        table_parts = [name for name in zf.namelist() if name.startswith("xl/tables/")]
    if table_parts:
        raise RuntimeError(f"Unexpected Excel table parts detected: {table_parts}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Non Inline Transactions" not in workbook.sheetnames:
            raise RuntimeError("Workbook validation failed: detail sheet missing.")
    finally:
        workbook.close()


async def _load_scope(args: argparse.Namespace, conn: Any) -> dict[str, Any]:
    page_size = max(50, int(args.page_size or 500))
    company_keywords = parse_csv_values(args.company_keywords)
    companies = await _select_companies(
        conn,
        company_keywords,
        company_limit=args.company_limit,
        seed=args.random_seed,
    )
    if not companies:
        raise RuntimeError(f"No companies found for keywords: {company_keywords}")
    company_ids = [int(row["id"]) for row in companies]

    model_fields = {
        "purchase.order.line": await _fields(conn, "purchase.order.line"),
        "product.product": await _fields(conn, "product.product"),
        "account.move.line": await _fields(conn, "account.move.line"),
        "stock.move": await _fields(conn, "stock.move"),
        "stock.valuation.layer": await _fields(conn, "stock.valuation.layer"),
    }
    po_fields = _filter_fields(PO_LINE_FIELDS, model_fields["purchase.order.line"])
    product_fields = _filter_fields(PRODUCT_FIELDS, model_fields["product.product"])
    bill_line_fields = _filter_fields(BILL_LINE_FIELDS, model_fields["account.move.line"])
    stock_move_fields = _filter_fields(STOCK_MOVE_FIELDS, model_fields["stock.move"])
    svl_fields = _filter_fields(SVL_FIELDS, model_fields["stock.valuation.layer"])

    po_domain: list[Any] = [
        ["company_id", "in", company_ids],
        ["product_id", "!=", False],
        ["product_uom_id", "!=", False],
    ]
    if "state" in model_fields["purchase.order.line"]:
        po_domain.append(["state", "in", ["purchase", "done"]])
    if args.date_from:
        po_domain.append(["date_planned", ">=", _date_to_datetime_start(args.date_from)])
    if args.date_to:
        po_domain.append(["date_planned", "<=", _date_to_datetime_end(args.date_to)])

    po_count = await _search_count(conn, "purchase.order.line", po_domain)
    po_lines = await _search_read_paged(
        conn,
        "purchase.order.line",
        po_domain,
        po_fields,
        page_size=page_size,
        order="id",
        max_rows=max(0, int(args.max_purchase_lines or 0)),
    )
    products = await _read_map(
        conn,
        "product.product",
        _unique_ints(row.get("product_id") for row in po_lines),
        product_fields,
    )

    base_uom_ids = []
    for line in po_lines:
        base_uom_ids.append(_id(line.get("product_uom_id")))
        base_uom_ids.append(_id(products.get(_id(line.get("product_id")), {}).get("uom_id")))
    uoms = await _fetch_uoms(conn, base_uom_ids)

    candidate_lines: list[dict[str, Any]] = []
    for line in po_lines:
        product = products.get(_id(line.get("product_id")), {})
        po_uom = uoms.get(_id(line.get("product_uom_id")))
        product_uom = uoms.get(_id(product.get("uom_id")))
        if po_uom and product_uom and not _is_inline(po_uom, product_uom):
            candidate_lines.append(line)

    invoice_line_ids: list[int] = []
    move_ids: list[int] = []
    for line in candidate_lines:
        invoice_line_ids.extend(_ids_from_x2many(line.get("invoice_lines")))
        move_ids.extend(_ids_from_x2many(line.get("move_ids")))

    bill_lines = await _read_map(conn, "account.move.line", invoice_line_ids, bill_line_fields) if invoice_line_ids else {}
    move_rows = await _read_map(conn, "stock.move", move_ids, stock_move_fields) if move_ids else {}
    svl_rows = []
    if move_rows:
        svl_rows = await _search_read_paged(
            conn,
            "stock.valuation.layer",
            [["stock_move_id", "in", list(move_rows)], ["company_id", "in", company_ids]],
            svl_fields,
            page_size=page_size,
            order="id",
        )

    extra_uom_ids = []
    for row in bill_lines.values():
        extra_uom_ids.append(_id(row.get("product_uom_id")))
    for row in move_rows.values():
        extra_uom_ids.append(_id(row.get("product_uom")))
    for row in svl_rows:
        extra_uom_ids.append(_id(row.get("uom_id")))
    if extra_uom_ids:
        uoms.update(await _fetch_uoms(conn, extra_uom_ids))

    return {
        "page_size": page_size,
        "company_keywords": company_keywords,
        "companies": companies,
        "company_ids": company_ids,
        "po_count": po_count,
        "po_lines": po_lines,
        "candidate_lines": candidate_lines,
        "products": products,
        "uoms": uoms,
        "bill_lines": bill_lines,
        "move_rows": move_rows,
        "svl_rows": svl_rows,
        "fields": {
            "po": po_fields,
            "product": product_fields,
            "bill_line": bill_line_fields,
            "stock_move": stock_move_fields,
            "svl": svl_fields,
        },
    }


def _build_non_inline_rows(scope: dict[str, Any]) -> list[dict[str, Any]]:
    uoms: dict[int, dict[str, Any]] = scope["uoms"]
    products: dict[int, dict[str, Any]] = scope["products"]
    bill_lines: dict[int, dict[str, Any]] = scope["bill_lines"]
    move_rows: dict[int, dict[str, Any]] = scope["move_rows"]
    svl_rows: list[dict[str, Any]] = scope["svl_rows"]

    bill_by_pol: dict[int, list[dict[str, Any]]] = defaultdict(list)
    for row in bill_lines.values():
        bill_by_pol[_id(row.get("purchase_line_id"))].append(row)
    move_by_pol: dict[int, list[dict[str, Any]]] = defaultdict(list)
    for row in move_rows.values():
        move_by_pol[_id(row.get("purchase_line_id"))].append(row)
    svl_by_move: dict[int, list[dict[str, Any]]] = defaultdict(list)
    for row in svl_rows:
        svl_by_move[_id(row.get("stock_move_id"))].append(row)

    rows: list[dict[str, Any]] = []
    for line in scope["candidate_lines"]:
        product = products.get(_id(line.get("product_id")), {})
        po_uom = uoms.get(_id(line.get("product_uom_id")))
        product_uom = uoms.get(_id(product.get("uom_id")))
        po_root = uoms.get(_uom_root_id(po_uom))
        product_root = uoms.get(_uom_root_id(product_uom))
        moves = move_by_pol.get(_id(line.get("id")), [])
        svls_for_line = [svl for move in moves for svl in svl_by_move.get(_id(move.get("id")), [])]
        bill_summary = _summarize_bill_lines(bill_by_pol.get(_id(line.get("id")), []), product_uom=product_uom, uoms=uoms)
        correct_qty = _convert(po_uom, product_uom, line.get("product_qty"))
        svl_qty = sum(_float(row.get("quantity")) for row in svls_for_line)
        svl_value = sum(_float(row.get("value")) for row in svls_for_line)
        rows.append(
            {
                "Company": _name(line.get("company_id")),
                "PO": _name(line.get("order_id")),
                "PO Line ID": _id(line.get("id")),
                "PO State": line.get("state", ""),
                "Expected Arrival": line.get("date_planned", ""),
                "Product Code": product.get("default_code", ""),
                "Product": product.get("display_name") or _name(line.get("product_id")),
                "PO Qty": line.get("product_qty"),
                "PO UoM": _uom_name(po_uom),
                "PO UoM Root": _uom_name(po_root),
                "Product/SVL UoM": _uom_name(product_uom),
                "Product/SVL UoM Root": _uom_name(product_root),
                "PO Total Qty": line.get("product_uom_qty"),
                "Correct Inline Qty": correct_qty or "Blocked - non-inline",
                "Qty Gap": _float(line.get("product_uom_qty")) - _float(correct_qty) if correct_qty else "",
                "Stock Picking": _stock_move_names(moves),
                "Stock Move IDs": ", ".join(str(_id(move.get("id"))) for move in moves),
                "SVL IDs": ", ".join(str(_id(svl.get("id"))) for svl in svls_for_line),
                "SVL JE IDs": _join_unique(
                    str(_id(svl.get("account_move_id"))) for svl in svls_for_line if _id(svl.get("account_move_id"))
                ),
                "SVL JE": _join_unique(_name(svl.get("account_move_id")) for svl in svls_for_line),
                "SVL Qty": svl_qty,
                "SVL Value": svl_value,
                "SVL Unit Cost": svl_value / svl_qty if svl_qty else "",
                "Risk": "Fatal: upstream UoM root differs from product/SVL UoM root.",
                "Suggested Action": (
                    "Do not repair by JE only; validate business quantity and fix stock/SVL quantity flow first. "
                    "Post value JE only if corrected total value has a delta."
                ),
                **bill_summary,
            }
        )
    return rows


async def _build_focus_rows(args: argparse.Namespace, conn: Any, scope: dict[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    fields = scope["fields"]
    uoms: dict[int, dict[str, Any]] = scope["uoms"]
    page_size = int(scope["page_size"])

    focus_products = await conn.client.search_read(
        "product.product",
        [["default_code", "=", args.product_code]],
        fields=fields["product"],
        limit=50,
        order="id",
        stage="UOM_NI_FOCUS_PRODUCT",
    )
    focus_product_ids = _unique_ints(row.get("id") for row in focus_products)
    focus_product_by_id = {int(row["id"]): row for row in focus_products if _id(row.get("id"))}

    focus_company_rows = await conn.client.search_read(
        "res.company",
        [["name", "ilike", args.bekasi_company_keyword]],
        fields=["id", "name"],
        limit=20,
        order="name",
        stage="UOM_NI_FOCUS_COMPANY",
    )
    if not focus_company_rows and args.bekasi_company_keyword.upper() != "BEKASI":
        focus_company_rows = await conn.client.search_read(
            "res.company",
            [["name", "ilike", "BEKASI"]],
            fields=["id", "name"],
            limit=20,
            order="name",
            stage="UOM_NI_FOCUS_COMPANY_FALLBACK",
        )
    focus_company_ids = _unique_ints(row.get("id") for row in focus_company_rows)

    focus_svl_rows: list[dict[str, Any]] = []
    if focus_product_ids and focus_company_ids:
        focus_svl_rows = await _search_read_paged(
            conn,
            "stock.valuation.layer",
            [["product_id", "in", focus_product_ids], ["company_id", "in", focus_company_ids]],
            fields["svl"],
            page_size=page_size,
            order="id",
        )

    focus_moves = await _read_map(
        conn,
        "stock.move",
        _unique_ints(row.get("stock_move_id") for row in focus_svl_rows),
        fields["stock_move"],
    )
    focus_po_lines = await _read_map(
        conn,
        "purchase.order.line",
        _unique_ints(row.get("purchase_line_id") for row in focus_moves.values()),
        fields["po"],
    )
    focus_invoice_ids: list[int] = []
    for line in focus_po_lines.values():
        focus_invoice_ids.extend(_ids_from_x2many(line.get("invoice_lines")))
    focus_bills = await _read_map(conn, "account.move.line", focus_invoice_ids, fields["bill_line"])

    focus_bill_by_pol: dict[int, list[dict[str, Any]]] = defaultdict(list)
    for bill in focus_bills.values():
        focus_bill_by_pol[_id(bill.get("purchase_line_id"))].append(bill)

    focus_uom_ids: list[int] = []
    for row in focus_products:
        focus_uom_ids.append(_id(row.get("uom_id")))
    for row in focus_svl_rows:
        focus_uom_ids.append(_id(row.get("uom_id")))
    for row in focus_moves.values():
        focus_uom_ids.append(_id(row.get("product_uom")))
    for row in focus_po_lines.values():
        focus_uom_ids.append(_id(row.get("product_uom_id")))
    for row in focus_bills.values():
        focus_uom_ids.append(_id(row.get("product_uom_id")))
    if focus_uom_ids:
        uoms.update(await _fetch_uoms(conn, focus_uom_ids))

    rows: list[dict[str, Any]] = []
    for svl in focus_svl_rows:
        product = focus_product_by_id.get(_id(svl.get("product_id")), {})
        product_uom = uoms.get(_id(product.get("uom_id"))) or uoms.get(_id(svl.get("uom_id")))
        move = focus_moves.get(_id(svl.get("stock_move_id")), {})
        po_line = focus_po_lines.get(_id(move.get("purchase_line_id")), {})
        po_uom = uoms.get(_id(po_line.get("product_uom_id")))
        po_root = uoms.get(_uom_root_id(po_uom)) if po_uom else None
        bill_summary = _summarize_bill_lines(
            focus_bill_by_pol.get(_id(po_line.get("id")), []),
            product_uom=product_uom,
            uoms=uoms,
        )
        inline_po = _is_inline(po_uom, product_uom) if po_uom else False
        rows.append(
            {
                "Company": _name(svl.get("company_id")),
                "SVL ID": _id(svl.get("id")),
                "PO Line ID": _id(po_line.get("id")),
                "Bill Line IDs": bill_summary["Bill Line IDs"],
                "Other AML IDs": bill_summary["Other AML IDs"],
                "SVL JE ID": _id(svl.get("account_move_id")) or "",
                "SVL Ref": svl.get("reference", ""),
                "SVL JE": _name(svl.get("account_move_id")),
                "Stock Move": _first_non_empty([move.get("reference"), _name(move.get("picking_id")), move.get("name")]),
                "PO": _name(po_line.get("order_id")),
                "Bill": bill_summary["Bill"],
                "Other Linked Moves": bill_summary["Other Linked Moves"],
                "SVL Description": svl.get("description", ""),
                "SVL Create Date": svl.get("create_date", ""),
                "Product Code": product.get("default_code", ""),
                "Product": product.get("display_name") or _name(svl.get("product_id")),
                "SVL Qty": svl.get("quantity", ""),
                "Remaining Qty": svl.get("remaining_qty", ""),
                "SVL Value": svl.get("value", ""),
                "Remaining Value": svl.get("remaining_value", ""),
                "Other Linked Net": bill_summary["Other Linked Net"],
                "Unit Cost": svl.get("unit_cost", ""),
                "SVL/Product UoM": _uom_name(product_uom),
                "Move Qty": _first_non_empty([move.get("quantity"), move.get("product_qty"), move.get("product_uom_qty")]),
                "Move UoM": _uom_name(uoms.get(_id(move.get("product_uom")))),
                "PO Qty": po_line.get("product_qty", ""),
                "PO UoM": _uom_name(po_uom),
                "PO UoM Root": _uom_name(po_root),
                "Inline PO->Product?": "Yes" if po_uom and inline_po else "No" if po_uom else "",
                "Bill Qty": bill_summary["Bill Qty"],
                "Bill UoM": bill_summary["Bill UoM"],
                "Inline Bill->Product?": bill_summary["Bill Inline?"],
                "Observation": "Upstream UoM not inline" if po_uom and not inline_po else "Upstream UoM inline or no PO link",
            }
        )
    return rows, focus_company_rows


async def _build_report(args: argparse.Namespace) -> tuple[dict[str, Any], Path]:
    conn = await open_tool_connection(
        database_profile_id=args.database_profile,
        max_concurrency=args.max_concurrency,
        verbose=args.verbose,
    )
    try:
        scope = await _load_scope(args, conn)
        non_inline_rows = _build_non_inline_rows(scope)
        focus_rows, focus_companies = await _build_focus_rows(args, conn, scope)

        today = datetime.now().strftime("%Y-%m-%d")
        output_path = Path(args.output).expanduser() if args.output else DEFAULT_ARTIFACT_DIR / f"UoM Not Inline Transactions - {today}.xlsx"
        if not output_path.is_absolute():
            output_path = Path.cwd() / output_path

        repair_guidance = [
            (
                "Cycle 7 proposal",
                "Treat UoM Not Inline as a blocking data-quality cycle, not an automatic JE repair case.",
            ),
            (
                "Journal repair",
                "A journal-only repair is not sufficient when total value is already balanced; the primary defect is stock/SVL quantity and cost basis.",
            ),
            (
                "DBKS/IN/00016 pattern",
                "If 3 CASE @24 CAN @320 ML becomes 23040 CAN @1 (EA), the accounting value can still balance while quantity/unit cost is wrong.",
            ),
            (
                "Corrective direction",
                "Validate true business quantity, correct stock receipt/SVL quantity through a controlled stock/UoM data fix, then post value JE only for any remaining value delta.",
            ),
        ]
        summary = [
            ("Database", conn.config.database),
            ("Generated At", datetime.now().isoformat(timespec="seconds")),
            ("Company Keywords", ", ".join(scope["company_keywords"])),
            ("Selected Company Count", len(scope["companies"])),
            ("Purchase Lines Matching Scope", scope["po_count"]),
            ("Purchase Lines Scanned", len(scope["po_lines"])),
            ("Non-Inline PO Lines", len(scope["candidate_lines"])),
            ("Excel Detail Rows", len(non_inline_rows)),
            ("Focused Product", args.product_code),
            ("Focused Bekasi Companies", ", ".join(row.get("name", "") for row in focus_companies)),
            ("Focused Product SVL Rows", len(focus_rows)),
            ("Scope Note", "Read-only. Starts from purchase.order.line / bill account.move.line upstream UoM."),
        ]
        _write_excel(
            output_path,
            summary=summary,
            companies=scope["companies"],
            non_inline_rows=non_inline_rows,
            a_belg_rows=focus_rows,
            repair_guidance=repair_guidance,
        )
        _validate_xlsx(output_path)

        payload = build_base_payload(conn, action="uom_not_inline_report")
        payload.update(
            {
                "output_path": str(output_path),
                "company_keywords": scope["company_keywords"],
                "selected_companies": scope["companies"],
                "purchase_line_count": scope["po_count"],
                "purchase_lines_scanned": len(scope["po_lines"]),
                "non_inline_po_line_count": len(scope["candidate_lines"]),
                "excel_detail_rows": len(non_inline_rows),
                "focused_product": args.product_code,
                "focused_bekasi_companies": focus_companies,
                "focused_product_svl_rows": len(focus_rows),
                "sample_non_inline_rows": non_inline_rows[:10],
                "guardrail": "Starts from upstream PO/Bill UoM and compares parent_path root to product/SVL UoM.",
            }
        )
        return payload, output_path
    finally:
        await conn.close()


async def amain() -> int:
    args = parse_args()
    payload, output_path = await _build_report(args)
    print(json.dumps(payload, ensure_ascii=False, indent=2, default=str))
    if not args.no_artifact:
        json_path = write_json_artifact(payload, output_dir=args.output_dir, prefix="uom_not_inline_report")
        print(f"JSON Artefact: {json_path}")
    print(f"Excel: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(asyncio.run(amain()))
