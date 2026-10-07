import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from smartscc_tools.features.svl_fix_je.dashboard_pcb_export import (
    export_pcb_company_audit_excel,
    export_pcb_cycle_detail_excel,
)
from smartscc_tools.features.svl_fix_je.dashboard_repair_export import (
    build_pcb_case1_export_payload,
    build_repair_summary_export_payload,
    export_pcb_case1_summary_excel,
    export_repair_summary_excel,
    export_repair_summary_html,
)
from smartscc_tools.features.svl_fix_je.models import (
    SvlDashboardCycleAccountRow,
    SvlDashboardCycleItemRow,
    SvlDashboardPcbAdjustmentAuditRow,
    SvlDashboardPurchaseCycle,
    SvlDashboardRepairRowResult,
    SvlDashboardSnapshot,
)
from smartscc_tools.modules.svl_fix_je_dashboard_page import SvlFixJeDashboardPage


def _sample_results() -> list[SvlDashboardRepairRowResult]:
    return [
        SvlDashboardRepairRowResult(
            row_key="row-1",
            status="POSTED",
            company_id=7,
            company_name="Alpha Company",
            item_code="A-LQSC-0001",
            item_name="BAILEYS",
            amount=125000.5,
            reference="Correction New JE + Relink SVL A-LQSC-0001",
            message="Move baru dibuat dan SVL direlink.",
            selected_target_mode="new_and_relink",
            effective_date="2026-03-19",
            move_id=9001,
            move_name="STJ/2026/09001",
            posted=True,
            relinked_svl_id=501,
            old_move_action="mark_only",
            old_move_id=435,
            old_move_name="STJ/2025/10/0435",
            svl_reference="WCGT/IN/00034",
            repair_source_kind="svl_linked_empty_move",
            repair_source_label="Linked JE Header Kosong",
            debit_account_code="1105001",
            debit_account_name="Persediaan Alkohol",
            credit_account_code="1108099",
            credit_account_name="Akun koreksi default",
        ),
        SvlDashboardRepairRowResult(
            row_key="row-2",
            status="CREATED",
            company_id=7,
            company_name="Alpha Company",
            item_code="B-DGSD-0003",
            item_name="AIR MINERAL",
            amount=55000.0,
            reference="Correction New JE + Relink SVL B-DGSD-0003",
            message="Move baru dibuat dalam draft only.",
            selected_target_mode="new_and_relink",
            effective_date="2026-03-19",
            move_id=9002,
            move_name="STJ/2026/09002",
            posted=False,
            relinked_svl_id=502,
            svl_reference="WCGT/IN/00035",
            repair_source_kind="svl_no_move",
            repair_source_label="SVL tanpa JE",
            debit_account_code="1105002",
            debit_account_name="Persediaan Minuman",
            credit_account_code="1108099",
            credit_account_name="Akun koreksi default",
        ),
    ]


def _sample_pcb_rows() -> list[dict[str, object]]:
    return [
        {
            "row_key": "case1::7001::901::bill_line::8101",
            "cycle_key": "case1::7001",
            "company_id": 7,
            "company_name": "Alpha Company",
            "picking_name": "LHPK/IN/7001",
            "source_label": "Bill Line #8101",
            "po_name": "PO/2026/0001",
            "bill_name": "BILL/2026/0001",
            "stj_refs": ["STJ/2026/0451"],
            "date": "2026-03-19",
            "reference": "Koreksi PCB: LHPK/IN/7001 / BILL/2026/0001 / SKU-001",
            "journal_code": "STJ",
            "amount": 120.0,
            "debit_account_code": "1108099",
            "credit_account_code": "2103006",
            "reconcile_ready": True,
            "reconcile_readiness_label": "Exact Auto",
            "reconcile_attempted": True,
            "reconcile_performed": True,
            "reconcile_skipped": False,
            "reconcile_message": "2103006: exact reconcile berhasil. | 1108099: exact reconcile berhasil.",
            "row_status": "repaired",
            "row_status_message": "JE STJ/2026/0901 dibuat dan dipost.",
            "result_status": "POSTED",
            "result_posted": True,
            "result_move_id": 901,
            "result_move_name": "STJ/2026/0901",
            "existing_move_detected": False,
            "product_id": 901,
            "item_code": "SKU-001",
            "item_name": "Produk A",
            "item_category_name": "Raw",
            "bill_line_id": 8101,
            "bill_move_id": 8201,
            "purchase_line_id": 8301,
            "stock_move_id": 8401,
            "stock_move_ids": [8401],
            "stj_move_ids": [8801],
            "picking_id": 7001,
            "payment_move_ids": [8501],
            "bank_move_ids": [8601],
            "suspend_target_aml_ids": [9102],
            "clearing_target_aml_ids": [9101],
            "partner_id": 77,
            "partner_name": "Vendor Alpha",
            "currency_id": 13,
            "product_uom_id": 11,
            "quantity": 2.0,
            "amount_currency": 120.0,
            "analytic_distribution": {"CC-01": 100.0},
            "bill_price_unit": 60.0,
            "gr_price_unit": 0.0,
            "price_gap_value": 120.0,
            "allocated_amount": 120.0,
            "result_error_kind": "",
            "reconcile_error_kind": "",
        }
    ]


def _sample_pcb_adjustment_audit_row() -> SvlDashboardPcbAdjustmentAuditRow:
    return SvlDashboardPcbAdjustmentAuditRow(
        move_id=9901,
        move_name="RAC/2026/01/0026",
        move_date="2026-03-27",
        move_ref="Adjustment clearing from standard price",
        product_id=901,
        item_code="SKU-001",
        item_name="Produk A",
        source_account_code="5101010",
        source_account_name="Selisih HPP / COGS Variance",
        clearing_account_code="1108099",
        repair_clearing_amount=100.0,
        origin_move_id=8801,
        origin_move_name="STJ/2026/03/0262",
        origin_basis="explicit.move_ref",
        origin_product_id=901,
        origin_purchase_line_id=8301,
        origin_stock_move_id=8401,
        verified_for_case34=True,
        matched_basis="exact.purchase_line_id",
        ambiguous=False,
    )


def _sample_pcb_cycle_for_export() -> SvlDashboardPurchaseCycle:
    audit_row = _sample_pcb_adjustment_audit_row()
    item_row = SvlDashboardCycleItemRow(
        product_id=901,
        product_name="ARTISAN DARK CHOCOLATE COUVERTURE 73%",
        default_code="F-DGRP-0002",
        valuation_method="automated",
        account_rows=[
            SvlDashboardCycleAccountRow(
                account_id=302,
                code="1108099",
                name="External Clearing",
                account_type="asset_current",
                account_group="asset",
                debit=201765.62,
                credit=0.0,
                net_balance=201765.62,
                status="acceptable",
            ),
            SvlDashboardCycleAccountRow(
                account_id=303,
                code="2103006",
                name="Hutang Suspensed Pengadaan Barang",
                account_type="liability_current",
                account_group="liability",
                debit=0.0,
                credit=2868000.0,
                net_balance=-2868000.0,
                status="problem",
            ),
        ],
        bill_refs=["BILL/2026/03/0082"],
        has_item_bill=True,
        stj_refs=["STJ/2026/03/0262"],
        has_item_stj=True,
        correction_stj_refs=["STJ/2026/03/0160"],
        has_stj_evidence=True,
        external_clearing_amount=100.0,
        external_clearing_refs=["RAC/2026/01/0026 <- STJ/2026/03/0262"],
        external_clearing_basis="external.origin_purchase_line_id",
        external_clearing_verified=True,
        verified_audit_clearing_amount=100.0,
        eligible_case34=True,
        primary_case="case4",
        secondary_flags=["needs_review"],
        case_reason="Mismatch suspend vs expense",
        auto_repairable=False,
        stj_state="direct",
        svl_zero_at_gr=False,
        bill_hit_role="expense",
        repair_basis_amount=2868000.0,
        repair_basis_source="standard_cost_x_qty",
        adjustment_audit_rows=[audit_row],
    )
    return SvlDashboardPurchaseCycle(
        picking_id=2791,
        picking_name="PIKKP/IN/02791 [+2]",
        gr_date="2026-03-02",
        partner_name="SENTRAL CK - WH",
        partner_id=77,
        purchase_orders=["PO/PIKKPH/2026/02/02455"],
        stj_refs=["STJ/2026/03/0262"],
        bill_move_ids=[8201],
        bill_refs=["BILL/2026/03/0082"],
        payment_move_ids=[8501],
        payment_refs=["PBK039/2026/00609"],
        bank_move_ids=[8601],
        bank_refs=["BK/2026/03/0017"],
        product_ids=[901],
        cycle_status="problem",
        primary_case="case4",
        case_counts={"case4": 1},
        edge_flags=["external_clearing"],
        mixed_case_summary="Case 4 dominant",
        issue_patterns=["Pola 1"],
        adjustment_warning_text="Warning: Clearing adjustment -> RAC/2026/01/0026",
        account_rows=[
            SvlDashboardCycleAccountRow(
                account_id=301,
                code="2103006",
                name="Hutang Suspensed Pengadaan Barang",
                account_type="liability_current",
                account_group="liability",
                debit=0.0,
                credit=2868000.0,
                net_balance=-2868000.0,
                status="problem",
            ),
            SvlDashboardCycleAccountRow(
                account_id=302,
                code="1108099",
                name="External Clearing",
                account_type="asset_current",
                account_group="asset",
                debit=201765.62,
                credit=201765.62,
                net_balance=0.0,
                status="balanced",
            ),
        ],
        item_rows=[item_row],
        adjustment_audit_rows=[audit_row],
        total_debit=3271531.24,
        total_credit=3473296.86,
        problem_account_count=1,
        has_info_accounts=False,
        info_account_count=0,
        picking_ids=[2791, 2855, 2856],
        picking_names=["PIKKP/IN/02791", "PIKKP/IN/02855", "PIKKP/IN/02856"],
        raw_lines=[
            {
                "tanggal": "2026-03-06",
                "kode_transaksi": "BILL/2026/03/0082",
                "jenis": "BILL",
                "tipe_akun": "liability_current",
                "akun_code": "2103006",
                "akun_name": "Hutang Suspensed Pengadaan Barang",
                "kode_item": "F-DGRP-0002",
                "nama_item": "ARTISAN DARK CHOCOLATE COUVERTURE 73%",
                "uom": "BAG @2,5",
                "qty_item": 4.0,
                "kategori_produk": "READY-TO-USE",
                "no_po": "PO/PIKKPH/2026/02/02455",
                "komunikasi": "Vendor bill",
                "debit": 0.0,
                "kredit": 2868000.0,
                "saldo": -2868000.0,
                "matching": "",
                "partner": "SENTRAL CK - WH",
            },
            {
                "tanggal": "2026-03-02",
                "kode_transaksi": "STJ/2026/03/0262",
                "jenis": "STJ",
                "tipe_akun": "asset_current",
                "akun_code": "1105003",
                "akun_name": "Persediaan Makanan / Food Inventory",
                "kode_item": "F-DGRP-0002",
                "nama_item": "ARTISAN DARK CHOCOLATE COUVERTURE 73%",
                "uom": "BAG @2,5",
                "qty_item": 4.0,
                "kategori_produk": "READY-TO-USE",
                "no_po": "PO/PIKKPH/2026/02/02455",
                "komunikasi": "Stock Journal",
                "debit": 2868000.0,
                "kredit": 0.0,
                "saldo": 2868000.0,
                "matching": "M-001",
                "partner": "SENTRAL CK - WH",
            },
        ],
    )


class DashboardRepairExportTest(unittest.TestCase):
    def test_pcb_case1_payload_carries_result_and_traceability_fields(self) -> None:
        payload = build_pcb_case1_export_payload(
            database="hwgroup_erp",
            rows=_sample_pcb_rows(),
            exported_at="2026-03-20T10:00:00",
        )

        self.assertEqual(payload["company_label"], "Alpha Company - [7]")
        self.assertEqual(payload["created_count"], 1)
        self.assertEqual(payload["existing_count"], 0)
        self.assertEqual(payload["posted_count"], 1)
        self.assertEqual(payload["reconciled_count"], 1)
        self.assertEqual(payload["rows"][0]["result_move_name"], "STJ/2026/0901")
        self.assertEqual(payload["rows"][0]["stj_refs"], "STJ/2026/0451")

    def test_pcb_case1_payload_treats_posted_warning_as_non_error_and_existing_separately(self) -> None:
        rows = _sample_pcb_rows()
        rows[0]["reconcile_performed"] = False
        rows[0]["reconcile_skipped"] = True
        rows[0]["reconcile_message"] = "account.partial.reconcile.create Odoo error"
        rows[0]["reconcile_error_kind"] = "reconcile_create_failed"
        rows[0]["row_status_message"] = "JE STJ/2026/0901 dibuat dan dipost. Reconcile warning: account.partial.reconcile.create Odoo error"

        warning_payload = build_pcb_case1_export_payload(
            database="hwgroup_erp",
            rows=rows,
            exported_at="2026-03-20T10:00:00",
        )
        self.assertEqual(warning_payload["error_count"], 0)
        self.assertEqual(warning_payload["reconcile_skipped_count"], 1)

        rows[0]["existing_move_detected"] = True
        existing_payload = build_pcb_case1_export_payload(
            database="hwgroup_erp",
            rows=rows,
            exported_at="2026-03-20T10:00:00",
        )
        self.assertEqual(existing_payload["created_count"], 0)
        self.assertEqual(existing_payload["existing_count"], 1)

    def test_pcb_case1_payload_derives_debit_codes_from_case3_audit_and_residual_planned_lines(self) -> None:
        rows = _sample_pcb_rows()
        rows[0].pop("debit_account_code", None)
        rows[0].pop("credit_account_code", None)
        rows[0]["planned_lines"] = [
            {
                "role": "problem_2103006",
                "account_code": "2103006",
                "amount": 120.0,
                "side": "credit",
            },
            {
                "role": "audit_clearing_1108099",
                "account_code": "1108099",
                "amount": 70.0,
                "side": "debit",
            },
            {
                "role": "selisih_hpp",
                "account_code": "5101001",
                "amount": 50.0,
                "side": "debit",
            },
        ]

        payload = build_pcb_case1_export_payload(
            database="hwgroup_erp",
            rows=rows,
            exported_at="2026-03-20T10:00:00",
        )

        self.assertEqual(payload["rows"][0]["debit_account_code"], "1108099, 5101001")
        self.assertEqual(payload["rows"][0]["credit_account_code"], "2103006")

    def test_pcb_case1_payload_exports_case4_trace_and_target_hint_fields(self) -> None:
        payload = build_pcb_case1_export_payload(
            database="hwgroup_erp",
            rows=[
                {
                    "row_key": "case4::7004::901",
                    "cycle_key": "case4::7004",
                    "company_id": 7,
                    "company_name": "Alpha Company",
                    "picking_name": "LHPK/IN/7004",
                    "source_label": "Item Balance",
                    "po_name": "PO/2026/0004",
                    "bill_name": "BILL/2026/0004",
                    "stj_refs": ["STJ/2026/0454"],
                    "date": "2026-03-19",
                    "reference": "Correction: Purchase Cycle Balance: LHPK/IN/7004 / BILL/2026/0004 / SKU-001",
                    "journal_code": "STJ",
                    "amount": 100.0,
                    "planned_lines": [
                        {"role": "problem_2103006", "account_code": "2103006", "amount": 100.0, "side": "debit"},
                        {"role": "selisih_hpp", "account_code": "5101004", "amount": 100.0, "side": "credit"},
                    ],
                    "row_status": "ready",
                    "row_status_message": "Ready",
                    "product_id": 901,
                    "item_code": "SKU-001",
                    "item_name": "Produk A",
                    "item_category_name": "Raw",
                    "bill_line_id": 8104,
                    "bill_move_id": 8204,
                    "purchase_line_id": 8304,
                    "stock_move_id": 8404,
                    "stock_move_ids": [8404, 8405],
                    "stj_move_ids": [8804],
                    "picking_id": 7004,
                    "payment_move_ids": [8504],
                    "bank_move_ids": [8604],
                    "suspend_target_aml_ids": [9104],
                    "clearing_target_aml_ids": [9105],
                    "partner_id": 77,
                    "partner_name": "Vendor Alpha",
                    "product_uom_id": 11,
                    "quantity": 2.0,
                    "currency_id": 13,
                    "amount_currency": 120.0,
                    "analytic_distribution": {"CC-01": 100.0},
                    "bill_price_unit": 60.0,
                    "gr_price_unit": 50.0,
                    "price_gap_value": 20.0,
                    "allocated_amount": 120.0,
                }
            ],
            exported_at="2026-03-20T10:00:00",
        )

        self.assertEqual(payload["rows"][0]["payment_move_ids"], "8504")
        self.assertEqual(payload["rows"][0]["bank_move_ids"], "8604")
        self.assertEqual(payload["rows"][0]["suspend_target_aml_ids"], "9104")
        self.assertEqual(payload["rows"][0]["clearing_target_aml_ids"], "9105")
        self.assertEqual(payload["rows"][0]["currency_id"], 13)

    def test_pcb_case1_payload_includes_resolve_and_review_metadata(self) -> None:
        rows = _sample_pcb_rows()
        rows[0]["resolve_account_code"] = "1108099"
        rows[0]["resolve_account_preview"] = "1108099 - Clearing Override [Global Extra]"
        rows[0]["review_required"] = True
        rows[0]["review_confirmed"] = False
        rows[0]["review_reason"] = "Manual review wajib sebelum execute."

        payload = build_pcb_case1_export_payload(
            database="hwgroup_erp",
            rows=rows,
            exported_at="2026-03-20T10:00:00",
        )

        self.assertEqual(payload["rows"][0]["resolve_account_code"], "1108099")
        self.assertIn("1108099", payload["rows"][0]["resolve_account_preview"])
        self.assertTrue(payload["rows"][0]["review_required"])
        self.assertFalse(payload["rows"][0]["review_confirmed"])
        self.assertEqual(payload["rows"][0]["review_reason"], "Manual review wajib sebelum execute.")

    def test_payload_groups_counts_by_effective_mode(self) -> None:
        payload = build_repair_summary_export_payload(database="hwgroup_erp", results=_sample_results(), exported_at="2026-03-20T10:00:00")

        self.assertEqual(payload["company_label"], "Alpha Company - [7]")
        self.assertEqual(payload["success_count"], 2)
        self.assertEqual(payload["error_count"], 0)
        self.assertEqual(
            payload["mode_counts"],
            [
                {"mode": "New JE + Relink SVL", "count": 1},
                {"mode": "Created New JE", "count": 1},
            ],
        )
        self.assertIn("dari STJ/2025/10/0435", payload["rows"][0]["detail_sentence"])
        self.assertIn("sejak awal belum ada Journal Entry", payload["rows"][1]["detail_sentence"])

    def test_export_repair_summary_html_contains_document_numbers(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = export_repair_summary_html(
                database="hwgroup_erp",
                results=_sample_results(),
                output_path=str(Path(tmp_dir) / "repair_summary.html"),
                exported_at="2026-03-20T10:00:00",
            )
            html = path.read_text(encoding="utf-8")

        self.assertIn("Alpha Company - [7]", html)
        self.assertIn("New JE + Relink SVL", html)
        self.assertIn("Created New JE", html)
        self.assertIn("STJ/2026/09001", html)
        self.assertIn("STJ/2025/10/0435", html)
        self.assertIn("WCGT/IN/00034", html)
        self.assertIn("Persediaan Alkohol", html)

    def test_export_repair_summary_excel_writes_summary_and_flat_detail(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = export_repair_summary_excel(
                database="hwgroup_erp",
                results=_sample_results(),
                output_path=str(Path(tmp_dir) / "repair_summary.xlsx"),
                exported_at="2026-03-20T10:00:00",
            )
            workbook = load_workbook(path)
            try:
                self.assertEqual(workbook.sheetnames, ["Summary", "Detail Flat"])
                summary_sheet = workbook["Summary"]
                detail_sheet = workbook["Detail Flat"]
                self.assertEqual(summary_sheet["B1"].value, "Alpha Company - [7]")
                self.assertEqual(summary_sheet["B2"].value, "hwgroup_erp")
                self.assertEqual(summary_sheet["B4"].value, 2)
                self.assertEqual(summary_sheet["A9"].value, "Mode")
                self.assertEqual(summary_sheet["A10"].value, "New JE + Relink SVL")
                self.assertEqual(summary_sheet["B10"].value, 1)
                self.assertEqual(detail_sheet["E2"].value, "New JE + Relink SVL")
                self.assertEqual(detail_sheet["J2"].value, "STJ/2025/10/0435")
                self.assertEqual(detail_sheet["K2"].value, "STJ/2026/09001")
                self.assertEqual(detail_sheet["J3"].value, None)
                self.assertEqual(detail_sheet["I3"].value, "WCGT/IN/00035")
            finally:
                workbook.close()

    def test_export_pcb_case1_summary_excel_writes_result_journal_and_trace_fields(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = export_pcb_case1_summary_excel(
                database="hwgroup_erp",
                rows=_sample_pcb_rows(),
                output_path=str(Path(tmp_dir) / "pcb_case1_summary.xlsx"),
                exported_at="2026-03-20T10:00:00",
            )
            workbook = load_workbook(path)
            try:
                self.assertEqual(workbook.sheetnames, ["Summary", "Detail Flat"])
                summary_sheet = workbook["Summary"]
                detail_sheet = workbook["Detail Flat"]
                self.assertEqual(summary_sheet["B1"].value, "Alpha Company - [7]")
                self.assertEqual(summary_sheet["B5"].value, 1)
                self.assertEqual(summary_sheet["B7"].value, 1)
                self.assertEqual(detail_sheet["O2"].value, "STJ/2026/0901")
                self.assertEqual(detail_sheet["J2"].value, "STJ/2026/0451")
                self.assertEqual(detail_sheet["AU2"].value, "9102")
                self.assertEqual(detail_sheet["AV2"].value, "9101")
            finally:
                workbook.close()

    def test_build_pcb_cycle_detail_export_payload_carries_raw_rows_notes_and_hidden_metadata(self) -> None:
        page = SvlFixJeDashboardPage.__new__(SvlFixJeDashboardPage)
        page._latest_snapshot = SvlDashboardSnapshot(
            database="hwgroup_erp",
            company_id=7,
            company_name="Alpha Company",
            generated_at="2026-03-20T10:00:00",
            period="2026-03-01..2026-03-31",
        )
        page._pcb_raw_sort_var = type("_SortVar", (), {"get": lambda self: "Sort by Process"})()

        payload = page._build_pcb_cycle_detail_export_payload(_sample_pcb_cycle_for_export())

        self.assertEqual(payload["company_name"], "Alpha Company")
        self.assertEqual(payload["cycle_name"], "PIKKP/IN/02791 [+2]")
        self.assertEqual(payload["sort_mode"], "Sort by Process")
        self.assertEqual(payload["raw_rows"][0]["process_group"], "GR / STJ")
        self.assertEqual(payload["raw_rows"][0]["kode_transaksi"], "STJ/2026/03/0262")
        self.assertTrue(all(row["cycle_name"] == "PIKKP/IN/02791 [+2]" for row in payload["raw_rows"]))
        self.assertTrue(any(row["row_kind"] == "item_evidence" for row in payload["detail_rows"]))
        self.assertTrue(any(row["row_kind"] == "external_clearing" for row in payload["detail_rows"]))
        self.assertTrue(any(row["field_name"] == "external_clearing_refs" for row in payload["hidden_rows"]))
        self.assertTrue(any(row["field_name"] == "case_counts" for row in payload["hidden_rows"]))

    def test_export_pcb_cycle_detail_excel_writes_raw_sheet_detail_sheet_and_hidden_sheet(self) -> None:
        page = SvlFixJeDashboardPage.__new__(SvlFixJeDashboardPage)
        page._latest_snapshot = SvlDashboardSnapshot(
            database="hwgroup_erp",
            company_id=7,
            company_name="Alpha Company",
            generated_at="2026-03-20T10:00:00",
            period="2026-03-01..2026-03-31",
        )
        page._pcb_raw_sort_var = type("_SortVar", (), {"get": lambda self: "Sort by Process"})()
        payload = page._build_pcb_cycle_detail_export_payload(_sample_pcb_cycle_for_export())

        with tempfile.TemporaryDirectory() as tmp_dir:
            path = export_pcb_cycle_detail_excel(
                payload=payload,
                output_path=str(Path(tmp_dir) / "pcb_cycle_detail.xlsx"),
            )
            workbook = load_workbook(path)
            try:
                self.assertEqual(workbook.sheetnames, ["Summary", "Detail JE", "Detail Akun & Catatan", "Data Tidak Tampil"])
                summary_sheet = workbook["Summary"]
                raw_sheet = workbook["Detail JE"]
                detail_sheet = workbook["Detail Akun & Catatan"]
                hidden_sheet = workbook["Data Tidak Tampil"]
                self.assertEqual(summary_sheet["B1"].value, "Alpha Company")
                self.assertEqual(summary_sheet["B4"].value, "PIKKP/IN/02791 [+2]")
                self.assertEqual(raw_sheet["J2"].value, "GR / STJ")
                self.assertEqual(raw_sheet["L2"].value, "STJ/2026/03/0262")
                self.assertEqual(raw_sheet["B2"].value, "Alpha Company")
                self.assertTrue(any(cell.value == "item_evidence" for cell in detail_sheet["C"]))
                self.assertTrue(any(cell.value == "external_clearing_refs" for cell in hidden_sheet["F"]))
            finally:
                workbook.close()

    def test_export_pcb_company_audit_enriches_raw_ledger_po_from_cycle(self) -> None:
        cycle = _sample_pcb_cycle_for_export()
        cycle.raw_lines.append(
            {
                "tanggal": "2026-03-06",
                "kode_transaksi": "BILL/2026/03/0082",
                "jenis": "BILL",
                "tipe_akun": "liability_current",
                "akun_code": "2103006",
                "akun_name": "Hutang Suspensed Pengadaan Barang",
                "kode_item": "",
                "nama_item": "",
                "uom": "",
                "qty_item": 0.0,
                "kategori_produk": "",
                "no_po": "",
                "komunikasi": "Vendor bill payable line",
                "debit": 0.0,
                "kredit": 2868000.0,
                "saldo": -2868000.0,
                "matching": "",
                "partner": "SENTRAL CK - WH",
            }
        )
        snapshot = SvlDashboardSnapshot(
            database="hwgroup_erp",
            company_id=7,
            company_name="Alpha Company",
            generated_at="2026-03-20T10:00:00",
            period="2026-03-01..2026-03-31",
            purchase_cycles=[cycle],
        )

        with tempfile.TemporaryDirectory() as tmp_dir:
            path = export_pcb_company_audit_excel(
                snapshot=snapshot,
                output_path=str(Path(tmp_dir) / "pcb_company_audit.xlsx"),
            )
            workbook = load_workbook(path)
            try:
                self.assertIn("5. Raw Ledger PO Bill", workbook.sheetnames)
                raw_sheet = workbook["4. Raw Ledger"]
                po_audit_sheet = workbook["5. Raw Ledger PO Bill"]
                self.assertEqual(raw_sheet["S2"].value, "No PO")
                self.assertEqual(raw_sheet["S5"].value, "PO/PIKKPH/2026/02/02455")
                self.assertEqual(po_audit_sheet["S5"].value, "PO/PIKKPH/2026/02/02455")
                self.assertIsNone(po_audit_sheet["Z5"].value)
                self.assertEqual(po_audit_sheet["AA5"].value, "cycle_bill_po")
                self.assertEqual(po_audit_sheet["AB5"].value, "PO/PIKKPH/2026/02/02455")
                self.assertEqual(po_audit_sheet["AC5"].value, "BILL/2026/03/0082")
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
