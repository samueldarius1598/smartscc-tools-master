import json
import tempfile
import unittest
from pathlib import Path

from smartscc_tools.features.item_journal.observability.audit import migrate_audit_jsonl, migrate_audit_sheet
from smartscc_tools.features.item_journal.workbook import AUDIT_SHEET_NAME
from openpyxl import Workbook, load_workbook


class AuditMigrateTest(unittest.TestCase):
    def test_migrate_jsonl_legacy_to_v2(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            src = Path(tmp_dir) / "legacy.jsonl"
            dst = Path(tmp_dir) / "v2.jsonl"
            payload = {
                "run_id": "r1",
                "timestamp": "2026-03-05 01:00:00",
                "row_number": 2,
                "group_key": "g1",
                "batch_seq": 1,
                "stage": "ROW_ERROR",
                "model": "stock.picking",
                "method": "create",
                "attempt": 1,
                "http_status": 200,
                "duration_ms": 240000,
                "result": "ERROR",
                "message": "stock.picking.create timeout.",
                "picking_id": 0,
            }
            src.write_text(json.dumps(payload, ensure_ascii=False) + "\n", encoding="utf-8")
            result = migrate_audit_jsonl(str(src), str(dst))
            self.assertEqual(result["migrated"], 1)
            migrated = json.loads(dst.read_text(encoding="utf-8").strip())
            self.assertEqual(migrated["schema_version"], "v2")
            self.assertEqual(migrated["run_id"], "r1")
            self.assertEqual(migrated["cause_code"], "transport_timeout")

    def test_migrate_sheet(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            src = Path(tmp_dir) / "legacy.xlsm"
            dst = Path(tmp_dir) / "migrated.xlsm"
            wb = Workbook()
            ws = wb.active
            ws.title = "Item Journal"
            ws["G1"] = "Product ID"
            ws["H1"] = "Product key (kode/nama)"
            audit_ws = wb.create_sheet(AUDIT_SHEET_NAME)
            audit_ws.append(
                [
                    "run_id",
                    "timestamp",
                    "row_number",
                    "group_key",
                    "batch_seq",
                    "stage",
                    "model",
                    "method",
                    "attempt",
                    "http_status",
                    "duration_ms",
                    "result",
                    "message",
                    "picking_id",
                ]
            )
            audit_ws.append(
                ["r2", "2026-03-05 01:00:00", 2, "g2", 1, "RUN_DONE", "stock.picking", "create", 0, 200, 100, "DONE", "ok", 1]
            )
            wb.save(src)
            wb.close()

            result = migrate_audit_sheet(str(src), str(dst))
            self.assertEqual(result["migrated"], 1)
            wb2 = load_workbook(filename=str(dst), keep_vba=True, keep_links=False)
            try:
                ws2 = wb2[AUDIT_SHEET_NAME]
                self.assertEqual(str(ws2["A1"].value), "schema_version")
                self.assertEqual(str(ws2["B1"].value), "run_id")
                self.assertEqual(str(ws2["A2"].value), "v2")
            finally:
                wb2.close()


if __name__ == "__main__":
    unittest.main()
