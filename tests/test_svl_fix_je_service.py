import asyncio
import logging
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from smartscc_tools.features.svl_fix_je.export_results import export_results_to_excel
from smartscc_tools.features.svl_fix_je.models import SvlFixJeRunRequest
from smartscc_tools.features.svl_fix_je.service import SvlFixJeServiceAsync, build_move_values


def _write_workbook(path: Path, rows: list[list[object]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SVL_Fix"
    sheet.append(
        [
            "No",
            "SVL ID",
            "SVL Reference",
            "default_code",
            "Qty",
            "UoM",
            "Unit Cost",
            "Total Value",
            "COA Credit",
            "COA Debit",
            "Journal Code",
            "Tanggal JE",
            "Keterangan",
        ]
    )
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


class _FakeRpc:
    def __init__(self) -> None:
        self.uid = 77
        self.login_calls = 0
        self.fields_map = {
            "account.account": {"company_ids": {"type": "many2many"}},
            "stock.valuation.layer": {"account_move_id": {"readonly": False}},
        }
        self.svls = {
            5139103: {
                "id": 5139103,
                "description": "WCGT/IN/00892 - Receipt",
                "product_id": [101, "Demo Product"],
                "quantity": 20.0,
                "value": 500.0,
                "unit_cost": 25.0,
                "account_move_id": False,
                "stock_move_id": [91, "MOVE/91"],
                "remaining_qty": 0.0,
                "remaining_value": 0.0,
                "company_id": [1, "HWG"],
            },
            5139104: {
                "id": 5139104,
                "description": "WCGT/OUT/00893 - Delivery",
                "product_id": [102, "Demo Product 2"],
                "quantity": 10.0,
                "value": -250.0,
                "unit_cost": -25.0,
                "account_move_id": False,
                "stock_move_id": [92, "MOVE/92"],
                "remaining_qty": 0.0,
                "remaining_value": 0.0,
                "company_id": [1, "HWG"],
            },
            5139105: {
                "id": 5139105,
                "description": "WCGT/IN/00894 - Receipt",
                "product_id": [103, "Demo Product 3"],
                "quantity": 5.0,
                "value": 150.0,
                "unit_cost": 30.0,
                "account_move_id": False,
                "stock_move_id": [93, "MOVE/93"],
                "remaining_qty": 0.0,
                "remaining_value": 0.0,
                "company_id": [1, "HWG"],
            },
        }
        self.moves = {
            91: {"id": 91, "reference": "WCGT/IN/00892"},
            92: {"id": 92, "reference": "WCGT/OUT/00893"},
            93: {"id": 93, "reference": "WCGT/IN/00894"},
        }
        self.products = {
            101: {"id": 101, "default_code": "F-FHVF-0167", "uom_id": [1, "Units"]},
            102: {"id": 102, "default_code": "F-FHVF-0168", "uom_id": [1, "Units"]},
            103: {"id": 103, "default_code": "F-FHVF-0169", "uom_id": [1, "Units"]},
        }
        self.accounts = [
            {"id": 10, "code": "1.1.03.01", "name": "Persediaan", "company_ids": [1]},
            {"id": 11, "code": "5.1.01.01", "name": "Selisih", "company_ids": [1]},
        ]
        self.journals = [{"id": 12, "code": "STJ", "name": "Stock Journal", "company_id": 1}]
        self.created_moves: list[dict] = []
        self.write_calls: list[tuple[str, list[int], dict]] = []
        self.fail_write = False
        self.batch_post_error: str | None = None
        self.per_move_post_errors: dict[int, str] = {}
        self.create_delay = 0.0

    async def ensure_login(self) -> int:
        self.login_calls += 1
        return self.uid

    async def fields_get(self, model, attributes=None, context=None, stage=""):  # noqa: ANN001
        return self.fields_map.get(model, {})

    async def read(self, model, ids, fields=None, context=None, stage="", excel_row=0):  # noqa: ANN001
        source = {}
        if model == "stock.valuation.layer":
            source = self.svls
        elif model == "stock.move":
            source = self.moves
        elif model == "product.product":
            source = self.products
        return [dict(source[item]) for item in ids if item in source]

    async def search_read(self, model, domain, fields=None, limit=None, context=None, stage="", order=None):  # noqa: ANN001
        records = self.accounts if model == "account.account" else self.journals
        code_eq = None
        code_in = None
        code_prefix = None
        company_id = None
        for field_name, operator, value in domain:
            if field_name == "code" and operator == "=":
                code_eq = value
            elif field_name == "code" and operator == "in":
                code_in = set(value)
            elif field_name == "code" and operator == "=like":
                code_prefix = str(value).replace("%", "")
            elif field_name in {"company_ids", "company_id"}:
                company_id = value[0] if isinstance(value, list) else value
        out = []
        for record in records:
            code = record.get("code")
            if code_eq is not None and code != code_eq:
                continue
            if code_in is not None and code not in code_in:
                continue
            if code_prefix is not None and not str(code).startswith(code_prefix):
                continue
            if company_id is not None:
                company_values = record.get("company_ids") or []
                if isinstance(company_values, int):
                    company_values = [company_values]
                if company_id not in company_values and record.get("company_id") != company_id:
                    continue
            out.append(dict(record))
        if limit:
            out = out[:limit]
        return out

    async def create(self, model, values, context=None, stage="", excel_row=0):  # noqa: ANN001
        if self.create_delay > 0:
            await asyncio.sleep(self.create_delay)
        move_id = 1000 + len(self.created_moves) + 1
        self.created_moves.append({"model": model, "values": values, "context": context})
        return move_id

    async def write(self, model, ids, values, context=None, stage="", excel_row=0):  # noqa: ANN001
        if self.fail_write:
            raise RuntimeError("field readonly")
        self.write_calls.append((model, list(ids), dict(values)))
        return True

    async def execute_kw(
        self,
        model,
        method,
        args=None,
        kwargs=None,
        stage="",
        excel_row=0,
        mutating=False,
    ):  # noqa: ANN001
        if model == "account.move" and method == "action_post":
            move_ids = list(args[0])
            if self.batch_post_error and len(move_ids) > 1:
                raise RuntimeError(self.batch_post_error)
            for move_id in move_ids:
                if move_id in self.per_move_post_errors:
                    raise RuntimeError(self.per_move_post_errors[move_id])
            return True
        raise AssertionError(f"Unexpected execute_kw: {model}.{method}")


class SvlFixJeServiceTest(unittest.IsolatedAsyncioTestCase):
    async def test_build_move_values_uses_sign_based_direction(self) -> None:
        from smartscc_tools.features.svl_fix_je.models import SvlFixJeExcelRow, SvlFixJeValidatedRow

        positive_row = SvlFixJeExcelRow(row_number=2, svl_id=1, total_value=500.0, je_date="2026-01-01")
        negative_row = SvlFixJeExcelRow(row_number=3, svl_id=2, total_value=-500.0, je_date="2026-01-01")
        validated = SvlFixJeValidatedRow(
            row=positive_row,
            svl_record={"id": 1},
            company_id=1,
            company_name="HWG",
            product_id=0,
            credit_account_id=10,
            debit_account_id=11,
            journal_id=12,
        )

        positive_vals = build_move_values(positive_row, validated, "FIX-SVL")
        negative_vals = build_move_values(negative_row, validated, "FIX-SVL")

        self.assertEqual(positive_vals["line_ids"][0][2]["account_id"], 10)
        self.assertEqual(positive_vals["line_ids"][1][2]["account_id"], 11)
        self.assertEqual(negative_vals["line_ids"][0][2]["account_id"], 11)
        self.assertEqual(negative_vals["line_ids"][1][2]["account_id"], 10)

    async def test_execute_creates_draft_move_and_link_failure_is_warning_only(self) -> None:
        rpc = _FakeRpc()
        rpc.fail_write = True
        logs: list[str] = []
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "svl_fix.xlsx"
            _write_workbook(
                path,
                [[1, 5139103, "WCGT/IN/00892", "F-FHVF-0167", 20, "Units", 25, 500, "1.1.03.01", "5.1.01.01", "STJ", "2026-01-07", ""]],
            )
            service = SvlFixJeServiceAsync(rpc=rpc, logger=logging.getLogger("test.svl_fix"), on_log=logs.append)
            summary = await service.execute(
                SvlFixJeRunRequest(
                    excel_path=str(path),
                    database="hwgroup_erp",
                    ref_prefix="FIX-SVL",
                    auto_post=False,
                    max_workers=1,
                )
            )

        self.assertEqual(summary.created_count, 1)
        self.assertEqual(summary.execution_errors, 0)
        self.assertEqual(summary.results[0].status, "CREATED")
        self.assertTrue(any("readonly" in item.lower() for item in logs))

    async def test_execute_auto_post_falls_back_per_move_and_marks_partial_failure(self) -> None:
        rpc = _FakeRpc()
        rpc.batch_post_error = "batch failed"
        rpc.per_move_post_errors = {1002: "lock date"}
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "svl_fix.xlsx"
            _write_workbook(
                path,
                [
                    [1, 5139103, "WCGT/IN/00892", "F-FHVF-0167", 20, "Units", 25, 500, "1.1.03.01", "5.1.01.01", "STJ", "2026-01-07", ""],
                    [2, 5139104, "WCGT/OUT/00893", "F-FHVF-0168", 10, "Units", -25, -250, "1.1.03.01", "5.1.01.01", "STJ", "2026-01-07", ""],
                ],
            )
            service = SvlFixJeServiceAsync(rpc=rpc, logger=logging.getLogger("test.svl_fix"))
            summary = await service.execute(
                SvlFixJeRunRequest(
                    excel_path=str(path),
                    database="hwgroup_erp",
                    ref_prefix="FIX-SVL",
                    auto_post=True,
                    max_workers=2,
                )
            )

        statuses = [item.status for item in summary.sorted_results()]
        self.assertEqual(statuses, ["POSTED", "ERROR"])
        self.assertEqual(summary.posted_count, 1)
        self.assertEqual(summary.execution_errors, 1)

    async def test_execute_stop_cancels_remaining_rows(self) -> None:
        rpc = _FakeRpc()
        rpc.create_delay = 0.02
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "svl_fix.xlsx"
            _write_workbook(
                path,
                [
                    [1, 5139103, "WCGT/IN/00892", "F-FHVF-0167", 20, "Units", 25, 500, "1.1.03.01", "5.1.01.01", "STJ", "2026-01-07", ""],
                    [2, 5139104, "WCGT/OUT/00893", "F-FHVF-0168", 10, "Units", -25, -250, "1.1.03.01", "5.1.01.01", "STJ", "2026-01-07", ""],
                    [3, 5139105, "WCGT/IN/00894", "F-FHVF-0169", 5, "Units", 30, 150, "1.1.03.01", "5.1.01.01", "STJ", "2026-01-07", ""],
                ],
            )
            service = SvlFixJeServiceAsync(rpc=rpc, logger=logging.getLogger("test.svl_fix"))

            def on_result(result) -> None:  # noqa: ANN001
                if result.status == "CREATED":
                    service.request_stop()

            service.on_result = on_result
            summary = await service.execute(
                SvlFixJeRunRequest(
                    excel_path=str(path),
                    database="hwgroup_erp",
                    ref_prefix="FIX-SVL",
                    auto_post=False,
                    max_workers=1,
                )
            )

        self.assertTrue(summary.stopped)
        self.assertEqual(summary.created_count, 1)
        self.assertEqual(summary.canceled_count, 2)
        self.assertEqual([item.status for item in summary.sorted_results()], ["CREATED", "CANCELED", "CANCELED"])

    async def test_export_results_to_excel_writes_result_rows(self) -> None:
        rpc = _FakeRpc()
        with tempfile.TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "svl_fix.xlsx"
            output_path = Path(tmp_dir) / "results.xlsx"
            _write_workbook(
                source_path,
                [[1, 5139103, "WCGT/IN/00892", "F-FHVF-0167", 20, "Units", 25, 500, "1.1.03.01", "5.1.01.01", "STJ", "2026-01-07", ""]],
            )
            service = SvlFixJeServiceAsync(rpc=rpc, logger=logging.getLogger("test.svl_fix"))
            summary = await service.validate(
                SvlFixJeRunRequest(
                    excel_path=str(source_path),
                    database="hwgroup_erp",
                    ref_prefix="FIX-SVL",
                    auto_post=False,
                    max_workers=1,
                )
            )

            exported = export_results_to_excel(summary, str(output_path))
            workbook = load_workbook(exported, data_only=True)
            sheet = workbook.active
            values = list(sheet.iter_rows(values_only=True))
            workbook.close()

        self.assertEqual(exported.name, "results.xlsx")
        self.assertEqual(values[1][1], 5139103)
        self.assertEqual(values[1][4], "VALID")
