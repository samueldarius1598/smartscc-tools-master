"""A valued receipt must stay visible and repairable before any journal exists."""
import logging
from pathlib import Path
import tempfile
import unittest
from unittest.mock import AsyncMock

from smartscc_tools.features.svl_fix_je.dashboard_service import SvlDashboardServiceAsync as Service
from smartscc_tools.features.svl_fix_je.dashboard_repair_service import SvlDashboardRepairServiceAsync as Repair
from smartscc_tools.features.svl_fix_je.models import SvlDashboardRequest, SvlDashboardSnapshot
from tests.test_pcb_normal_receipt import NormalRpc


def fixture(*, value=100, valuation_method="automated", input_code="2103006", linked=False):
    move = dict(id=4, company_id=[1, "Company"], picking_id=[6, "TEST/IN/01"], state="done",
                product_id=[1, "Supplies"], purchase_line_id=[3, "PO"], product_qty=2,
                product_uom=[7, "Units"], account_move_ids=[])
    layer = dict(id=9, company_id=[1, "Company"], product_id=[1, "Supplies"],
                 stock_move_id=[4, "Move"], quantity=2, value=value,
                 account_move_id=[8, "Existing JE"] if linked else False)
    trace = dict(picking_rows_by_id={6: dict(id=6, name="TEST/IN/01", state="done", scheduled_date="2026-01-01", partner_id=[77, "Vendor"])},
                 po_id_by_picking={6: 2}, product_ids_by_picking={6: [1]},
                 inventory_types_by_picking={6: ["purchase"]},
                 purchase_order_rows_by_id={2: dict(id=2, name="PO", partner_id=[77, "Vendor"], invoice_ids=[])},
                 purchase_line_rows_by_id={3: dict(id=3, order_id=[2, "PO"], product_id=[1, "Supplies"], product_qty=2, price_unit=25)},
                 purchase_line_product_map={3: 1}, purchase_line_po_name_map={3: "PO"},
                 stock_move_rows_by_id={4: move}, svl_rows_by_stock_move_id={4: [layer]})
    product = {1: dict(default_code="TEST", name="Supplies", valuation_method=valuation_method,
                       valuation_account_code="1105006", valuation_account_name="Inventory",
                       input_account_code=input_code, standard_price=75, uom_id=7, uom_name="Units")}
    service = Service(rpc=object(), logger=logging.getLogger("test"))
    cycles = service._build_purchase_cycles(trace=trace, all_ledger_rows=[], account_info_map={},
        move_info_map={}, product_info_map=product, warnings=[],
        problem_codes=frozenset({"2103006", "1108099"}), info_codes=frozenset({"11120003"}))
    return cycles, trace, layer


class UnbilledReceiptTest(unittest.TestCase):
    def test_no_journals_still_has_cycle_item_and_reviewed_recovery(self):
        cycles, _, _ = fixture()
        self.assertEqual(len(cycles), 1)
        cycle = cycles[0]
        self.assertEqual((cycle.primary_case, cycle.cycle_status), ("case5", "problem"))
        self.assertNotIn("case10", cycle.edge_flags)
        self.assertEqual(cycle.raw_lines, [])
        self.assertEqual(cycle.account_rows, [])
        self.assertEqual((cycle.total_debit, cycle.total_credit), (0, 0))
        item = cycle.item_rows[0]
        self.assertEqual((item.product_name, item.valuation_method, item.gr_quantity), ("Supplies", "automated", 2))
        self.assertEqual((item.stj_state, item.svl_zero_at_gr), ("missing_actual_svl", False))
        row = cycle.case2_repair_rows[0]
        self.assertEqual([(line.account_code, line.side, line.amount) for line in row.planned_lines],
                         [("1105006", "debit", 100), ("2103006", "credit", 100)])
        self.assertEqual((row.bill_move_id, row.bill_line_id, row.quantity), (0, 0, 2))
        self.assertTrue(row.review_required)
        self.assertFalse(row.review_confirmed)
        self.assertEqual(row.case_evidence["additional_expense"], 0)
        self.assertEqual(row.case_evidence["cost_change_offset"], 0)
        self.assertEqual(row.case_evidence["bill_standard_gap"], 0)

    def test_empty_or_zero_svl_does_not_add_empty_cycle(self):
        cycles, _, _ = fixture(value=0)
        self.assertEqual(cycles, [])

    def test_manual_valuation_is_visible_without_automatic_receipt_repair(self):
        cycles, _, _ = fixture(valuation_method="manual")
        self.assertEqual(len(cycles), 1)
        self.assertEqual(cycles[0].case2_repair_rows, [])
        self.assertNotEqual(cycles[0].primary_case, "case5")

    def test_input_account_mismatch_blocks_default_plan(self):
        cycles, _, _ = fixture(input_code="2103999")
        row = cycles[0].case2_repair_rows[0]
        self.assertEqual(row.planned_lines, [])
        self.assertIn("receipt_recovery_blocked", row.guard_flags)

    def test_already_linked_receipt_is_not_offered_again(self):
        cycles, _, _ = fixture(linked=True)
        self.assertEqual(cycles[0].case2_repair_rows, [])
        self.assertNotEqual(cycles[0].primary_case, "case5")

    def test_collection_review_message_describes_actual_svl_basis(self):
        from smartscc_tools.modules.svl_fix_je_dashboard_pcb_ui import _pcb_case2_guard_reason_preview
        cycles, _, _ = fixture()
        row = cycles[0].case2_repair_rows[0]
        message = _pcb_case2_guard_reason_preview(dict(case_evidence=row.case_evidence,
            guard_flags=row.guard_flags, guard_messages=row.guard_messages))
        self.assertEqual(message, row.review_reason)
        self.assertIn("Bill belum tersedia", message)

    def test_native_excel_includes_receipt_in_overview_and_items_without_fake_ledger(self):
        from openpyxl import load_workbook
        from smartscc_tools.features.svl_fix_je.dashboard_pcb_export import export_pcb_company_audit_excel
        cycles, _, _ = fixture()
        snapshot = SvlDashboardSnapshot(database="test", company_id=1, company_name="Company",
            generated_at="2026-10-12T10:00:00", period="All", dataset_mode="purchase_cycle_balance", purchase_cycles=cycles)
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "pcb.xlsx"
            export_pcb_company_audit_excel(snapshot=snapshot, output_path=str(path))
            book = load_workbook(path, read_only=True)
            try:
                hits = [any(cell == "TEST/IN/01" for row in sheet.iter_rows(values_only=True) for cell in row) for sheet in book]
                self.assertEqual(hits[:2], [True, True])
                self.assertEqual(hits[3:], [False, False])
            finally:
                book.close()


class UnbilledReceiptPipelineTest(unittest.IsolatedAsyncioTestCase):
    async def test_analysis_fetches_receipt_product_even_when_ledger_is_empty(self):
        _, trace, _ = fixture()
        product = {1: dict(name="Supplies", default_code="TEST", valuation_method="automated",
                          valuation_account_code="1105006", input_account_code="2103006")}
        service = Service(rpc=object(), logger=logging.getLogger("test"))
        service._fetch_company_name = AsyncMock(return_value="Company")
        service._collect_account_balance_trace = AsyncMock(return_value=trace)
        service._find_pcb_correction_move_ids = AsyncMock(return_value=[])
        service._search_read_in_chunks = AsyncMock(return_value=[])
        service._fetch_account_info_map = AsyncMock(return_value={})
        service._fetch_pcb_move_info = AsyncMock(return_value={})
        service._fetch_pcb_product_info_map = AsyncMock(return_value=product)
        service._fetch_pcb_adjustment_audit_moves = AsyncMock(return_value=([], {}, product))
        snapshot = await service._analyze_purchase_cycle_balance(
            request=SvlDashboardRequest(company_id=1, database="test", dataset_mode="purchase_cycle_balance"),
            date_from="", date_to="", context={}, capabilities={"aml_fields": {}}, warnings=[])
        self.assertEqual(service._fetch_pcb_product_info_map.call_args.kwargs["product_ids"], [1])
        self.assertEqual(snapshot.purchase_cycles[0].primary_case, "case5")
        self.assertEqual(len(snapshot.purchase_cycles[0].case2_repair_rows), 1)


class UnbilledRpc(NormalRpc):
    def __init__(self, layer, trace):
        super().__init__(layer)
        self.stock = dict(trace["stock_move_rows_by_id"][4])
        self.purchase = dict(id=3, company_id=[1, "Company"], product_id=[1, "Supplies"], order_id=[2, "PO"], invoice_lines=[])
        self.order = dict(id=2, name="PO", company_id=[1, "Company"], invoice_ids=[])
        self.origin_bill = False
        self.category = dict(id=5, property_valuation="real_time", property_stock_valuation_account_id=[101, "Inventory"],
                             property_stock_account_input_categ_id=[301, "Suspense"])
        self.fields_map["stock.move"] = {key: {} for key in self.stock}
        self.fields_map["product.category"] = {key: {"type": "many2one" if "account" in key else "selection"} for key in self.category}
        self.other_journal = False
        self.sibling_link = False

    async def read(self, model, ids, **kwargs):
        if model == "stock.move":
            return [dict(self.stock)] if ids == [4] else []
        if model == "product.product":
            return [dict(id=1, categ_id=[5, "Category"])]
        if model == "product.category":
            return [dict(self.category)]
        if model == "purchase.order.line":
            return [dict(self.purchase)]
        if model == "purchase.order":
            return [dict(self.order)]
        return await super().read(model, ids, **kwargs)

    async def search_read(self, model, domain, **kwargs):
        if kwargs.get("stage") == "PCB_RECEIPT_BILL_ORIGIN_PREFLIGHT":
            return [dict(id=99)] if self.origin_bill else []
        if kwargs.get("stage") == "PCB_RECEIPT_JE_PREFLIGHT":
            return [dict(id=99, name="OTHER/JE")] if self.other_journal else []
        if kwargs.get("stage") == "PCB_RECEIPT_SIBLING_PREFLIGHT":
            return [dict(id=10)] if self.sibling_link else []
        return await super().search_read(model, domain, **kwargs)


class UnbilledReceiptExecuteTest(unittest.IsolatedAsyncioTestCase):
    def setup_row(self):
        cycles, trace, layer = fixture()
        row = cycles[0].case2_repair_rows[0]
        row.company_id = 1
        row.partner_id = 77
        row.date = "2026-09-01"
        row.reference = "Correction: Purchase Cycle Balance: TEST/IN/01 / TEST"
        rpc = UnbilledRpc(layer, trace)
        return row, rpc, Repair(rpc=rpc, logger=logging.getLogger("test"))

    async def test_review_gate_draft_post_relink_and_retry(self):
        row, rpc, repair = self.setup_row()
        with self.assertRaisesRegex(Exception, "Review|review"):
            await repair._execute_pcb_case2_row(row, posting_mode="draft")
        self.assertEqual(rpc.create_calls, [])
        row.review_confirmed = True
        draft = await repair._execute_pcb_case2_row(row, posting_mode="draft")
        self.assertFalse(draft.posted)
        self.assertFalse(rpc.svls[9]["account_move_id"])
        posted = await repair._execute_pcb_case2_row(row, posting_mode="post")
        self.assertTrue(posted.posted)
        self.assertEqual(posted.move_id, draft.move_id)
        self.assertEqual(rpc.svls[9]["account_move_id"][0], posted.move_id)
        self.assertEqual((rpc.svls[9]["value"], rpc.svls[9]["quantity"]), (100, 2))
        self.assertEqual([line["quantity"] for line in rpc.moves[posted.move_id]["line_ids"]], [2, 2])
        retry = await repair._execute_pcb_case2_row(row, posting_mode="post")
        self.assertEqual(retry.move_id, posted.move_id)
        self.assertEqual(sum(model == "account.move" for model, _ in rpc.create_calls), 1)
        self.assertFalse(any(model == "stock.valuation.layer" for model, _ in rpc.create_calls))

    async def test_source_drift_or_existing_journal_prevents_create(self):
        for changed in ("bill", "po_bill", "bill_origin", "stock_je", "other_je", "sibling", "input", "valuation", "company", "nan"):
            with self.subTest(changed=changed):
                row, rpc, repair = self.setup_row()
                row.review_confirmed = True
                if changed == "bill": rpc.purchase["invoice_lines"] = [10]
                elif changed == "po_bill": rpc.order["invoice_ids"] = [99]
                elif changed == "bill_origin": rpc.origin_bill = True
                elif changed == "stock_je": rpc.stock["account_move_ids"] = [99]
                elif changed == "other_je": rpc.other_journal = True
                elif changed == "sibling": rpc.sibling_link = True
                elif changed == "input": rpc.category["property_stock_account_input_categ_id"] = [202, "Expense"]
                elif changed == "valuation": rpc.category["property_valuation"] = "manual_periodic"
                elif changed == "company": rpc.stock["company_id"] = [2, "Other"]
                elif changed == "nan": rpc.svls[9]["value"] = float("nan")
                with self.assertRaises(Exception):
                    await repair._execute_pcb_case2_row(row, posting_mode="draft")
                self.assertEqual(rpc.create_calls, [])


if __name__ == "__main__":
    unittest.main()
