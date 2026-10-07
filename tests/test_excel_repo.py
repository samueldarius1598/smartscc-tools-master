import tempfile
import unittest
from unittest import mock
from datetime import datetime
from pathlib import Path
import re
import zipfile

from openpyxl import Workbook

import smartscc_tools.features.item_journal.workbook as workbook_module

from smartscc_tools.features.item_journal.workbook import (
    AUDIT_SHEET_NAME,
    COL_COMPANY_ID,
    COL_DATE_DONE,
    COL_DEST_LOC,
    COL_OP_TYPE,
    COL_PROD_ID,
    COL_PROD_KEY,
    COL_QTY,
    COL_RESULT,
    COL_SRC_LOC,
    COL_STJ,
    COL_UOM,
    EXCEL_OPEN_PATH_LIMIT,
    ItemJournalWorkbookRepo,
    build_summary_copy_name_stem,
    resolve_copy_save_target,
)


class ExcelRepoTest(unittest.TestCase):
    def test_should_keep_vba_depends_on_extension(self) -> None:
        self.assertTrue(ItemJournalWorkbookRepo._should_keep_vba(Path("sample.xlsm")))
        self.assertTrue(ItemJournalWorkbookRepo._should_keep_vba(Path("sample.xltm")))
        self.assertFalse(ItemJournalWorkbookRepo._should_keep_vba(Path("sample.xlsx")))
        self.assertFalse(ItemJournalWorkbookRepo._should_keep_vba(Path("sample.xltx")))

    def test_build_summary_copy_name_stem(self) -> None:
        stem = build_summary_copy_name_stem(
            company_name='CECILIA:Outlet/Bar*',
            processed_rows=321,
            when=datetime(2026, 3, 8, 5, 31, 42),
        )
        self.assertEqual(stem, "CECILIA_Outlet_Bar_ - 08-03-26 05.31.42 - 321")

    def test_load_workbook_disables_external_links(self) -> None:
        path = Path(r"C:\Temp\sample.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Item Journal"
        ws[f"{COL_PROD_ID}1"] = "Product ID"
        ws[f"{COL_PROD_KEY}1"] = "Product key (kode/nama)"

        with mock.patch.object(workbook_module, "load_workbook", return_value=wb) as mocked_load:
            repo = ItemJournalWorkbookRepo(str(path))
            self.assertIsNotNone(repo.sheet)
            mocked_load.assert_called_once_with(
                filename=str(path),
                keep_vba=False,
                keep_links=False,
            )

    def test_inplace_save_xlsx_keeps_non_macro_content(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "plain.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Item Journal"
            ws[f"{COL_PROD_ID}1"] = "Product ID"
            ws[f"{COL_PROD_KEY}1"] = "Product key (kode/nama)"
            ws[f"{COL_DATE_DONE}2"] = "2026-03-01"
            ws[f"{COL_COMPANY_ID}2"] = 10
            ws[f"{COL_OP_TYPE}2"] = "Internal Transfer"
            ws[f"{COL_SRC_LOC}2"] = "WH/Stock"
            ws[f"{COL_DEST_LOC}2"] = "WH/Output"
            ws[f"{COL_PROD_ID}2"] = 1001
            ws[f"{COL_QTY}2"] = 2
            ws[f"{COL_UOM}2"] = "Units"
            wb.save(path)

            repo = ItemJournalWorkbookRepo(str(path))
            rows = repo.read_rows()
            repo.write_rows(rows)
            save_result = repo.save("in-place")
            self.assertEqual(save_result.path, path)
            repo.workbook.close()

            with zipfile.ZipFile(path, "r") as zf:
                content_types = zf.read("[Content_Types].xml").decode("utf-8")
                workbook_rels = zf.read("xl/_rels/workbook.xml.rels").decode("utf-8")
            self.assertIn("spreadsheetml.sheet.main+xml", content_types)
            self.assertNotIn("macroEnabled.main+xml", content_types)
            self.assertNotIn("vbaProject", workbook_rels)

    def test_read_rows_and_active_logic(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "test.xlsm"
            wb = Workbook()
            ws = wb.active
            ws.title = "Item Journal"
            ws[f"{COL_PROD_ID}1"] = "Product ID"
            ws[f"{COL_PROD_KEY}1"] = "Product key (kode/nama)"

            ws[f"{COL_DATE_DONE}2"] = "2026-03-01"
            ws[f"{COL_COMPANY_ID}2"] = 10
            ws[f"{COL_OP_TYPE}2"] = "Internal Transfer"
            ws[f"{COL_SRC_LOC}2"] = "WH/Stock"
            ws[f"{COL_DEST_LOC}2"] = "WH/Output"
            ws[f"{COL_PROD_ID}2"] = 1001
            ws[f"{COL_QTY}2"] = 2
            ws[f"{COL_UOM}2"] = "Units"

            ws[f"{COL_DATE_DONE}3"] = "2026-03-01"
            ws[f"{COL_COMPANY_ID}3"] = 10
            ws[f"{COL_PROD_KEY}3"] = "SKU-1"
            ws[f"{COL_QTY}3"] = 1
            ws[f"{COL_RESULT}3"] = "DONE/001"
            ws[f"{COL_STJ}3"] = ""

            wb.save(path)

            repo = ItemJournalWorkbookRepo(str(path))
            rows = repo.read_rows()
            self.assertEqual(len(rows), 2)
            self.assertTrue(rows[0].is_active())
            self.assertFalse(rows[1].is_active())
            repo.workbook.close()

    def test_mark_rows_stopped_and_append_audit(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "test_stop.xlsm"
            wb = Workbook()
            ws = wb.active
            ws.title = "Item Journal"
            ws[f"{COL_PROD_ID}1"] = "Product ID"
            ws[f"{COL_PROD_KEY}1"] = "Product key (kode/nama)"

            ws[f"{COL_DATE_DONE}2"] = "2026-03-01"
            ws[f"{COL_COMPANY_ID}2"] = 10
            ws[f"{COL_OP_TYPE}2"] = "Internal Transfer"
            ws[f"{COL_SRC_LOC}2"] = "WH/Stock"
            ws[f"{COL_DEST_LOC}2"] = "WH/Output"
            ws[f"{COL_PROD_ID}2"] = 1001
            ws[f"{COL_QTY}2"] = 2
            ws[f"{COL_UOM}2"] = "Units"
            wb.save(path)

            repo = ItemJournalWorkbookRepo(str(path))
            rows = repo.read_rows()
            repo.mark_rows_stopped(rows, "Stop test")
            repo.write_rows(rows)
            repo.append_audit_events(
                [
                    {
                        "run_id": "r1",
                        "timestamp": "2026-03-04 10:00:00",
                        "row_number": 2,
                        "group_key": "g1",
                        "batch_seq": 1,
                        "stage": "ROW_STOPPED",
                        "model": "stock.picking",
                        "method": "create",
                        "attempt": 0,
                        "http_status": 0,
                        "duration_ms": 0,
                        "result": "STOPPED",
                        "message": "Stop test",
                        "picking_id": 0,
                    }
                ]
            )
            save_result = repo.save("in-place")
            self.assertEqual(save_result.mode_used, "in-place")
            self.assertFalse(save_result.used_temp_fallback)
            self.assertEqual(save_result.path, path)
            repo.workbook.close()

            repo2 = ItemJournalWorkbookRepo(str(path))
            rows2 = repo2.read_rows()
            self.assertEqual(rows2[0].result, "[Stopped]")
            self.assertIn("Stop test", rows2[0].error)
            self.assertIn(AUDIT_SHEET_NAME, repo2.workbook.sheetnames)
            repo2.workbook.close()

    def test_resolve_copy_save_target_normal_path(self) -> None:
        source = Path(r"C:\Work\item_journal.xlsm")
        ts = "20260305_010646"
        result = resolve_copy_save_target(source_path=source, timestamp=ts, excel_limit=EXCEL_OPEN_PATH_LIMIT)
        self.assertFalse(result.used_temp_fallback)
        self.assertEqual(result.warning, "")
        self.assertTrue(str(result.path).endswith(f"_processed_{ts}.xlsm"))
        self.assertLessEqual(len(str(result.path)), EXCEL_OPEN_PATH_LIMIT)

    def test_resolve_copy_save_target_name_override(self) -> None:
        source = Path(r"C:\Work\item_journal.xlsm")
        ts = "20260305_010646"
        result = resolve_copy_save_target(
            source_path=source,
            timestamp=ts,
            excel_limit=EXCEL_OPEN_PATH_LIMIT,
            name_stem_override="CECILIA - 08-03-26 05.31.42 - 321",
        )
        self.assertFalse(result.used_temp_fallback)
        self.assertEqual(result.warning, "")
        self.assertEqual(result.path.name, "CECILIA - 08-03-26 05.31.42 - 321.xlsm")
        self.assertLessEqual(len(str(result.path)), EXCEL_OPEN_PATH_LIMIT)

    def test_resolve_copy_save_target_shorten_name(self) -> None:
        parent = Path("C:/") / ("A" * 160)
        source = parent / ("Internal Transfer Cecillia Wijaya Food - Feb 2026.xlsx")
        ts = "20260305_010646"
        result = resolve_copy_save_target(source_path=source, timestamp=ts, excel_limit=EXCEL_OPEN_PATH_LIMIT)
        self.assertFalse(result.used_temp_fallback)
        self.assertIn("dipendekkan otomatis", result.warning)
        self.assertLessEqual(len(str(result.path)), EXCEL_OPEN_PATH_LIMIT)
        self.assertRegex(result.path.name, re.compile(r"_p_20260305_010646_[0-9a-f]{8}\.xlsx$", re.IGNORECASE))

    def test_resolve_copy_save_target_fallback_temp(self) -> None:
        parent = Path("C:/") / ("B" * 230)
        source = parent / "ij.xlsx"
        ts = "20260305_010646"
        temp_dir = Path(r"C:\Temp\ij")
        result = resolve_copy_save_target(
            source_path=source,
            timestamp=ts,
            excel_limit=EXCEL_OPEN_PATH_LIMIT,
            temp_dir=temp_dir,
        )
        self.assertTrue(result.used_temp_fallback)
        self.assertIn("disimpan ke folder temp", result.warning)
        self.assertEqual(result.path.parent, temp_dir)
        self.assertLessEqual(len(str(result.path)), EXCEL_OPEN_PATH_LIMIT)


if __name__ == "__main__":
    unittest.main()
