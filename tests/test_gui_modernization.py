import json
import re
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

from openpyxl import Workbook

import smartscc_tools.features.item_journal.entrypoints.gui as gui_module
from smartscc_tools.features.item_journal.entrypoints.gui import (
    AUTO_GROUP_BATCH_LIMIT,
    ERROR_COLOR,
    PersistentState,
    SUCCESS_COLOR,
    WARNING_COLOR,
    advance_slow_wait_preview,
    build_action_progress_layout,
    build_business_overrides,
    build_connection_mode_section_summary,
    build_grouping_section_summary,
    build_lock_dashboard_state,
    build_lock_dashboard_section_summary,
    build_output_dir_section_summary,
    build_progress_snapshot,
    build_slow_wait_preview_state,
    build_workbook_section_summary,
    default_section_open,
    render_slow_wait_preview,
)
from smartscc_tools.services.odoo.profiles import FOLLOW_GLOBAL_PROFILE_ID
from smartscc_tools.features.item_journal.workbook import COL_PROD_ID, COL_PROD_KEY, EXCEL_OPEN_PATH_LIMIT, ItemJournalWorkbookRepo


class _FakeTextWidget:
    def __init__(self) -> None:
        self.configure_calls: list[dict[str, object]] = []
        self.insert_calls: list[tuple[str, str]] = []
        self.see_calls: list[str] = []

    def configure(self, **kwargs) -> None:
        self.configure_calls.append(kwargs)

    def insert(self, index: str, text: str) -> None:
        self.insert_calls.append((index, text))

    def index(self, index: str) -> str:
        return "1.0"

    def delete(self, start: str, end: str) -> None:
        pass

    def see(self, index: str) -> None:
        self.see_calls.append(index)


class _FakeVar:
    def __init__(self, value: str = "") -> None:
        self.value = value

    def set(self, value: str) -> None:
        self.value = value

    def get(self) -> str:
        return self.value


class GuiModernizationTest(unittest.TestCase):
    def test_persistent_state_load_falls_back_when_output_dir_missing(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            state_path = Path(tmp_dir) / "gui_state.json"
            fallback_dir = Path(tmp_dir) / "Documents"
            fallback_dir.mkdir()
            state_path.write_text(
                json.dumps(
                    {
                        "workbook_path": r"C:\temp\sample.xlsm",
                        "output_dir": r"C:\does-not-exist\missing",
                        "db_override": "demo-db",
                        "grouping_mode": "custom",
                        "custom_limit": 220,
                        "section_open": {"logs": True, "actions": False},
                    }
                ),
                encoding="utf-8",
            )

            with mock.patch.object(gui_module, "default_output_dir", return_value=fallback_dir):
                state = PersistentState(path=state_path).load()

            self.assertEqual(state.output_dir, str(fallback_dir))
            self.assertEqual(state.database_profile_id, "demo-db")
            self.assertEqual(state.grouping_mode, "custom")
            self.assertEqual(state.custom_limit, 220)
            self.assertTrue(state.section_open["logs"])
            self.assertFalse(state.section_open["actions"])
            self.assertFalse(state.section_open["workbook"])

    def test_persistent_state_load_survives_corrupt_json(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            state_path = Path(tmp_dir) / "gui_state.json"
            state_path.write_text("{invalid", encoding="utf-8")
            state = PersistentState(path=state_path).load()
            self.assertEqual(state.grouping_mode, "auto")
            self.assertGreaterEqual(state.custom_limit, 1)
            self.assertTrue(state.output_dir)
            self.assertEqual(state.section_open, default_section_open())

    def test_persistent_state_save_persists_section_state(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            state_path = Path(tmp_dir) / "gui_state.json"
            store = PersistentState(path=state_path)
            state = gui_module.GuiState(section_open={"actions": True, "logs": True})

            with mock.patch.object(gui_module, "default_output_dir", return_value=Path(tmp_dir)):
                store.save(state)
                saved = json.loads(state_path.read_text(encoding="utf-8"))

            self.assertTrue(saved["section_open"]["actions"])
            self.assertTrue(saved["section_open"]["logs"])
            self.assertFalse(saved["section_open"]["workbook"])

    def test_build_business_overrides_auto_uses_sentinel(self) -> None:
        overrides = build_business_overrides("auto", 150, 2, 500, 600000)
        self.assertIn(f"TRANSACTION_VOLUME_PER_BATCH={AUTO_GROUP_BATCH_LIMIT}", overrides)

    def test_build_business_overrides_custom_uses_limit(self) -> None:
        overrides = build_business_overrides("custom", 175, 2, 500, 600000)
        self.assertIn("TRANSACTION_VOLUME_PER_BATCH=175", overrides)

    def test_build_lock_dashboard_state_green_when_all_dates_after_lock(self) -> None:
        state = build_lock_dashboard_state(
            unique_dates=[date(2026, 3, 10), date(2026, 3, 11)],
            lock_date=date(2026, 3, 9),
        )
        self.assertEqual(state.status_text, "Aman / Transfer Bisa Dijalankan")
        self.assertEqual(state.status_color, SUCCESS_COLOR)

    def test_build_lock_dashboard_state_red_when_any_date_on_or_before_lock(self) -> None:
        state = build_lock_dashboard_state(
            unique_dates=[date(2026, 3, 8), date(2026, 3, 10)],
            lock_date=date(2026, 3, 8),
        )
        self.assertEqual(state.status_text, "Minta Accounting Buka Lock Period Terlebih Dahulu")
        self.assertEqual(state.status_color, ERROR_COLOR)

    def test_build_lock_dashboard_state_warning_when_lock_missing(self) -> None:
        state = build_lock_dashboard_state(unique_dates=[date(2026, 3, 10)], lock_date=None)
        self.assertEqual(state.status_color, WARNING_COLOR)
        self.assertIn("Lock date kosong", state.status_text)

    def test_build_progress_snapshot_formats_real_progress(self) -> None:
        snapshot = build_progress_snapshot(current=45, total=100)
        self.assertEqual(snapshot.value, 45.0)
        self.assertEqual(snapshot.percent_text, "45%")
        self.assertEqual(snapshot.count_text, "45/100 item")

    def test_action_progress_layout_keeps_summary_visible_when_collapsed(self) -> None:
        expanded = build_action_progress_layout(True)
        collapsed = build_action_progress_layout(False)

        self.assertTrue(expanded.show_detail)
        self.assertTrue(expanded.show_buttons)
        self.assertFalse(collapsed.show_detail)
        self.assertFalse(collapsed.show_buttons)
        self.assertTrue(collapsed.show_compact_progress)

    def test_section_summary_helpers_format_collapsed_headers(self) -> None:
        self.assertEqual(
            build_workbook_section_summary(r"C:\Users\User\Documents\File Upload Cecillia Wijaya 2.xlsx"),
            "Documents - File Upload Cecillia Wijaya 2.xlsx",
        )
        self.assertEqual(
            build_output_dir_section_summary(r"C:\Users\User\2026 HPP\2026.02 HPP Februari"),
            "2026.02 HPP Februari",
        )
        self.assertEqual(
            build_connection_mode_section_summary("hwg-live", False),
            "hwg-live - No Dry Run",
        )
        self.assertEqual(
            build_connection_mode_section_summary("", True),
            "Use GAS Default - Dry Run",
        )
        self.assertEqual(
            build_grouping_section_summary("auto", 150, 1, 500, 600000),
            "Sesuai COA (Auto) - No Advanced Setting",
        )
        self.assertEqual(
            build_grouping_section_summary("custom", 175, 3, 1000, 120000),
            "Custom Limit 175 - Retry 3x, Delay 1000 ms, Journal Wait 120000 ms",
        )
        self.assertEqual(
            build_lock_dashboard_section_summary("Aman / Transfer Bisa Dijalankan"),
            "Aman / Transfer Bisa Dijalankan",
        )

    def test_append_log_updates_latest_preview_to_single_line(self) -> None:
        panel = gui_module.ItemJournalPanel.__new__(gui_module.ItemJournalPanel)
        panel.log_text = _FakeTextWidget()
        panel.latest_log_line_var = _FakeVar("Belum ada log.")

        gui_module.ItemJournalPanel._append_log_batch(
            panel,
            ["OK: Upload selesai untuk batch yang sangat panjang dengan detail tambahan\nbaris kedua diabaikan"],
        )

        self.assertEqual(
            panel.latest_log_line_var.get(),
            "OK: Upload selesai untuk batch yang sangat panjang dengan detail tambahan baris kedua diabaikan",
        )
        self.assertEqual(panel.log_text.insert_calls, [("end", "OK: Upload selesai untuk batch yang sangat panjang dengan detail tambahan\nbaris kedua diabaikan\n")])

    def test_slow_wait_preview_helper_loops_and_renders(self) -> None:
        state = build_slow_wait_preview_state(
            {
                "company_name": "CECILIA",
                "src": "WH/Stock",
                "dest": "WH/Output",
                "item_count": 3,
                "preview_items": ["A x1", "B x2", "C x3"],
            }
        )
        assert state is not None
        self.assertIn("Item 1/3", render_slow_wait_preview(state))

        state = advance_slow_wait_preview(state)
        self.assertEqual(state.current_index, 1)
        self.assertEqual(state.shown_once_count, 2)
        self.assertIn("Item 2/3", render_slow_wait_preview(state))

        state = advance_slow_wait_preview(state)
        state = advance_slow_wait_preview(state)
        self.assertEqual(state.current_index, 0)
        self.assertEqual(state.loop_count, 1)
        self.assertEqual(state.shown_once_count, 3)
        self.assertIn("loop 2", render_slow_wait_preview(state))

    def test_slow_wait_preview_helper_resets_on_new_fingerprint(self) -> None:
        state_a = build_slow_wait_preview_state(
            {
                "company_name": "CECILIA",
                "src": "WH/Stock",
                "dest": "WH/Output",
                "item_count": 2,
                "preview_items": ["A x1", "B x2"],
            }
        )
        state_b = build_slow_wait_preview_state(
            {
                "company_name": "CECILIA",
                "src": "WH/Stock",
                "dest": "WH/Transit",
                "item_count": 2,
                "preview_items": ["A x1", "B x2"],
            }
        )
        assert state_a is not None
        assert state_b is not None
        self.assertNotEqual(state_a.fingerprint, state_b.fingerprint)
        self.assertEqual(state_b.current_index, 0)

    def test_repo_save_copy_uses_output_dir(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "plain.xlsx"
            output_dir = Path(tmp_dir) / "results"

            wb = Workbook()
            ws = wb.active
            ws.title = "Item Journal"
            ws[f"{COL_PROD_ID}1"] = "Product ID"
            ws[f"{COL_PROD_KEY}1"] = "Product key (kode/nama)"
            wb.save(source_path)

            repo = ItemJournalWorkbookRepo(str(source_path))
            save_result = repo.save(
                mode="copy",
                copy_name_stem="CECILIA - 08-03-26 05.31.42 - 321",
                output_dir=output_dir,
            )
            repo.workbook.close()

            self.assertEqual(save_result.path.parent, output_dir)
            self.assertEqual(save_result.path.name, "CECILIA - 08-03-26 05.31.42 - 321.xlsx")
            self.assertTrue(save_result.path.exists())

    def test_repo_save_copy_keeps_excel_guard_under_output_dir(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            source_path = Path(tmp_dir) / "plain.xlsx"
            long_output_dir = Path(tmp_dir) / ("A" * 180)

            wb = Workbook()
            ws = wb.active
            ws.title = "Item Journal"
            ws[f"{COL_PROD_ID}1"] = "Product ID"
            ws[f"{COL_PROD_KEY}1"] = "Product key (kode/nama)"
            wb.save(source_path)

            repo = ItemJournalWorkbookRepo(str(source_path))
            save_result = repo.save(
                mode="copy",
                copy_name_stem="Internal Transfer Cecillia Wijaya Food - Feb 2026",
                output_dir=long_output_dir,
            )
            repo.workbook.close()

            self.assertLessEqual(len(str(save_result.path)), EXCEL_OPEN_PATH_LIMIT)
            self.assertRegex(save_result.path.name, re.compile(r"(\.xlsx|_[0-9a-f]{8}\.xlsx)$", re.IGNORECASE))


if __name__ == "__main__":
    unittest.main()
