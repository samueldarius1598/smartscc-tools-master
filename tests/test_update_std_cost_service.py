import asyncio
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

from smartscc_tools.features.update_std_cost.models import UpdateStdCostRunRequest
from smartscc_tools.features.update_std_cost.service import UpdateStdCostServiceAsync


def _make_workbook(path: Path, rows: list[dict[str, object]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Preparation Cost"
    ws["A1"] = "note-1"
    ws["A2"] = "Area"
    ws["B2"] = "Product Code"
    ws["C2"] = "Product Name"
    ws["H2"] = "Value per Unit"
    for index, row in enumerate(rows, start=3):
        ws.cell(row=index, column=1).value = row.get("area")
        ws.cell(row=index, column=2).value = row.get("code")
        ws.cell(row=index, column=3).value = row.get("name")
        ws.cell(row=index, column=8).value = row.get("price")
    cfg = wb.create_sheet("konfigurasi-ClientSide")
    cfg["H4"] = 1
    cfg["I4"] = "Company 1"
    cfg["H5"] = 2
    cfg["I5"] = "Company 2"
    wb.save(path)
    wb.close()


class _FakeRpc:
    def __init__(self) -> None:
        self.write_calls: list[tuple[str, tuple[int, ...], float, int]] = []
        self.stop_after_first_write = False
        self.stop_callback = None
        self.variant_lookup_delay = 0.0
        self.active_variant_lookups = 0
        self.max_active_variant_lookups = 0

    async def search_read(self, model, domain, fields=None, context=None, stage=""):  # noqa: ANN001, ANN201
        _ = (stage,)
        if model != "product.product":
            return []
        code_values = domain[0][2]
        if fields == ["default_code", "product_tmpl_id"]:
            rows = []
            for code in code_values:
                if code == "DUP":
                    rows.append({"default_code": "DUP", "product_tmpl_id": [10, "A"]})
                    rows.append({"default_code": "DUP", "product_tmpl_id": [20, "B"]})
                elif code != "MISS":
                    rows.append({"default_code": code, "product_tmpl_id": [100 + len(code), code]})
            return rows

        company_id = int((context or {}).get("company_id") or 0)
        if self.variant_lookup_delay:
            self.active_variant_lookups += 1
            self.max_active_variant_lookups = max(self.max_active_variant_lookups, self.active_variant_lookups)
            try:
                await asyncio.sleep(self.variant_lookup_delay)
            finally:
                self.active_variant_lookups -= 1
        rows = []
        for code in code_values:
            if code == "MISS":
                continue
            if company_id == 2 and code == "B2":
                continue
            rows.append({"id": company_id * 1000 + len(code), "default_code": code})
        return rows

    async def write(self, model, ids, values, context=None, stage=""):  # noqa: ANN001, ANN201
        company_id = int((context or {}).get("company_id") or 0)
        self.write_calls.append((model, tuple(ids), float(values["standard_price"]), company_id))
        if self.stop_after_first_write and len(self.write_calls) == 1 and self.stop_callback is not None:
            self.stop_callback()
        return True


class UpdateStdCostServiceTest(unittest.TestCase):
    def test_validate_template_writes_status_and_log_without_rpc_write(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "std_cost.xlsm"
            _make_workbook(
                workbook_path,
                [
                    {"area": "Area1", "code": "A1", "name": "Prod A1", "price": 100},
                    {"area": "Unknown", "code": "A2", "name": "Prod A2", "price": 200},
                ],
            )
            rpc = _FakeRpc()
            service = UpdateStdCostServiceAsync(rpc=rpc, logger=mock.Mock())

            with mock.patch("smartscc_tools.features.update_std_cost.service.fetch_company_area_map", return_value={"1": "Area1", "2": "Area1"}):
                summary = asyncio.run(
                    service.validate(
                        UpdateStdCostRunRequest(
                            workbook_path=str(workbook_path),
                            database="hwgroup_erp",
                            mode="template",
                        )
                    )
                )

            wb = load_workbook(workbook_path)
            ws = wb["Preparation Cost"]
            self.assertEqual(ws["J3"].value, "Yes")
            self.assertEqual(ws["J4"].value, "No")
            self.assertIn("LOG", wb.sheetnames)
            self.assertEqual(summary.updated_rows, 1)
            self.assertEqual(summary.failed_rows, 1)
            self.assertEqual(rpc.write_calls, [])
            wb.close()

    def test_execute_variant_writes_company_scoped_standard_price(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "std_cost.xlsm"
            _make_workbook(
                workbook_path,
                [
                    {"area": "Area1", "code": "A1", "name": "Prod A1", "price": 100},
                    {"area": "Area1", "code": "B2", "name": "Prod B2", "price": 200},
                ],
            )
            rpc = _FakeRpc()
            service = UpdateStdCostServiceAsync(rpc=rpc, logger=mock.Mock())

            with mock.patch("smartscc_tools.features.update_std_cost.service.fetch_company_area_map", return_value={"1": "Area1", "2": "Area1"}):
                summary = asyncio.run(
                    service.execute(
                        UpdateStdCostRunRequest(
                            workbook_path=str(workbook_path),
                            database="hwgroup_erp",
                            mode="variant",
                            dry_run=False,
                        )
                    )
                )

            self.assertEqual(summary.updated_rows, 2)
            self.assertEqual(summary.failed_rows, 0)
            self.assertEqual(summary.mode_summaries[0].missing_products, 1)
            self.assertTrue(any(call[0] == "product.product" and call[3] == 1 for call in rpc.write_calls))

    def test_execute_variant_fetches_company_product_maps_in_parallel(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "std_cost.xlsm"
            _make_workbook(
                workbook_path,
                [
                    {"area": "Area1", "code": "A1", "name": "Prod A1", "price": 100},
                    {"area": "Area1", "code": "B2", "name": "Prod B2", "price": 200},
                ],
            )
            rpc = _FakeRpc()
            rpc.variant_lookup_delay = 0.01
            service = UpdateStdCostServiceAsync(rpc=rpc, logger=mock.Mock())

            with mock.patch("smartscc_tools.features.update_std_cost.service.fetch_company_area_map", return_value={"1": "Area1", "2": "Area1"}):
                summary = asyncio.run(
                    service.execute(
                        UpdateStdCostRunRequest(
                            workbook_path=str(workbook_path),
                            database="hwgroup_erp",
                            mode="variant",
                            dry_run=False,
                        )
                    )
                )

            self.assertEqual(summary.updated_rows, 2)
            self.assertGreaterEqual(rpc.max_active_variant_lookups, 2)

    def test_execute_both_supports_cooperative_stop(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "std_cost.xlsm"
            _make_workbook(
                workbook_path,
                [
                    {"area": "Area1", "code": "A1", "name": "Prod A1", "price": 100},
                    {"area": "Area1", "code": "A2", "name": "Prod A2", "price": 120},
                ],
            )
            rpc = _FakeRpc()
            service = UpdateStdCostServiceAsync(rpc=rpc, logger=mock.Mock())
            rpc.stop_after_first_write = True
            rpc.stop_callback = service.request_stop

            with mock.patch("smartscc_tools.features.update_std_cost.service.fetch_company_area_map", return_value={"1": "Area1", "2": "Area1"}):
                summary = asyncio.run(
                    service.execute(
                        UpdateStdCostRunRequest(
                            workbook_path=str(workbook_path),
                            database="hwgroup_erp",
                            mode="both",
                            dry_run=False,
                        )
                    )
                )

            self.assertTrue(summary.stopped)
            self.assertTrue(summary.mode_summaries[0].stopped or summary.mode_summaries[-1].stopped)
            self.assertGreaterEqual(len(rpc.write_calls), 1)


if __name__ == "__main__":
    unittest.main()
