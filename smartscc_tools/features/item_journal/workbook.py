"""Excel repository for Item Journal workbook read/write."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
import hashlib
import os
from pathlib import Path
import tempfile
from typing import Any, Callable, Dict, Iterable, List

from smartscc_tools.features.item_journal.observability.audit import migrate_legacy_event
from openpyxl import load_workbook
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from smartscc_tools.features.item_journal.utils import excel_value_to_date, normalize_text, to_float, to_int


ITEM_JOURNAL_SHEET_NAME = "Item Journal"
AUDIT_SHEET_NAME = "Item Journal Audit"
STATUS_CELL = "N1"
FISCALYEAR_LOCK_DATE_CELL = "P1"
DATA_START_ROW = 2

COL_DATE_DONE = "A"
COL_COMPANY_ID = "B"
COL_COMPANY_NAME = "C"
COL_OP_TYPE = "D"
COL_SRC_LOC = "E"
COL_DEST_LOC = "F"
COL_PROD_ID = "G"
COL_PROD_KEY = "H"
COL_QTY = "J"
COL_UOM = "K"
COL_RESULT = "L"
COL_STJ = "M"
COL_ERROR = "N"
EXCEL_OPEN_PATH_LIMIT = 218

SHORTEN_COPY_WARNING = "Nama file copy dipendekkan otomatis karena batas path Excel."
TEMP_FALLBACK_WARNING_PREFIX = "Path folder asal terlalu panjang; file disimpan ke folder temp: "

AUDIT_V2_HEADERS = [
    "schema_version",
    "run_id",
    "event_id",
    "timestamp_utc",
    "event_type",
    "stage",
    "severity",
    "cause_code",
    "rpc_model",
    "rpc_method",
    "duration_ms",
    "batch_seq",
    "row_number",
    "group_key",
    "attempt",
    "http_status",
    "message",
    "picking_id",
]
def _normalize_stj_input(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return ""
    text = normalize_text(value)
    if text.lower() in {"", "false", "true"}:
        return ""
    return text


@dataclass
class ItemJournalRow:
    row_number: int
    date_done_raw: Any
    company_id: int
    company_name: str
    op_type: str
    src_loc: str
    dest_loc: str
    prod_id_text: str
    prod_key: str
    qty: float
    uom: str
    result: str
    stj: str
    error: str = ""

    def is_active(self) -> bool:
        if self.result.strip() or self.stj.strip():
            return False
        if (
            self.company_id <= 0
            and not self.op_type
            and not self.src_loc
            and not self.dest_loc
            and not self.prod_id_text
            and not self.prod_key
            and not normalize_text(self.date_done_raw)
        ):
            return False
        return True

    def parsed_date(self) -> date | None:
        return excel_value_to_date(self.date_done_raw)

    def append_error(self, message: str) -> None:
        text = normalize_text(message)
        if not text:
            return
        if self.error:
            self.error = f"{self.error} | {text}"
        else:
            self.error = text

    def mark_error(self, message: str) -> None:
        self.result = "[Error]"
        self.stj = ""
        self.append_error(message)


@dataclass
class SaveResult:
    path: Path
    mode_used: str
    warning: str = ""
    used_temp_fallback: bool = False


@dataclass
class CopyPathResolution:
    path: Path
    warning: str = ""
    used_temp_fallback: bool = False


def _path_len(path_value: Path) -> int:
    return len(str(path_value))


def _suffix_or_default(source_path: Path, suffix_override: str | None = None) -> str:
    clean = (suffix_override or "").strip()
    if clean:
        return clean if clean.startswith(".") else f".{clean}"
    source_suffix = (source_path.suffix or "").strip()
    if source_suffix:
        return source_suffix
    return ".xlsm"


def _sanitize_filename_component(value: str, fallback: str) -> str:
    text = normalize_text(value)
    if not text:
        return fallback
    cleaned = text
    for token in '<>:"/\\|?*':
        cleaned = cleaned.replace(token, "_")
    cleaned = " ".join(cleaned.replace("\r", " ").replace("\n", " ").split())
    cleaned = cleaned.rstrip(" .")
    return cleaned or fallback


def build_summary_copy_name_stem(
    company_name: str,
    processed_rows: int,
    when: datetime | None = None,
) -> str:
    safe_company = _sanitize_filename_component(company_name, fallback="Company")
    safe_company = safe_company[:64].rstrip(" .") or "Company"
    safe_rows = max(0, int(processed_rows or 0))
    stamp = (when or datetime.now()).strftime("%d-%m-%y %H.%M.%S")
    return f"{safe_company} - {stamp} - {safe_rows}"


def resolve_copy_save_target(
    source_path: Path,
    timestamp: str,
    excel_limit: int = EXCEL_OPEN_PATH_LIMIT,
    temp_dir: Path | None = None,
    suffix_override: str | None = None,
    name_stem_override: str | None = None,
    output_dir: Path | None = None,
) -> CopyPathResolution:
    target_parent = Path(output_dir) if output_dir is not None else source_path.parent
    suffix = _suffix_or_default(source_path=source_path, suffix_override=suffix_override)
    requested_stem = normalize_text(name_stem_override)
    if requested_stem:
        base_stem = _sanitize_filename_component(requested_stem, fallback=source_path.stem or "Workbook")
    else:
        base_stem = f"{source_path.stem}_processed_{timestamp}"

    original_name = f"{base_stem}{suffix}"
    original_candidate = target_parent / original_name
    if _path_len(original_candidate) <= excel_limit:
        return CopyPathResolution(path=original_candidate)

    hash8 = hashlib.sha1(f"{source_path}|{target_parent}|{timestamp}|{base_stem}".encode("utf-8")).hexdigest()[:8]
    compact_tail = f"_p_{timestamp}_{hash8}{suffix}"
    max_name_len = excel_limit - _path_len(target_parent) - 1
    available_stem_len = max_name_len - len(compact_tail)
    if available_stem_len >= 1:
        short_stem = base_stem[:available_stem_len]
        short_candidate = target_parent / f"{short_stem}{compact_tail}"
        if _path_len(short_candidate) <= excel_limit:
            return CopyPathResolution(path=short_candidate, warning=SHORTEN_COPY_WARNING)

    target_temp_dir = temp_dir or (Path(tempfile.gettempdir()) / "smartscc_tools")
    fallback_name = f"IJ_{timestamp}_{hash8}{suffix}"
    fallback_candidate = target_temp_dir / fallback_name
    if _path_len(fallback_candidate) > excel_limit:
        fallback_candidate = target_temp_dir / f"IJ_{hash8}{suffix}"
    warning = f"{TEMP_FALLBACK_WARNING_PREFIX}{fallback_candidate}"
    return CopyPathResolution(path=fallback_candidate, warning=warning, used_temp_fallback=True)


class ItemJournalWorkbookRepo:
    @staticmethod
    def _should_keep_vba(path: Path) -> bool:
        return path.suffix.lower() in {".xlsm", ".xltm", ".xlam"}

    def __init__(self, workbook_path: str) -> None:
        self.workbook_path = Path(workbook_path)
        self.workbook: Workbook = load_workbook(
            filename=str(self.workbook_path),
            keep_vba=self._should_keep_vba(self.workbook_path),
            keep_links=False,
        )
        self.sheet: Worksheet = self._get_item_journal_sheet()
        self.validate_layout_or_fail()

    def _get_item_journal_sheet(self) -> Worksheet:
        if ITEM_JOURNAL_SHEET_NAME not in self.workbook.sheetnames:
            raise RuntimeError(f"Sheet '{ITEM_JOURNAL_SHEET_NAME}' tidak ditemukan.")
        return self.workbook[ITEM_JOURNAL_SHEET_NAME]

    def validate_layout_or_fail(self) -> None:
        header_prod_id = normalize_text(self.sheet[f"{COL_PROD_ID}1"].value)
        header_prod_key = normalize_text(self.sheet[f"{COL_PROD_KEY}1"].value)
        if header_prod_id.lower() != "product id" or header_prod_key.lower() != "product key (kode/nama)":
            raise RuntimeError(
                "Layout Item Journal wajib versi baru: "
                "kolom G='Product ID' dan kolom H='Product key (kode/nama)'. "
                f"Header saat ini: G='{header_prod_id}', H='{header_prod_key}'."
            )

    def get_last_data_row(self) -> int:
        cols = [
            COL_COMPANY_ID,
            COL_COMPANY_NAME,
            COL_OP_TYPE,
            COL_SRC_LOC,
            COL_DEST_LOC,
            COL_PROD_ID,
            COL_PROD_KEY,
            COL_QTY,
            COL_UOM,
        ]
        max_row = 1
        for col in cols:
            row = self.sheet.max_row
            while row > 1:
                value = self.sheet[f"{col}{row}"].value
                if normalize_text(value):
                    break
                row -= 1
            if row > max_row:
                max_row = row
        return max_row

    def read_rows(self) -> List[ItemJournalRow]:
        last_row = self.get_last_data_row()
        if last_row < DATA_START_ROW:
            return []

        rows: List[ItemJournalRow] = []
        for row in range(DATA_START_ROW, last_row + 1):
            rows.append(
                ItemJournalRow(
                    row_number=row,
                    date_done_raw=self.sheet[f"{COL_DATE_DONE}{row}"].value,
                    company_id=to_int(self.sheet[f"{COL_COMPANY_ID}{row}"].value),
                    company_name=normalize_text(self.sheet[f"{COL_COMPANY_NAME}{row}"].value),
                    op_type=normalize_text(self.sheet[f"{COL_OP_TYPE}{row}"].value),
                    src_loc=normalize_text(self.sheet[f"{COL_SRC_LOC}{row}"].value),
                    dest_loc=normalize_text(self.sheet[f"{COL_DEST_LOC}{row}"].value),
                    prod_id_text=normalize_text(self.sheet[f"{COL_PROD_ID}{row}"].value),
                    prod_key=normalize_text(self.sheet[f"{COL_PROD_KEY}{row}"].value),
                    qty=to_float(self.sheet[f"{COL_QTY}{row}"].value),
                    uom=normalize_text(self.sheet[f"{COL_UOM}{row}"].value),
                    result=normalize_text(self.sheet[f"{COL_RESULT}{row}"].value),
                    stj=_normalize_stj_input(self.sheet[f"{COL_STJ}{row}"].value),
                    error=normalize_text(self.sheet[f"{COL_ERROR}{row}"].value),
                )
            )
        return rows

    def write_rows(self, rows: Iterable[ItemJournalRow]) -> None:
        for item in rows:
            row = item.row_number
            self.sheet[f"{COL_RESULT}{row}"] = item.result
            self.sheet[f"{COL_STJ}{row}"] = item.stj
            self.sheet[f"{COL_ERROR}{row}"] = item.error

    def write_status(self, text: str) -> None:
        self.sheet[STATUS_CELL] = text

    def mark_rows_stopped(self, rows: Iterable[ItemJournalRow], message: str) -> None:
        for item in rows:
            if item.result.strip():
                continue
            item.result = "[Stopped]"
            item.stj = ""
            item.append_error(message)

    def write_lock_date(self, value: date | None) -> None:
        cell = self.sheet[FISCALYEAR_LOCK_DATE_CELL]
        if value is None:
            cell.value = None
        else:
            cell.value = value
            cell.number_format = "yyyy-mm-dd"

    def _preferred_copy_suffix(self) -> str:
        mime_type = str(getattr(self.workbook, "mime_type", "") or "").lower()
        is_template = bool(getattr(self.workbook, "template", False))
        if "macroenabled" in mime_type:
            return ".xltm" if is_template else ".xlsm"
        if is_template:
            return ".xltx"
        source_suffix = (self.workbook_path.suffix or "").strip().lower()
        if source_suffix:
            return source_suffix
        return ".xlsm"

    def _atomic_save_to_path(self, target_path: Path) -> None:
        target_path.parent.mkdir(parents=True, exist_ok=True)
        fd, temp_name = tempfile.mkstemp(
            prefix=f".{target_path.stem}_",
            suffix=target_path.suffix or ".tmp",
            dir=str(target_path.parent),
        )
        os.close(fd)
        temp_path = Path(temp_name)
        try:
            self.workbook.save(str(temp_path))
            temp_path.replace(target_path)
        except Exception:
            try:
                if temp_path.exists():
                    temp_path.unlink()
            except Exception:
                pass
            raise

    def save(
        self,
        mode: str,
        ask_decision: Callable[[], str] | None = None,
        copy_name_stem: str | None = None,
        output_dir: str | Path | None = None,
    ) -> SaveResult:
        mode_clean = mode.strip().lower()
        target_path = self.workbook_path
        if mode_clean == "ask":
            if ask_decision is None:
                raise RuntimeError("Save mode 'ask' membutuhkan callback keputusan.")
            mode_clean = ask_decision().strip().lower()
            if mode_clean not in {"in-place", "copy"}:
                raise RuntimeError("Keputusan save mode tidak valid.")

        if mode_clean == "copy":
            target_output_dir = Path(output_dir).expanduser() if output_dir is not None and str(output_dir).strip() else None
            resolution = resolve_copy_save_target(
                source_path=self.workbook_path,
                timestamp=datetime.now().strftime("%Y%m%d_%H%M%S"),
                excel_limit=EXCEL_OPEN_PATH_LIMIT,
                suffix_override=self._preferred_copy_suffix(),
                name_stem_override=copy_name_stem,
                output_dir=target_output_dir,
            )
            target_path = resolution.path
            self._atomic_save_to_path(target_path)
            return SaveResult(
                path=target_path,
                mode_used="copy",
                warning=resolution.warning,
                used_temp_fallback=resolution.used_temp_fallback,
            )

        if mode_clean != "in-place":
            raise RuntimeError(f"Save mode tidak valid: {mode}")

        self._atomic_save_to_path(self.workbook_path)
        return SaveResult(path=self.workbook_path, mode_used="in-place")

    def append_audit_events(self, events: Iterable[Dict[str, Any]]) -> None:
        ws = self._get_or_create_audit_sheet()
        self._ensure_audit_v2_header(ws)
        for event in events:
            migrated = migrate_legacy_event(dict(event))
            ws.append(
                [
                    migrated.get("schema_version", ""),
                    migrated.get("run_id", ""),
                    migrated.get("event_id", ""),
                    migrated.get("timestamp_utc", ""),
                    migrated.get("event_type", ""),
                    migrated.get("stage", ""),
                    migrated.get("severity", ""),
                    migrated.get("cause_code", ""),
                    migrated.get("rpc_model", ""),
                    migrated.get("rpc_method", ""),
                    migrated.get("duration_ms", 0),
                    migrated.get("batch_seq", 0),
                    migrated.get("row_number", 0),
                    migrated.get("group_key", ""),
                    migrated.get("attempt", 0),
                    migrated.get("http_status", 0),
                    migrated.get("message", ""),
                    migrated.get("picking_id", 0),
                ]
            )

    def append_perf_summary(self, summary: Dict[str, Any]) -> None:
        ws = self._get_or_create_audit_sheet()
        self._ensure_audit_v2_header(ws)
        ws.append([])
        ws.append(["PERF_SUMMARY_V2", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""])
        ws.append(["run_id", summary.get("run_id", ""), "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""])
        ws.append(["profile", summary.get("profile", ""), "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""])
        ws.append(["event_count", summary.get("event_count", 0), "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""])
        ws.append(["threshold_breach_count", summary.get("threshold_breach_count", 0), "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""])
        ws.append(["top_causes", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""])
        for cause in summary.get("top_causes", [])[:10]:
            ws.append(
                [
                    "cause",
                    cause.get("cause_code", ""),
                    "count",
                    cause.get("count", 0),
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                ]
            )
        ws.append(["top_rpc_by_duration", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""])
        for rpc_item in summary.get("top_rpc_by_duration", [])[:10]:
            ws.append(
                [
                    "rpc",
                    rpc_item.get("rpc", ""),
                    "total_ms",
                    rpc_item.get("total_duration_ms", 0),
                    "avg_ms",
                    rpc_item.get("avg_duration_ms", 0),
                    "count",
                    rpc_item.get("count", 0),
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                ]
            )

    def _get_or_create_audit_sheet(self) -> Worksheet:
        if AUDIT_SHEET_NAME in self.workbook.sheetnames:
            return self.workbook[AUDIT_SHEET_NAME]
        ws = self.workbook.create_sheet(AUDIT_SHEET_NAME)
        ws.append(AUDIT_V2_HEADERS)
        return ws

    def _ensure_audit_v2_header(self, ws: Worksheet) -> None:
        if ws.max_row < 1:
            ws.append(AUDIT_V2_HEADERS)
            return
        existing = [str(ws.cell(row=1, column=idx + 1).value or "").strip() for idx in range(len(AUDIT_V2_HEADERS))]
        if existing == AUDIT_V2_HEADERS:
            return
        ws.delete_rows(1, ws.max_row)
        ws.append(AUDIT_V2_HEADERS)


def build_copy_path(source_path: Path, output_dir: Path | None = None) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    resolution = resolve_copy_save_target(
        source_path=source_path,
        timestamp=timestamp,
        excel_limit=EXCEL_OPEN_PATH_LIMIT,
        output_dir=output_dir,
    )
    return resolution.path


def summarize_active_company_ids(rows: Iterable[ItemJournalRow]) -> List[int]:
    company_ids = sorted({item.company_id for item in rows if item.is_active() and item.company_id > 0})
    return company_ids


def summarize_unique_dates(rows: Iterable[ItemJournalRow]) -> List[date]:
    dates = sorted({d for item in rows if item.is_active() for d in [item.parsed_date()] if d is not None})
    return dates
