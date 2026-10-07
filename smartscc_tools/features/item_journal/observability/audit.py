"""Audit event model, sinks, and migration helpers."""

from __future__ import annotations

from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
import json
from pathlib import Path
import shutil
from threading import Lock
import uuid
from typing import Any, Dict, Iterable, List

from openpyxl import load_workbook

from smartscc_tools.features.item_journal.observability.perf import SCHEMA_VERSION, classify_cause
from smartscc_tools.features.item_journal.utils import normalize_text, utc_now_iso


MANDATORY_V2_KEYS = [
    "schema_version",
    "run_id",
    "event_id",
    "timestamp_utc",
    "event_type",
    "stage",
    "severity",
    "cause_code",
    "rpc_model",
    "rpc_method",
    "duration_ms",
    "batch_seq",
    "row_number",
]


def _severity_from_result(result: str) -> str:
    clean = normalize_text(result).upper()
    if clean in {"ERROR", "FAIL"}:
        return "error"
    if clean in {"RETRY", "SPLIT", "STOPPED", "WARNING", "WARN"}:
        return "warn"
    return "info"


def migrate_legacy_event(payload: Dict[str, Any], run_id_override: str = "") -> Dict[str, Any]:
    run_id = normalize_text(run_id_override) or normalize_text(payload.get("run_id"))
    if not run_id:
        run_id = datetime.now().strftime("%Y%m%d_%H%M%S")

    stage = normalize_text(payload.get("stage"))
    message = normalize_text(payload.get("message"))
    event_type = normalize_text(payload.get("event_type")).lower() or stage.lower() or "event"
    severity = normalize_text(payload.get("severity")).lower() or _severity_from_result(
        str(payload.get("result", ""))
    )
    duration_ms = int(payload.get("duration_ms", 0) or 0)
    rpc_model = normalize_text(payload.get("rpc_model") or payload.get("model"))
    rpc_method = normalize_text(payload.get("rpc_method") or payload.get("method"))
    cause_code = normalize_text(payload.get("cause_code")) or classify_cause(
        message=message,
        stage=stage,
        severity=severity,
        event_type=event_type,
        duration_ms=duration_ms,
    )
    timestamp_utc = normalize_text(payload.get("timestamp_utc"))
    if not timestamp_utc:
        timestamp_legacy = normalize_text(payload.get("timestamp"))
        if timestamp_legacy:
            try:
                dt = datetime.strptime(timestamp_legacy, "%Y-%m-%d %H:%M:%S")
                timestamp_utc = dt.replace(tzinfo=timezone.utc).isoformat().replace("+00:00", "Z")
            except ValueError:
                timestamp_utc = utc_now_iso()
        else:
            timestamp_utc = utc_now_iso()

    return {
        "schema_version": SCHEMA_VERSION,
        "run_id": run_id,
        "event_id": normalize_text(payload.get("event_id")) or str(uuid.uuid4()),
        "timestamp_utc": timestamp_utc,
        "event_type": event_type,
        "stage": stage,
        "severity": severity or "info",
        "cause_code": cause_code or "unknown",
        "rpc_model": rpc_model,
        "rpc_method": rpc_method,
        "duration_ms": duration_ms,
        "batch_seq": int(payload.get("batch_seq", 0) or 0),
        "row_number": int(payload.get("row_number", 0) or 0),
        "group_key": normalize_text(payload.get("group_key")),
        "attempt": int(payload.get("attempt", 0) or 0),
        "http_status": int(payload.get("http_status", 0) or 0),
        "message": message,
        "picking_id": int(payload.get("picking_id", 0) or 0),
        "details": payload.get("details") if isinstance(payload.get("details"), dict) else {},
    }


@dataclass
class AuditEvent:
    run_id: str
    event_type: str
    stage: str
    severity: str
    cause_code: str
    rpc_model: str
    rpc_method: str
    duration_ms: int
    batch_seq: int
    row_number: int
    message: str = ""
    group_key: str = ""
    attempt: int = 0
    http_status: int = 0
    picking_id: int = 0
    details: Dict[str, Any] = field(default_factory=dict)
    schema_version: str = SCHEMA_VERSION
    event_id: str = field(default_factory=lambda: str(uuid.uuid4()))
    timestamp_utc: str = field(default_factory=utc_now_iso)

    @classmethod
    def now(
        cls,
        run_id: str,
        row_number: int = 0,
        group_key: str = "",
        batch_seq: int = 0,
        stage: str = "",
        model: str = "",
        method: str = "",
        attempt: int = 0,
        http_status: int = 0,
        duration_ms: int = 0,
        result: str = "",
        message: str = "",
        picking_id: int = 0,
    ) -> "AuditEvent":
        event_type = normalize_text(stage).lower() or "event"
        severity = _severity_from_result(result)
        cause_code = classify_cause(
            message=message,
            stage=stage,
            severity=severity,
            event_type=event_type,
            duration_ms=int(duration_ms or 0),
        )
        return cls(
            run_id=run_id,
            event_type=event_type,
            stage=stage,
            severity=severity,
            cause_code=cause_code,
            rpc_model=model,
            rpc_method=method,
            duration_ms=int(duration_ms or 0),
            batch_seq=int(batch_seq or 0),
            row_number=int(row_number or 0),
            message=message,
            group_key=group_key,
            attempt=int(attempt or 0),
            http_status=int(http_status or 0),
            picking_id=int(picking_id or 0),
        )

    def to_dict(self) -> Dict[str, Any]:
        payload = asdict(self)
        for key in MANDATORY_V2_KEYS:
            if key not in payload:
                raise RuntimeError(f"Audit event missing mandatory key: {key}")
        return payload


class AuditSink:
    def __init__(
        self,
        run_id: str,
        audit_level: str = "item",
        audit_file: str | None = None,
        flush_every: int = 200,
        persist_file: bool = True,
    ) -> None:
        self.run_id = run_id
        self.audit_level = (audit_level or "item").strip().lower()
        self.flush_every = max(1, int(flush_every))
        self.lock = Lock()
        self.buffer: List[AuditEvent] = []
        self.closed = False
        self.persist_file = bool(persist_file or normalize_text(audit_file))
        self.path: Path | None = None
        self.fp = None

        if normalize_text(audit_file):
            self.path = Path(str(audit_file))
        elif self.persist_file:
            log_dir = Path("logs")
            log_dir.mkdir(parents=True, exist_ok=True)
            self.path = log_dir / f"audit_{run_id}.jsonl"
        if self.path is not None and self.persist_file:
            self.path.parent.mkdir(parents=True, exist_ok=True)
            self.fp = self.path.open("a", encoding="utf-8")

    def emit(self, event: AuditEvent) -> None:
        if self.closed:
            return
        if self.audit_level not in {"item", "summary"}:
            return
        if self.audit_level == "summary" and event.stage not in {
            "RUN_START",
            "RUN_DONE",
            "RUN_STOPPED",
            "PRECHECK",
            "BATCH_SUBMIT",
            "BATCH_ERROR",
        }:
            return

        with self.lock:
            if self.fp is not None:
                self.fp.write(json.dumps(event.to_dict(), ensure_ascii=False) + "\n")
                self.fp.flush()
            self.buffer.append(event)

    def pop_buffered_events(self) -> List[AuditEvent]:
        with self.lock:
            if not self.buffer:
                return []
            out = self.buffer[:]
            self.buffer.clear()
            return out

    def should_flush(self) -> bool:
        with self.lock:
            return len(self.buffer) >= self.flush_every

    def close(self) -> None:
        with self.lock:
            if self.closed:
                return
            self.closed = True
            if self.fp is not None:
                try:
                    self.fp.flush()
                finally:
                    self.fp.close()


def migrate_audit_jsonl(input_path: str, output_path: str) -> Dict[str, Any]:
    src = Path(input_path)
    dst = Path(output_path)
    if not src.exists():
        raise RuntimeError(f"Input audit file tidak ditemukan: {src}")
    dst.parent.mkdir(parents=True, exist_ok=True)

    total = 0
    migrated = 0
    with src.open("r", encoding="utf-8") as fin, dst.open("w", encoding="utf-8") as fout:
        for raw in fin:
            line = raw.strip()
            if not line:
                continue
            total += 1
            try:
                payload = json.loads(line)
            except Exception:
                continue
            out = migrate_legacy_event(payload)
            fout.write(json.dumps(out, ensure_ascii=False) + "\n")
            migrated += 1
    return {"input": str(src), "output": str(dst), "total": total, "migrated": migrated}


def _rows_from_sheet(ws) -> List[Dict[str, Any]]:  # noqa: ANN001
    if ws.max_row < 1:
        return []
    headers = [str(ws.cell(row=1, column=idx).value or "").strip() for idx in range(1, ws.max_column + 1)]
    rows: List[Dict[str, Any]] = []
    for row_idx in range(2, ws.max_row + 1):
        obj: Dict[str, Any] = {}
        non_empty = False
        for col_idx, key in enumerate(headers, start=1):
            if not key:
                continue
            value = ws.cell(row=row_idx, column=col_idx).value
            if value not in {None, ""}:
                non_empty = True
            obj[key] = value
        if non_empty:
            rows.append(obj)
    return rows


def _write_rows_to_sheet_v2(ws, events: Iterable[Dict[str, Any]]) -> int:  # noqa: ANN001
    from smartscc_tools.features.item_journal.workbook import AUDIT_V2_HEADERS

    ws.delete_rows(1, ws.max_row if ws.max_row > 0 else 1)
    ws.append(AUDIT_V2_HEADERS)
    count = 0
    for payload in events:
        ws.append(
            [
                payload.get("schema_version", ""),
                payload.get("run_id", ""),
                payload.get("event_id", ""),
                payload.get("timestamp_utc", ""),
                payload.get("event_type", ""),
                payload.get("stage", ""),
                payload.get("severity", ""),
                payload.get("cause_code", ""),
                payload.get("rpc_model", ""),
                payload.get("rpc_method", ""),
                payload.get("duration_ms", 0),
                payload.get("batch_seq", 0),
                payload.get("row_number", 0),
                payload.get("group_key", ""),
                payload.get("attempt", 0),
                payload.get("http_status", 0),
                payload.get("message", ""),
                payload.get("picking_id", 0),
            ]
        )
        count += 1
    return count


def migrate_audit_sheet(sheet_input: str, sheet_output: str) -> Dict[str, Any]:
    from smartscc_tools.features.item_journal.workbook import AUDIT_SHEET_NAME, AUDIT_V2_HEADERS

    src = Path(sheet_input)
    dst = Path(sheet_output)
    if not src.exists():
        raise RuntimeError(f"Workbook input tidak ditemukan: {src}")
    dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, dst)

    wb = load_workbook(filename=str(dst), keep_vba=True, keep_links=False)
    try:
        if AUDIT_SHEET_NAME not in wb.sheetnames:
            ws = wb.create_sheet(AUDIT_SHEET_NAME)
            ws.append(AUDIT_V2_HEADERS)
            wb.save(str(dst))
            return {"input": str(src), "output": str(dst), "total": 0, "migrated": 0}

        ws = wb[AUDIT_SHEET_NAME]
        rows = _rows_from_sheet(ws)
        migrated_rows = [migrate_legacy_event(row) for row in rows]
        migrated_count = _write_rows_to_sheet_v2(ws, migrated_rows)
        wb.save(str(dst))
        return {"input": str(src), "output": str(dst), "total": len(rows), "migrated": migrated_count}
    finally:
        wb.close()


__all__ = [
    "AuditEvent",
    "AuditSink",
    "MANDATORY_V2_KEYS",
    "migrate_audit_jsonl",
    "migrate_audit_sheet",
    "migrate_legacy_event",
]
