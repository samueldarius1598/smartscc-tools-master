"""Excel helpers for Update Standard Cost workbook flow."""

from __future__ import annotations

import os
from dataclasses import dataclass
from datetime import datetime
from typing import Any, List, Sequence, Tuple

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet


PREP_HEADER_ROW = 2
PREP_DATA_START_ROW = 3

LOG_HEADERS = [
    "Timestamp",
    "User",
    "Level",
    "Message",
    "Area",
    "Company ID",
    "Company Name",
    "Product Code",
    "Product Name",
    "Value per Unit",
    "Note",
]


@dataclass
class PrepRow:
    row_number: int
    area: Any
    product_code: Any
    product_name: Any
    value_per_unit: Any


def load_workbook_keep_vba(path: str) -> Workbook:
    return load_workbook(filename=path, keep_vba=True)


def get_required_sheet(workbook: Workbook, sheet_name: str) -> Worksheet:
    if sheet_name not in workbook.sheetnames:
        raise RuntimeError(f"Sheet '{sheet_name}' tidak ditemukan.")
    return workbook[sheet_name]


def get_or_create_sheet(workbook: Workbook, sheet_name: str) -> Worksheet:
    if sheet_name in workbook.sheetnames:
        return workbook[sheet_name]
    return workbook.create_sheet(sheet_name)


def determine_prep_last_row(ws: Worksheet) -> int:
    return max(last_row_by_col(ws, 1), last_row_by_col(ws, 2), last_row_by_col(ws, 8))


def last_row_by_col(ws: Worksheet, col_idx: int) -> int:
    max_row = max(ws.max_row, 1)
    for row in range(max_row, 0, -1):
        value = ws.cell(row=row, column=col_idx).value
        if value is None:
            continue
        if isinstance(value, str) and value.strip() == "":
            continue
        return row
    return 1


def read_preparation_rows(ws: Worksheet) -> Tuple[List[PrepRow], int]:
    last_row = determine_prep_last_row(ws)
    if last_row < PREP_DATA_START_ROW:
        return [], last_row

    rows: list[PrepRow] = []
    for row_number in range(PREP_DATA_START_ROW, last_row + 1):
        rows.append(
            PrepRow(
                row_number=row_number,
                area=ws.cell(row=row_number, column=1).value,
                product_code=ws.cell(row=row_number, column=2).value,
                product_name=ws.cell(row=row_number, column=3).value,
                value_per_unit=ws.cell(row=row_number, column=8).value,
            )
        )
    return rows, last_row


def ensure_update_status_header(ws: Worksheet) -> None:
    ws.cell(row=PREP_HEADER_ROW, column=10).value = "Update Status"
    for column in range(1, 11):
        cell = ws.cell(row=PREP_HEADER_ROW, column=column)
        cell.font = Font(name=cell.font.name, size=cell.font.size, bold=True)


def write_statuses(ws: Worksheet, statuses: Sequence[str], start_row: int = PREP_DATA_START_ROW) -> None:
    for index, status in enumerate(statuses):
        ws.cell(row=start_row + index, column=10).value = status
    _autofit_column(ws, 10, start_row, start_row + max(len(statuses), 1))


def ensure_log_header(ws: Worksheet) -> None:
    has_any_header = any(str(ws.cell(row=1, column=index).value or "").strip() for index in range(1, 12))
    if has_any_header:
        return
    for index, header in enumerate(LOG_HEADERS, start=1):
        ws.cell(row=1, column=index).value = header
        cell = ws.cell(row=1, column=index)
        cell.font = Font(name=cell.font.name, size=cell.font.size, bold=True)


def append_log_summary(ws: Worksheet, user: str, message: str, note: str) -> int:
    next_row = last_row_by_col(ws, 1) + 1
    if next_row < 2:
        next_row = 2
    ws.cell(row=next_row, column=1).value = datetime.now()
    ws.cell(row=next_row, column=2).value = user
    ws.cell(row=next_row, column=3).value = "INFO"
    ws.cell(row=next_row, column=4).value = message
    ws.cell(row=next_row, column=11).value = note
    ws.cell(row=next_row, column=1).number_format = "yyyy-mm-dd hh:mm:ss"
    for column in (1, 2, 3, 4, 11):
        _autofit_column(ws, column, next_row, next_row)
    return next_row


def read_note(ws: Worksheet, cell_address: str = "A1") -> str:
    value = ws[cell_address].value
    if value is None:
        return ""
    return str(value).strip()


def read_companies_from_client_config(
    workbook: Workbook,
    sheet_name: str = "konfigurasi-ClientSide",
) -> list[tuple[int, str]]:
    if sheet_name not in workbook.sheetnames:
        return []
    ws = workbook[sheet_name]
    last_row = last_row_by_col(ws, 8)
    if last_row < 4:
        return []

    output: list[tuple[int, str]] = []
    for row_number in range(4, last_row + 1):
        raw_id = ws.cell(row=row_number, column=8).value
        if raw_id is None:
            continue
        try:
            company_id = int(float(raw_id))
        except (TypeError, ValueError):
            continue
        if company_id <= 0:
            continue
        raw_name = ws.cell(row=row_number, column=9).value
        output.append((company_id, str(raw_name).strip() if raw_name is not None else ""))
    return output


def resolve_log_user() -> str:
    username = (os.environ.get("USERNAME") or "").strip()
    if username:
        return username
    return "python-user"


def _autofit_column(ws: Worksheet, col_idx: int, start_row: int, end_row: int) -> None:
    max_len = 0
    for row_number in range(start_row, end_row + 1):
        value = ws.cell(row=row_number, column=col_idx).value
        if value is None:
            continue
        max_len = max(max_len, len(str(value)))
    if max_len == 0:
        return
    column_letter = ws.cell(row=1, column=col_idx).column_letter
    ws.column_dimensions[column_letter].width = max(12, min(max_len + 2, 80))
