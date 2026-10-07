"""Excel parsing for SVL Fix JE template."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from smartscc_tools.features.item_journal.utils import normalize_text
from .config import SHEET_NAME
from .models import SvlFixJeExcelRow


def parse_optional_float(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, str):
        clean = normalize_text(value).replace(",", "")
        if not clean:
            return None
        value = clean
    return float(value)


def parse_svl_id(value: Any) -> int:
    clean = normalize_text(value)
    if not clean:
        raise ValueError("SVL ID kosong")
    return int(float(clean))


def normalize_je_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    clean = normalize_text(value)
    if clean:
        return clean
    return datetime.today().strftime("%Y-%m-%d")


def read_excel(filepath: str, sheet_name: str = SHEET_NAME) -> list[SvlFixJeExcelRow]:
    path = Path(filepath).expanduser().resolve()
    if not path.exists():
        raise FileNotFoundError(f"File tidak ditemukan: {path}")

    workbook = load_workbook(path, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Sheet '{sheet_name}' tidak ditemukan. Sheet yang tersedia: {workbook.sheetnames}"
            )

        worksheet = workbook[sheet_name]
        rows: list[SvlFixJeExcelRow] = []
        for row_idx, cells in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(cell not in (None, "") for cell in cells):
                continue
            if cells[1] in (None, ""):
                continue
            rows.append(
                SvlFixJeExcelRow(
                    row_number=row_idx,
                    svl_id=parse_svl_id(cells[1]),
                    svl_ref=normalize_text(cells[2]),
                    default_code=normalize_text(cells[3]),
                    qty=parse_optional_float(cells[4]),
                    uom=normalize_text(cells[5]),
                    unit_cost=parse_optional_float(cells[6]),
                    total_value=parse_optional_float(cells[7]),
                    coa_credit=normalize_text(cells[8]),
                    coa_debit=normalize_text(cells[9]),
                    journal_code=normalize_text(cells[10]),
                    je_date=normalize_je_date(cells[11]),
                    note=normalize_text(cells[12]),
                )
            )
        return rows
    finally:
        workbook.close()
