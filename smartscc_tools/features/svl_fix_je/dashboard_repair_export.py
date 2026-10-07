"""Export helpers for dashboard repair summary HTML and Excel."""

from __future__ import annotations

from datetime import datetime
from html import escape
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from smartscc_tools.features.item_journal.utils import normalize_text

from .models import SvlDashboardRepairRowResult


def repair_result_effective_mode_label(row: SvlDashboardRepairRowResult) -> str:
    clean_mode = normalize_text(row.selected_target_mode).lower()
    if clean_mode == "fill_existing":
        return "Rewrite Existing JE"
    if clean_mode == "new_and_relink":
        if normalize_text(row.old_move_name) or int(row.old_move_id or 0) > 0:
            return "New JE + Relink SVL"
        return "Created New JE"
    if clean_mode:
        return clean_mode.replace("_", " ").title()
    return "-"


def repair_result_transaction_label(move_name: str, move_id: int) -> str:
    clean_name = normalize_text(move_name)
    if clean_name:
        return clean_name
    return str(int(move_id or 0)) if int(move_id or 0) > 0 else ""


def repair_result_company_label(row: SvlDashboardRepairRowResult) -> str:
    company_name = normalize_text(row.company_name)
    company_id = int(row.company_id or 0)
    if company_name and company_id > 0:
        return f"{company_name} - [{company_id}]"
    if company_name:
        return company_name
    if company_id > 0:
        return f"Company - [{company_id}]"
    return "-"


def repair_result_account_label(code: str, name: str) -> str:
    clean_code = normalize_text(code)
    clean_name = normalize_text(name)
    if clean_code and clean_name:
        return f"{clean_name} [{clean_code}]"
    return clean_name or clean_code or "-"


def repair_result_detail_sentence(row: SvlDashboardRepairRowResult) -> str:
    svl_reference = normalize_text(row.svl_reference) or "-"
    new_move = repair_result_transaction_label(row.move_name, row.move_id) or "-"
    old_move = repair_result_transaction_label(row.old_move_name, row.old_move_id) or "-"
    debit_label = repair_result_account_label(row.debit_account_code, row.debit_account_name)
    credit_label = repair_result_account_label(row.credit_account_code, row.credit_account_name)
    amount_label = f"{abs(float(row.amount or 0.0)):,.2f}"
    effective_mode = repair_result_effective_mode_label(row)
    status_text = normalize_text(row.status).upper()

    if status_text == "ERROR":
        item_bits = [normalize_text(row.item_code), normalize_text(row.item_name)]
        item_label = " - ".join(part for part in item_bits if part) or normalize_text(row.row_key) or "-"
        return (
            f"Gagal memproses Stock Valuation {svl_reference} untuk item {item_label}"
            f" pada mode {effective_mode}: {normalize_text(row.message) or '-'}."
        )

    if effective_mode == "New JE + Relink SVL":
        return (
            f"Dibuat Journal Entry Valuation baru untuk Stock Valuation {svl_reference} "
            f"dari {old_move} karena Debit dan Kredit valuation lama tidak ada / di archive, "
            f"menjadi {new_move}, dengan Debit {debit_label}, dan Kredit {credit_label}, "
            f"senilai {amount_label}."
        )
    if effective_mode == "Created New JE":
        return (
            f"Dibuat Journal Entry Valuation baru untuk Stock Valuation {svl_reference} "
            f"karena sejak awal belum ada Journal Entry yang ter-assign, menjadi {new_move}, "
            f"dengan Debit {debit_label}, dan Kredit {credit_label}, senilai {amount_label}."
        )
    if effective_mode == "Rewrite Existing JE":
        return (
            f"Journal Entry Valuation existing {new_move} diperbarui untuk Stock Valuation {svl_reference}, "
            f"dengan Debit {debit_label}, dan Kredit {credit_label}, senilai {amount_label}."
        )
    return normalize_text(row.message) or "-"


def build_repair_summary_export_payload(
    *,
    database: str,
    results: list[SvlDashboardRepairRowResult],
    exported_at: str | None = None,
) -> dict[str, Any]:
    exported_at = normalize_text(exported_at) or datetime.now().isoformat(timespec="seconds")
    success_count = sum(1 for row in results if normalize_text(row.status).upper() != "ERROR")
    error_count = sum(1 for row in results if normalize_text(row.status).upper() == "ERROR")
    posted_count = sum(1 for row in results if bool(row.posted))

    company_labels: list[str] = []
    mode_counts: dict[str, int] = {}
    detail_rows: list[dict[str, Any]] = []

    for index, row in enumerate(results, start=1):
        company_label = repair_result_company_label(row)
        if company_label not in company_labels:
            company_labels.append(company_label)
        effective_mode = repair_result_effective_mode_label(row)
        mode_counts[effective_mode] = int(mode_counts.get(effective_mode, 0)) + 1
        detail_rows.append(
            {
                "no": index,
                "company_label": company_label,
                "company_name": normalize_text(row.company_name),
                "company_id": int(row.company_id or 0),
                "status": normalize_text(row.status).upper() or "-",
                "effective_mode": effective_mode,
                "transaction_no": repair_result_transaction_label(row.move_name, row.move_id) or "-",
                "old_transaction_no": repair_result_transaction_label(row.old_move_name, row.old_move_id),
                "svl_reference": normalize_text(row.svl_reference) or "-",
                "amount": abs(float(row.amount or 0.0)),
                "amount_label": f"{abs(float(row.amount or 0.0)):,.2f}",
                "item_code": normalize_text(row.item_code),
                "item_name": normalize_text(row.item_name),
                "item_label": " - ".join(part for part in (normalize_text(row.item_code), normalize_text(row.item_name)) if part) or "-",
                "effective_date": normalize_text(row.effective_date) or "-",
                "reference": normalize_text(row.reference),
                "repair_source_kind": normalize_text(row.repair_source_kind),
                "repair_source_label": normalize_text(row.repair_source_label) or "-",
                "debit_account_code": normalize_text(row.debit_account_code),
                "debit_account_name": normalize_text(row.debit_account_name),
                "credit_account_code": normalize_text(row.credit_account_code),
                "credit_account_name": normalize_text(row.credit_account_name),
                "debit_account_label": repair_result_account_label(row.debit_account_code, row.debit_account_name),
                "credit_account_label": repair_result_account_label(row.credit_account_code, row.credit_account_name),
                "detail_sentence": repair_result_detail_sentence(row),
                "posted": bool(row.posted),
                "error_kind": normalize_text(row.error_kind),
                "move_id": int(row.move_id or 0),
                "old_move_id": int(row.old_move_id or 0),
                "svl_id": int(row.relinked_svl_id or 0),
                "raw_message": normalize_text(row.message),
            }
        )

    return {
        "database": normalize_text(database),
        "exported_at": exported_at,
        "company_label": company_labels[0] if len(company_labels) == 1 else " / ".join(company_labels) or "-",
        "company_labels": company_labels,
        "total_rows": len(results),
        "success_count": success_count,
        "error_count": error_count,
        "posted_count": posted_count,
        "mode_counts": [{"mode": mode, "count": count} for mode, count in mode_counts.items()],
        "rows": detail_rows,
    }


def export_repair_summary_excel(
    *,
    database: str,
    results: list[SvlDashboardRepairRowResult],
    output_path: str,
    exported_at: str | None = None,
) -> Path:
    payload = build_repair_summary_export_payload(database=database, results=results, exported_at=exported_at)
    path = Path(output_path).expanduser().resolve()
    workbook = Workbook()
    try:
        summary_sheet = workbook.active
        summary_sheet.title = "Summary"
        summary_sheet.append(["Company", payload["company_label"]])
        summary_sheet.append(["Database", payload["database"] or "-"])
        summary_sheet.append(["Exported At", payload["exported_at"]])
        summary_sheet.append(["Total Transaksi", payload["total_rows"]])
        summary_sheet.append(["Berhasil", payload["success_count"]])
        summary_sheet.append(["Gagal", payload["error_count"]])
        summary_sheet.append(["Posted", payload["posted_count"]])
        summary_sheet.append([])
        summary_sheet.append(["Mode", "Jumlah Transaksi"])
        for mode_row in payload["mode_counts"]:
            summary_sheet.append([mode_row["mode"], mode_row["count"]])

        detail_sheet = workbook.create_sheet("Detail Flat")
        detail_sheet.append(
            [
                "No",
                "Company",
                "Company ID",
                "Database",
                "Effective Mode",
                "Status",
                "Tanggal",
                "Nominal",
                "SVL Reference",
                "Old JE No",
                "New JE No",
                "Item Code",
                "Item Name",
                "Repair Source",
                "Repair Reference",
                "Debit Account Code",
                "Debit Account Name",
                "Credit Account Code",
                "Credit Account Name",
                "Detail",
                "Posted",
                "Error Kind",
                "Move ID",
                "Old Move ID",
                "Relinked SVL ID",
            ]
        )
        for row in payload["rows"]:
            detail_sheet.append(
                [
                    row["no"],
                    row["company_name"] or row["company_label"],
                    row["company_id"],
                    payload["database"] or "-",
                    row["effective_mode"],
                    row["status"],
                    row["effective_date"],
                    row["amount"],
                    row["svl_reference"],
                    row["old_transaction_no"] or "",
                    row["transaction_no"],
                    row["item_code"],
                    row["item_name"],
                    row["repair_source_label"],
                    row["reference"],
                    row["debit_account_code"],
                    row["debit_account_name"],
                    row["credit_account_code"],
                    row["credit_account_name"],
                    row["detail_sentence"],
                    "Yes" if row["posted"] else "No",
                    row["error_kind"],
                    row["move_id"],
                    row["old_move_id"],
                    row["svl_id"],
                ]
            )

        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
        return path
    finally:
        workbook.close()


def _pcb_case1_render_int_list(value: Any) -> str:
    if not isinstance(value, list):
        return ""
    clean_values = [str(int(item or 0)) for item in value if int(item or 0) > 0]
    return ", ".join(clean_values)


def _pcb_case1_render_text_list(value: Any) -> str:
    if not isinstance(value, list):
        return ""
    clean_values = [normalize_text(item) for item in value if normalize_text(item)]
    return ", ".join(clean_values)


def _pcb_case1_render_mapping(value: Any) -> str:
    if isinstance(value, dict):
        return str(value) if value else ""
    return normalize_text(value)


def _pcb_planned_line_account_codes(value: Any, *, side: str) -> str:
    if not isinstance(value, list):
        return ""
    clean_side = normalize_text(side).lower()
    seen: set[str] = set()
    codes: list[str] = []
    for entry in value:
        if not isinstance(entry, dict):
            entry = {
                field_name: getattr(entry, field_name)
                for field_name in getattr(entry, "__dataclass_fields__", {})
            }
        if normalize_text(entry.get("side")).lower() != clean_side:
            continue
        account_code = normalize_text(entry.get("account_code")).upper()
        if not account_code or account_code in seen:
            continue
        seen.add(account_code)
        codes.append(account_code)
    return ", ".join(codes)


def build_pcb_case1_export_payload(
    *,
    database: str,
    rows: list[dict[str, Any]],
    exported_at: str | None = None,
) -> dict[str, Any]:
    exported_at = normalize_text(exported_at) or datetime.now().isoformat(timespec="seconds")
    detail_rows: list[dict[str, Any]] = []
    company_labels: list[str] = []

    for index, row in enumerate(rows, start=1):
        company_id = int(row.get("company_id") or 0)
        company_name = normalize_text(row.get("company_name"))
        if company_name and company_id > 0:
            company_label = f"{company_name} - [{company_id}]"
        elif company_name:
            company_label = company_name
        elif company_id > 0:
            company_label = f"Company - [{company_id}]"
        else:
            company_label = "-"
        if company_label not in company_labels:
            company_labels.append(company_label)
        amount = abs(float(row.get("amount") or 0.0))
        debit_account_code = normalize_text(row.get("debit_account_code")) or _pcb_planned_line_account_codes(
            row.get("planned_lines"),
            side="debit",
        )
        credit_account_code = normalize_text(row.get("credit_account_code")) or _pcb_planned_line_account_codes(
            row.get("planned_lines"),
            side="credit",
        )
        detail_rows.append(
            {
                "no": index,
                "company_label": company_label,
                "company_name": company_name,
                "company_id": company_id,
                "cycle_key": normalize_text(row.get("cycle_key")),
                "picking_name": normalize_text(row.get("picking_name")),
                "source_label": normalize_text(row.get("source_label")),
                "po_name": normalize_text(row.get("po_name")),
                "bill_name": normalize_text(row.get("bill_name")),
                "stj_refs": _pcb_case1_render_text_list(row.get("stj_refs")),
                "date": normalize_text(row.get("date")),
                "reference": normalize_text(row.get("reference")),
                "journal_code": normalize_text(row.get("journal_code")),
                "amount": amount,
                "amount_label": f"{amount:,.2f}",
                "debit_account_code": debit_account_code,
                "credit_account_code": credit_account_code,
                "resolve_account_code": normalize_text(row.get("resolve_account_code")),
                "resolve_account_preview": normalize_text(row.get("resolve_account_preview")),
                "result_status": normalize_text(row.get("result_status")) or normalize_text(row.get("row_status")).upper(),
                "result_move_id": int(row.get("result_move_id") or 0),
                "result_move_name": normalize_text(row.get("result_move_name")),
                "result_posted": bool(row.get("result_posted")),
                "existing_move_detected": bool(row.get("existing_move_detected")),
                "result_error_kind": normalize_text(row.get("result_error_kind")),
                "reconcile_ready": bool(row.get("reconcile_ready")),
                "reconcile_readiness_label": normalize_text(row.get("reconcile_readiness_label")),
                "reconcile_attempted": bool(row.get("reconcile_attempted")),
                "reconcile_performed": bool(row.get("reconcile_performed")),
                "reconcile_skipped": bool(row.get("reconcile_skipped")),
                "reconcile_message": normalize_text(row.get("reconcile_message")),
                "reconcile_error_kind": normalize_text(row.get("reconcile_error_kind")),
                "row_status": normalize_text(row.get("row_status")),
                "row_status_message": normalize_text(row.get("row_status_message")),
                "review_required": bool(row.get("review_required")),
                "review_confirmed": bool(row.get("review_confirmed")),
                "review_reason": normalize_text(row.get("review_reason")),
                "product_id": int(row.get("product_id") or 0),
                "item_code": normalize_text(row.get("item_code")),
                "item_name": normalize_text(row.get("item_name")),
                "item_category_name": normalize_text(row.get("item_category_name")),
                "bill_line_id": int(row.get("bill_line_id") or 0),
                "bill_move_id": int(row.get("bill_move_id") or 0),
                "purchase_line_id": int(row.get("purchase_line_id") or 0),
                "stock_move_id": int(row.get("stock_move_id") or 0),
                "stock_move_ids": _pcb_case1_render_int_list(row.get("stock_move_ids")),
                "stj_move_ids": _pcb_case1_render_int_list(row.get("stj_move_ids")),
                "picking_id": int(row.get("picking_id") or 0),
                "payment_move_ids": _pcb_case1_render_int_list(row.get("payment_move_ids")),
                "bank_move_ids": _pcb_case1_render_int_list(row.get("bank_move_ids")),
                "suspend_target_aml_ids": _pcb_case1_render_int_list(row.get("suspend_target_aml_ids")),
                "clearing_target_aml_ids": _pcb_case1_render_int_list(row.get("clearing_target_aml_ids")),
                "partner_id": int(row.get("partner_id") or 0),
                "partner_name": normalize_text(row.get("partner_name")),
                "currency_id": int(row.get("currency_id") or 0),
                "product_uom_id": int(row.get("product_uom_id") or 0),
                "quantity": float(row.get("quantity") or 0.0),
                "amount_currency": float(row.get("amount_currency") or 0.0),
                "analytic_distribution": _pcb_case1_render_mapping(row.get("analytic_distribution")),
                "bill_price_unit": float(row.get("bill_price_unit") or 0.0),
                "gr_price_unit": float(row.get("gr_price_unit") or 0.0),
                "price_gap_value": float(row.get("price_gap_value") or 0.0),
                "allocated_amount": float(row.get("allocated_amount") or 0.0),
            }
        )

    return {
        "database": normalize_text(database),
        "exported_at": exported_at,
        "company_label": company_labels[0] if len(company_labels) == 1 else " / ".join(company_labels) or "-",
        "company_labels": company_labels,
        "total_rows": len(detail_rows),
        "created_count": sum(1 for row in detail_rows if int(row["result_move_id"]) > 0 and not row["existing_move_detected"]),
        "existing_count": sum(1 for row in detail_rows if row["existing_move_detected"]),
        "posted_count": sum(1 for row in detail_rows if row["result_posted"]),
        "reconciled_count": sum(1 for row in detail_rows if row["reconcile_performed"]),
        "reconcile_skipped_count": sum(1 for row in detail_rows if row["reconcile_skipped"]),
        "error_count": sum(
            1
            for row in detail_rows
            if not bool(row["result_posted"])
            and int(row["result_move_id"]) <= 0
            and (
                normalize_text(row["result_status"]).upper() == "ERROR"
                or normalize_text(row["row_status"]).lower() == "error"
            )
        ),
        "rows": detail_rows,
    }


def export_pcb_case1_summary_excel(
    *,
    database: str,
    rows: list[dict[str, Any]],
    output_path: str,
    exported_at: str | None = None,
) -> Path:
    payload = build_pcb_case1_export_payload(database=database, rows=rows, exported_at=exported_at)
    path = Path(output_path).expanduser().resolve()
    workbook = Workbook()
    try:
        summary_sheet = workbook.active
        summary_sheet.title = "Summary"
        summary_sheet.append(["Company", payload["company_label"]])
        summary_sheet.append(["Database", payload["database"] or "-"])
        summary_sheet.append(["Exported At", payload["exported_at"]])
        summary_sheet.append(["Total Rows", payload["total_rows"]])
        summary_sheet.append(["JE Created", payload["created_count"]])
        summary_sheet.append(["Posted", payload["posted_count"]])
        summary_sheet.append(["Reconciled", payload["reconciled_count"]])
        summary_sheet.append(["Reconcile Skipped", payload["reconcile_skipped_count"]])
        summary_sheet.append(["Error", payload["error_count"]])
        summary_sheet.append(["Existing Reused", payload["existing_count"]])

        detail_sheet = workbook.create_sheet("Detail Flat")
        detail_sheet.append(
            [
                "No",
                "Company",
                "Company ID",
                "Database",
                "Cycle Key",
                "Picking",
                "Source",
                "PO",
                "Bill",
                "Related STJ",
                "Date",
                "Reference",
                "Journal Code",
                "Result Status",
                "Result JE No",
                "Result JE ID",
                "Posted",
                "Amount",
                "Debit Account Code",
                "Credit Account Code",
                "Resolve Account Code",
                "Resolve Account Preview",
                "Reconcile Ready",
                "Reconcile Readiness",
                "Reconcile Attempted",
                "Reconcile Performed",
                "Reconcile Skipped",
                "Reconcile Message",
                "Row Status",
                "Row Status Message",
                "Review Required",
                "Review Confirmed",
                "Review Reason",
                "Product ID",
                "Item Code",
                "Item Name",
                "Item Category",
                "Bill Line ID",
                "Bill Move ID",
                "Purchase Line ID",
                "Stock Move ID",
                "Stock Move IDs",
                "STJ Move IDs",
                "Picking ID",
                "Payment Move IDs",
                "Bank Move IDs",
                "Suspend Target AML IDs",
                "Clearing Target AML IDs",
                "Partner ID",
                "Partner Name",
                "Currency ID",
                "UoM ID",
                "Quantity",
                "Amount Currency",
                "Analytic Distribution",
                "Bill Price Unit",
                "GR Price Unit",
                "Price Gap Value",
                "Allocated Amount",
                "Result Error Kind",
                "Existing Move Detected",
                "Reconcile Error Kind",
            ]
        )
        for row in payload["rows"]:
            detail_sheet.append(
                [
                    row["no"],
                    row["company_name"] or row["company_label"],
                    row["company_id"],
                    payload["database"] or "-",
                    row["cycle_key"],
                    row["picking_name"],
                    row["source_label"],
                    row["po_name"],
                    row["bill_name"],
                    row["stj_refs"],
                    row["date"],
                    row["reference"],
                    row["journal_code"],
                    row["result_status"],
                    row["result_move_name"],
                    row["result_move_id"],
                    "Yes" if row["result_posted"] else "No",
                    row["amount"],
                    row["debit_account_code"],
                    row["credit_account_code"],
                    row["resolve_account_code"],
                    row["resolve_account_preview"],
                    "Yes" if row["reconcile_ready"] else "No",
                    row["reconcile_readiness_label"],
                    "Yes" if row["reconcile_attempted"] else "No",
                    "Yes" if row["reconcile_performed"] else "No",
                    "Yes" if row["reconcile_skipped"] else "No",
                    row["reconcile_message"],
                    row["row_status"],
                    row["row_status_message"],
                    "Yes" if row["review_required"] else "No",
                    "Yes" if row["review_confirmed"] else "No",
                    row["review_reason"],
                    row["product_id"],
                    row["item_code"],
                    row["item_name"],
                    row["item_category_name"],
                    row["bill_line_id"],
                    row["bill_move_id"],
                    row["purchase_line_id"],
                    row["stock_move_id"],
                    row["stock_move_ids"],
                    row["stj_move_ids"],
                    row["picking_id"],
                    row["payment_move_ids"],
                    row["bank_move_ids"],
                    row["suspend_target_aml_ids"],
                    row["clearing_target_aml_ids"],
                    row["partner_id"],
                    row["partner_name"],
                    row["currency_id"],
                    row["product_uom_id"],
                    row["quantity"],
                    row["amount_currency"],
                    row["analytic_distribution"],
                    row["bill_price_unit"],
                    row["gr_price_unit"],
                    row["price_gap_value"],
                    row["allocated_amount"],
                    row["result_error_kind"],
                    "Yes" if row["existing_move_detected"] else "No",
                    row["reconcile_error_kind"],
                ]
            )

        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
        return path
    finally:
        workbook.close()


def export_repair_summary_html(
    *,
    database: str,
    results: list[SvlDashboardRepairRowResult],
    output_path: str,
    exported_at: str | None = None,
) -> Path:
    payload = build_repair_summary_export_payload(database=database, results=results, exported_at=exported_at)
    html = _build_repair_summary_html(payload)
    path = Path(output_path).expanduser().resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(html, encoding="utf-8")
    return path


def _build_repair_summary_html(payload: dict[str, Any]) -> str:
    mode_rows = "".join(
        f"<tr><td>{escape(str(row['mode']))}</td><td class='r'>{int(row['count'])}</td></tr>"
        for row in payload["mode_counts"]
    ) or "<tr><td colspan='2'>Tidak ada transaksi.</td></tr>"
    detail_rows = "".join(
        (
            "<tr>"
            f"<td>{int(row['no'])}</td>"
            f"<td>{escape(str(row['status']))}</td>"
            f"<td>{escape(str(row['effective_mode']))}</td>"
            f"<td>{escape(str(row['transaction_no']))}</td>"
            f"<td>{escape(str(row['old_transaction_no'] or '-'))}</td>"
            f"<td>{escape(str(row['svl_reference']))}</td>"
            f"<td class='r'>{escape(str(row['amount_label']))}</td>"
            f"<td>{escape(str(row['item_label']))}</td>"
            f"<td>{escape(str(row['effective_date']))}</td>"
            f"<td>{escape(str(row['debit_account_label']))}</td>"
            f"<td>{escape(str(row['credit_account_label']))}</td>"
            f"<td>{escape(str(row['detail_sentence']))}</td>"
            "</tr>"
        )
        for row in payload["rows"]
    ) or "<tr><td colspan='12'>Tidak ada detail repair.</td></tr>"
    return f"""<!DOCTYPE html>
<html lang="id">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Repair Summary Export</title>
<style>
body {{ font-family: Segoe UI, Arial, sans-serif; background: #f4f6f8; color: #18212b; margin: 0; }}
.page {{ max-width: 1680px; margin: 0 auto; padding: 24px; }}
.card {{ background: #fff; border: 1px solid #d8dde3; border-radius: 12px; padding: 18px; margin-bottom: 18px; }}
.banner {{ background: #0f9fa3; color: #fff; }}
.meta {{ display: flex; flex-wrap: wrap; gap: 12px; margin-top: 10px; font-size: 13px; }}
.pill {{ background: rgba(255,255,255,.14); border-radius: 999px; padding: 6px 10px; }}
.kpis {{ display: grid; grid-template-columns: repeat(4, minmax(180px, 1fr)); gap: 12px; }}
.kpi {{ background: #fff; border: 1px solid #d8dde3; border-radius: 12px; padding: 14px; }}
.kpi-label {{ font-size: 11px; text-transform: uppercase; color: #607080; }}
.kpi-value {{ margin-top: 6px; font-size: 24px; font-weight: 700; }}
table {{ width: 100%; border-collapse: collapse; font-size: 12px; }}
th, td {{ border-bottom: 1px solid #e6eaef; padding: 8px 10px; text-align: left; vertical-align: top; }}
th {{ background: #f7f9fb; font-size: 11px; text-transform: uppercase; color: #607080; }}
.r {{ text-align: right; }}
</style>
</head>
<body>
<div class="page">
  <div class="card banner">
    <h1 style="margin:0;">Repair Summary Export</h1>
    <div class="meta">
      <span class="pill">{escape(str(payload['company_label']))}</span>
      <span class="pill">{escape(str(payload['database'] or '-'))}</span>
      <span class="pill">Exported: {escape(str(payload['exported_at']))}</span>
    </div>
  </div>
  <div class="kpis">
    <div class="kpi"><div class="kpi-label">Total Transaksi</div><div class="kpi-value">{int(payload['total_rows'])}</div></div>
    <div class="kpi"><div class="kpi-label">Berhasil</div><div class="kpi-value">{int(payload['success_count'])}</div></div>
    <div class="kpi"><div class="kpi-label">Gagal</div><div class="kpi-value">{int(payload['error_count'])}</div></div>
    <div class="kpi"><div class="kpi-label">Posted</div><div class="kpi-value">{int(payload['posted_count'])}</div></div>
  </div>
  <div class="card">
    <h2 style="margin-top:0;">Jumlah Transaksi per Mode</h2>
    <table>
      <thead><tr><th>Mode</th><th class="r">Jumlah</th></tr></thead>
      <tbody>{mode_rows}</tbody>
    </table>
  </div>
  <div class="card">
    <h2 style="margin-top:0;">Detail Repair</h2>
    <table>
      <thead>
        <tr>
          <th>No</th>
          <th>Status</th>
          <th>Mode</th>
          <th>No JE Baru</th>
          <th>No JE Lama</th>
          <th>No SVL</th>
          <th class="r">Nominal</th>
          <th>Item</th>
          <th>Tanggal</th>
          <th>Akun Debit</th>
          <th>Akun Kredit</th>
          <th>Detail</th>
        </tr>
      </thead>
      <tbody>{detail_rows}</tbody>
    </table>
  </div>
</div>
</body>
</html>
"""
